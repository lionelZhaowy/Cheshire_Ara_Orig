#!/usr/bin/env python3
"""Structured SVG -> native, editable PPTX. Unsupported features fail explicitly."""
import argparse
from collections import Counter
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET
import zipfile

NS = "http://www.w3.org/2000/svg"
ET.register_namespace("", NS)
TOKEN = r"[MmLlHhVvCcQqZz]|[+-]?(?:\d*\.\d+|\d+\.?)(?:[eE][+-]?\d+)?"
DEFAULTS = {"fill": "#000000", "stroke": "none", "stroke-width": "1",
            "font-family": "Arial", "font-size": "16", "font-weight": "normal",
            "font-style": "normal", "text-anchor": "start"}
UNSUPPORTED = {"style", "class", "transform", "filter", "clip-path", "mask",
               "marker-start", "marker-mid", "marker-end", "stroke-dasharray",
               "stroke-dashoffset", "textLength", "lengthAdjust", "rotate", "dx", "dy",
               "visibility", "display", "vector-effect", "stroke-miterlimit",
               "dominant-baseline", "alignment-baseline", "letter-spacing", "word-spacing",
               "baseline-shift", "writing-mode", "font-stretch", "font-variant"}


def number(value):
    value = str(value).strip()
    if value.endswith("px"):
        value = value[:-2]
    result = float(value)
    if not math.isfinite(result):
        raise ValueError("Non-finite coordinate")
    return result


def colour(value):
    if value == "none":
        return value
    if re.fullmatch(r"#[0-9a-fA-F]{3}", value):
        return "#" + "".join(c * 2 for c in value[1:])
    if not re.fullmatch(r"#[0-9a-fA-F]{6}", value):
        raise ValueError(f"Use #RGB/#RRGGBB/none, got {value!r}")
    return value


def path_commands(data):
    if re.sub(TOKEN, "", data).strip(" ,\t\r\n"):
        raise ValueError("Path supports M/L/H/V/C/Q/Z only; simplify unsupported commands")
    tokens = re.findall(TOKEN, data)
    commands, position, start = [], (0.0, 0.0), None
    index, command = 0, None
    while index < len(tokens):
        if tokens[index].isalpha():
            command = tokens[index]
            index += 1
        if command is None:
            raise ValueError("Path missing command")
        kind, relative = command.upper(), command.islower()
        if kind == "Z":
            if start is None:
                raise ValueError("Close before move")
            commands.append(("Z", []))
            position, command = start, None
            continue
        size = {"M": 2, "L": 2, "H": 1, "V": 1, "C": 6, "Q": 4}[kind]
        if index + size > len(tokens) or any(t.isalpha() for t in tokens[index:index+size]):
            raise ValueError("Incomplete path command")
        values = [number(t) for t in tokens[index:index+size]]
        index += size
        if kind in {"H", "V"}:
            value = values[0]
            if kind == "H":
                points = [(value + (position[0] if relative else 0), position[1])]
            else:
                points = [(position[0], value + (position[1] if relative else 0))]
            kind = "L"
        else:
            points = list(zip(values[::2], values[1::2]))
            if relative:
                points = [(x + position[0], y + position[1]) for x, y in points]
        if kind == "M":
            if start is not None:
                raise ValueError("Compound path not supported; use independent elements")
            start = points[0]
            command = "l" if relative else "L"
        elif start is None:
            raise ValueError("Path must start with M")
        if kind == "Q":
            q, end = points
            points = [(position[0] + 2*(q[0]-position[0])/3,
                       position[1] + 2*(q[1]-position[1])/3),
                      (end[0] + 2*(q[0]-end[0])/3, end[1] + 2*(q[1]-end[1])/3), end]
            kind = "C"
        commands.append((kind, points))
        position = points[-1]
    if not commands:
        raise ValueError("Empty path")
    return commands


def read_svg(path, font_override=None):
    root = ET.parse(path).getroot()
    if root.tag != f"{{{NS}}}svg":
        raise ValueError(f"{path}: missing SVG namespace")
    view = [number(t) for t in re.split(r"[ ,]+", root.get("viewBox", "").strip()) if t]
    if not view:
        view = [0, 0, number(root.get("width")), number(root.get("height"))]
    if len(view) != 4 or view[2] <= 0 or view[3] <= 0:
        raise ValueError("Invalid viewBox")
    normal = ET.Element(f"{{{NS}}}svg", {"version": "1.1", "width": str(view[2]),
        "height": str(view[3]), "viewBox": " ".join(map(str, view))})
    items = []

    def visit(node, inherited):
        tag = node.tag.split("}")[-1]
        if tag in {"title", "desc"}:
            return
        label = node.get("id", tag)
        def fail(reason):
            raise ValueError(f"{path.name} / {label}: {reason}")
        bad = UNSUPPORTED.intersection(node.attrib)
        if bad:
            fail(f"Unsupported {sorted(bad)}; normalize to explicit primitive attributes")
        if node.get("fill-rule", "nonzero") != "nonzero":
            fail("Only nonzero fill rule supported")
        for key in ["opacity", "fill-opacity", "stroke-opacity"]:
            if key in node.attrib and number(node.get(key)) != 1:
                fail("Transparency not yet supported; explicitly choose opaque colours")
        if node.get("stroke-linecap", "butt") != "butt" or node.get("stroke-linejoin", "miter") != "miter":
            fail("Only butt line caps and miter joins supported")
        paint = {**inherited, **{k: v for k, v in node.attrib.items() if k in DEFAULTS}}
        if tag in {"svg", "g"}:
            for child in node:
                visit(child, paint)
            return
        if tag not in {"rect", "circle", "ellipse", "line", "polyline", "polygon", "path", "text"}:
            fail(f"Unsupported element {tag}")
        if tag in {"line", "polyline"}:
            paint["fill"] = "none"
        if list(node):
            fail("Nested visual elements/tspan are not supported; use independent text lines")
        a = {**paint, **node.attrib}
        a["fill"], a["stroke"] = colour(a["fill"]), colour(a["stroke"])
        if number(a["stroke-width"]) < 0:
            fail("Negative stroke width")
        commands = None
        if tag == "rect":
            a.setdefault("x", "0"); a.setdefault("y", "0")
            for k in ["width", "height"]:
                if number(a[k]) <= 0:
                    fail("Rectangle dimensions must be positive")
            rx = number(a.get("rx", a.get("ry", "0")))
            ry = number(a.get("ry", a.get("rx", "0")))
            if rx != ry or rx < 0:
                fail("Only nonnegative equal rx/ry supported")
            a["rx"] = a["ry"] = str(min(rx, number(a["width"])/2, number(a["height"])/2))
        elif tag in {"circle", "ellipse"}:
            a.setdefault("cx", "0"); a.setdefault("cy", "0")
            for k in (["r"] if tag == "circle" else ["rx", "ry"]):
                if number(a[k]) <= 0:
                    fail("Ellipse radii must be positive")
        elif tag == "line":
            for k in ["x1", "y1", "x2", "y2"]:
                a.setdefault(k, "0")
        elif tag == "path":
            commands = path_commands(a["d"])
        elif tag in {"polygon", "polyline"}:
            values = [number(t) for t in re.split(r"[ ,\s]+", a["points"].strip()) if t]
            if len(values) < 4 or len(values) % 2:
                fail("Invalid point list")
            pts = list(zip(values[::2], values[1::2]))
            commands = [("M", pts[:1])] + [("L", [p]) for p in pts[1:]]
            if tag == "polygon":
                commands.append(("Z", []))
        elif tag == "text":
            if not node.text or number(a["font-size"]) <= 0:
                fail("Text must be nonempty with positive font-size")
            if "\n" in node.text or "\r" in node.text:
                fail("Use separate text elements for multiline text")
            a.setdefault("x", "0"); a.setdefault("y", "0")
            if a["text-anchor"] not in {"start", "middle", "end"}:
                fail("Unsupported text-anchor")
            if a["font-style"] not in {"normal", "italic"}:
                fail("Unsupported font-style")
            if a["font-weight"] not in {"normal", "bold"}:
                if not a["font-weight"].isdigit():
                    fail("Unsupported font-weight")
                a["font-weight"] = "bold" if int(a["font-weight"]) >= 600 else "normal"
            if a["stroke"] != "none" or a["fill"] == "none":
                fail("Text must have a fill and no outline")
            if font_override:
                a["font-family"] = font_override
            else:
                a["font-family"] = a["font-family"].split(",")[0].strip(" '\"")
        a["id"] = f"{tag}-{len(items)+1:03d}"
        for key in ["fill", "stroke"]:
            a[key] = colour(a[key])
        elem = ET.SubElement(normal, f"{{{NS}}}{tag}", a)
        elem.text = node.text if tag == "text" else None
        items.append({"tag": tag, "attrs": a, "text": elem.text, "commands": commands})
    visit(root, DEFAULTS)
    if not items:
        raise ValueError("No visual objects")
    ET.indent(normal, space="  ")
    return {"source": path, "view": view, "svg": normal, "items": items}


def create_deck(figures, destination, aspect, font_file=None):
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.oxml.ns import qn
        from PIL import ImageFont
    except ImportError as error:
        raise RuntimeError("Install requirements.txt in an isolated Python environment") from error
    prs = Presentation()
    prs.slide_width = 12801600  # 14 inches
    ratio = {"16:9": 9/16, "4:3": 3/4}.get(aspect)
    if ratio is None:
        ratio = max(f["view"][3]/f["view"][2] for f in figures)
    prs.slide_height = round(prs.slide_width * ratio)
    reports = []
    for figure in figures:
        x0, y0, width, height = figure["view"]
        scale = min(prs.slide_width/width, prs.slide_height/height)
        ox = (prs.slide_width - width*scale)/2 - x0*scale
        oy = (prs.slide_height - height*scale)/2 - y0*scale
        def length(v): return round(number(v)*scale)
        def x(v): return round(number(v)*scale + ox)
        def y(v): return round(number(v)*scale + oy)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        sizes = []

        def geometry(commands):
            points = [p for _, pts in commands for p in pts]
            lx, ly = min(p[0] for p in points), min(p[1] for p in points)
            w = max(1, length(max(p[0] for p in points)-lx))
            h = max(1, length(max(p[1] for p in points)-ly))
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x(lx), y(ly), w, h)
            sppr = shape._element.spPr
            sppr.remove(sppr.prstGeom)
            geom = OxmlElement("a:custGeom")
            for name in ["avLst", "gdLst", "ahLst", "cxnLst"]:
                geom.append(OxmlElement(f"a:{name}"))
            rect = OxmlElement("a:rect")
            for k, v in {"l":"0", "t":"0", "r":"r", "b":"b"}.items(): rect.set(k,v)
            geom.append(rect)
            paths, path = OxmlElement("a:pathLst"), OxmlElement("a:path")
            path.set("w",str(w)); path.set("h",str(h))
            for kind, pts in commands:
                cmd = OxmlElement("a:" + {"M":"moveTo", "L":"lnTo", "C":"cubicBezTo", "Z":"close"}[kind])
                for px, py in pts:
                    point = OxmlElement("a:pt")
                    point.set("x",str(length(px-lx))); point.set("y",str(length(py-ly)))
                    cmd.append(point)
                path.append(cmd)
            paths.append(path); geom.append(paths); sppr.insert(1,geom)
            return shape

        for item in figure["items"]:
            tag, a = item["tag"], item["attrs"]
            if tag == "rect":
                radius = number(a["rx"])
                kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
                shape = slide.shapes.add_shape(kind, x(a["x"]), y(a["y"]), length(a["width"]), length(a["height"]))
                if radius: shape.adjustments[0] = radius/min(number(a["width"]), number(a["height"]))
            elif tag in {"circle", "ellipse"}:
                rx = number(a["r"] if tag == "circle" else a["rx"])
                ry = number(a["r"] if tag == "circle" else a["ry"])
                shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, x(number(a["cx"])-rx), y(number(a["cy"])-ry), length(2*rx), length(2*ry))
            elif tag == "line":
                shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x(a["x1"]), y(a["y1"]), x(a["x2"]), y(a["y2"]))
            elif tag in {"path", "polygon", "polyline"}:
                shape = geometry(item["commands"])
            else:
                size = number(a["font-size"])
                candidate = font_file
                if not candidate and a["font-family"] == "Microsoft YaHei":
                    name = "msyhbd.ttc" if a["font-weight"] == "bold" else "msyh.ttc"
                    fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" if os.name == "nt" else Path("/mnt/c/Windows/Fonts")
                    candidate = fonts / name
                if candidate and Path(candidate).is_file():
                    advance = ImageFont.truetype(str(candidate), round(size*10)).getlength(item["text"])/10
                else:
                    advance = sum(size if ord(c)>255 else size*.65 for c in item["text"])
                box_width = advance + size
                anchor = a["text-anchor"]
                left = number(a["x"]) - ({"start":0,"middle":.5,"end":1}[anchor])*box_width
                shape = slide.shapes.add_textbox(x(left), y(number(a["y"])-size*1.08), length(box_width), length(size*1.6))
                tf = shape.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.TOP
                p = tf.paragraphs[0]
                p.alignment = {"start":PP_ALIGN.LEFT,"middle":PP_ALIGN.CENTER,"end":PP_ALIGN.RIGHT}[anchor]
                p.space_before = p.space_after = 0; p.line_spacing = 1.0
                run = p.add_run(); run.text = item["text"]
                run.font.name = a["font-family"]; run.font.size = length(size)
                run.font.bold = a["font-weight"] == "bold"; run.font.italic = a["font-style"] == "italic"
                run.font.color.rgb = RGBColor.from_string(a["fill"][1:])
                for script in ["ea", "cs"]:
                    node = OxmlElement(f"a:{script}"); node.set("typeface",a["font-family"])
                    run._r.get_or_add_rPr().append(node)
                sizes.append(length(size)/12700)
            shape.name = a["id"] + (": " + item["text"][:40] if tag == "text" else "")
            if tag != "text":
                if hasattr(shape,"fill"):
                    if a["fill"] == "none": shape.fill.background()
                    else:
                        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(a["fill"][1:])
                if a["stroke"] == "none": shape.line.fill.background()
                else:
                    shape.line.color.rgb = RGBColor.from_string(a["stroke"][1:])
                    shape.line.width = length(a["stroke-width"])
            style = shape._element.find(qn("p:style"))
            if style is not None: shape._element.remove(style)
            shape._element.spPr.append(OxmlElement("a:effectLst"))
            if tag != "text" and a["stroke"] != "none":
                shape._element.spPr.get_or_add_ln().set("cap","flat")
        expected = [i["text"] for i in figure["items"] if i["tag"] == "text"]
        reports.append({"objects":len(slide.shapes), "text_objects":len(expected),
                        "min_text_pt":min(sizes) if sizes else None})
    prs.save(destination)
    reopened = Presentation(destination)
    for figure, slide in zip(figures,reopened.slides):
        expected = [i["text"] for i in figure["items"] if i["tag"] == "text"]
        actual = [s.text for s in slide.shapes if s.has_text_frame and s.text]
        if actual != expected or len(slide.shapes) != len(figure["items"]):
            raise RuntimeError("PPTX text/object preservation check failed")
    with zipfile.ZipFile(destination) as package:
        if package.testzip() is not None or any(n.startswith("ppt/media/") for n in package.namelist()):
            raise RuntimeError("Package integrity or unexpected raster/SVG image")
    return reports


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, nargs="+", required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    parser.add_argument("--aspect", choices=["source","16:9","4:3"], default="source")
    parser.add_argument("--font", help="Explicitly substitute all text fonts")
    parser.add_argument("--font-file", type=Path, help="Local font for text width measurement")
    parser.add_argument("--svg-only", action="store_true")
    args = parser.parse_args()
    if len({p.stem for p in args.input}) != len(args.input):
        parser.error("Input basenames must be unique")
    if args.out_dir.exists() and any(args.out_dir.iterdir()):
        parser.error("Output directory is not empty; choose a fresh directory")
    figures = [read_svg(p,args.font) for p in args.input]
    report = {"figures":[], "svg_text_outlined":False, "pptx_created":not args.svg_only,
              "office_application_tested":False, "svg_convert_to_shape_tested":False}
    # Validate the full subset before creating any output directory/files.
    args.out_dir.mkdir(parents=True, exist_ok=True)
    if not args.svg_only:
        report["slides"] = create_deck(figures,args.out_dir/"diagrams.pptx",args.aspect,args.font_file)
    for f in figures:
        target = args.out_dir/(f["source"].stem+".editable.svg")
        ET.ElementTree(f["svg"]).write(target,encoding="utf-8",xml_declaration=True)
        report["figures"].append({"source":str(f["source"]),
            "sha256":hashlib.sha256(f["source"].read_bytes()).hexdigest(),
            "output_svg":target.name,
            "counts":dict(Counter(i["tag"] for i in f["items"]))})
    (args.out_dir/"export_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__ == "__main__":
    try:
        main()
    except (ValueError,KeyError,RuntimeError) as error:
        print(f"ERROR: {error}",file=sys.stderr)
        sys.exit(2)
