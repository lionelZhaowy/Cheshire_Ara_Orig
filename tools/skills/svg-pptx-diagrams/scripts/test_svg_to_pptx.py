#!/usr/bin/env python3
"""Behavioral checks for the exported SVG/PPTX contract."""
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
import svg_to_pptx as tool

ROOT = Path(__file__).resolve().parent.parent


class ExportContract(unittest.TestCase):
    def test_launcher_waits_and_propagates_failure(self):
        with tempfile.TemporaryDirectory(prefix="svg-launcher-contract-") as directory:
            output = Path(directory)/"export"
            args = [sys.executable,str(ROOT/"scripts/run.py"),"--input",
                    str(ROOT/"assets/example.svg"),"--out-dir",str(output)]
            first = subprocess.run(args,capture_output=True)
            self.assertEqual(first.returncode,0,first.stderr)
            self.assertTrue((output/"diagrams.pptx").is_file())
            self.assertTrue((output/"export_report.json").is_file())
            second = subprocess.run(args,capture_output=True)
            self.assertEqual(second.returncode,2,second.stderr)

    def test_relative_and_quadratic_paths(self):
        commands = tool.path_commands("m10 20 h30 v40 l-10 -20 z")
        self.assertEqual(commands, [("M",[(10,20)]),("L",[(40,20)]),
            ("L",[(40,60)]),("L",[(30,40)]),("Z",[])])
        self.assertEqual(tool.path_commands("M0 0 Q3 3 6 0"),
            [("M",[(0,0)]),("C",[(2,2),(4,2),(6,0)])])
        with self.assertRaisesRegex(ValueError,"Compound"):
            tool.path_commands("M0 0 L1 1 M2 2 L3 3")

    def test_unsupported_features_fail_without_outputs(self):
        with tempfile.TemporaryDirectory(prefix="svg-contract-") as directory:
            base = Path(directory)
            for body in [
                '<path d="M0 0 A10 10 0 0 0 20 20"/>',
                '<text x="0" y="20"><tspan>not flattened</tspan></text>',
                '<rect width="10" height="10" transform="rotate(30)"/>',
                '<image href="secret.png"/>',
                '<rect width="10" height="10" filter="url(#shadow)"/>',
            ]:
                source = base/"unsupported.svg"
                source.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100">'+body+'</svg>')
                output = base/"no-output"
                result = subprocess.run([sys.executable,str(ROOT/"scripts/svg_to_pptx.py"),
                    "--input",str(source),"--out-dir",str(output)],capture_output=True)
                self.assertEqual(result.returncode,2,result.stderr)
                self.assertFalse(output.exists())

    def test_native_objects_text_geometry_and_overwrite_guard(self):
        from pptx import Presentation
        from pptx.oxml.ns import qn
        with tempfile.TemporaryDirectory(prefix="svg-native-contract-") as directory:
            output = Path(directory)/"export"
            args = [sys.executable,str(ROOT/"scripts/svg_to_pptx.py"),"--input",
                    str(ROOT/"assets/example.svg"),"--out-dir",str(output)]
            result = subprocess.run(args,capture_output=True)
            self.assertEqual(result.returncode,0,result.stderr)
            report = json.loads((output/"export_report.json").read_text())
            self.assertEqual(report["slides"][0]["objects"],19)
            figure = tool.read_svg(ROOT/"assets/example.svg")
            expected = [i["text"] for i in figure["items"] if i["tag"]=="text"]
            deck = Presentation(output/"diagrams.pptx")
            slide = deck.slides[0]
            self.assertEqual([s.text for s in slide.shapes if s.has_text_frame and s.text],expected)
            rounded = next(s for s in slide.shapes if s.name.startswith("rect-") and s.left>0)
            self.assertAlmostEqual(rounded.adjustments[0],16/180,places=4)
            centred = next(s for s in slide.shapes if s.has_text_frame and s.text=="输入数据")
            scale = deck.slide_width/1200
            self.assertAlmostEqual(centred.left+centred.width/2,185*scale,delta=2)
            curve = next(s for s in slide.shapes if s.name.startswith("path-"))
            self.assertEqual(len(curve._element.findall(".//"+qn("a:cubicBezTo"))),1)
            with zipfile.ZipFile(output/"diagrams.pptx") as package:
                self.assertFalse(any(n.startswith("ppt/media/") for n in package.namelist()))
            before = hashlib.sha256((output/"diagrams.pptx").read_bytes()).hexdigest()
            again = subprocess.run(args,capture_output=True)
            self.assertEqual(again.returncode,2)
            self.assertEqual(before,hashlib.sha256((output/"diagrams.pptx").read_bytes()).hexdigest())


if __name__ == "__main__":
    unittest.main()
