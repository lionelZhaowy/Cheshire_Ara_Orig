"""Documentation/graphics checks only; no RTL simulation or target claims."""
from pathlib import Path
from collections import Counter
from urllib.parse import unquote
import json,re,hashlib,xml.etree.ElementTree as ET,subprocess,zipfile
ROOT=Path(__file__).resolve().parent;REPO=ROOT.parents[2];NS={'s':'http://www.w3.org/2000/svg'}
errors=[];geometry=[];links=0;sources={};line_refs=0
manifest=json.loads((ROOT/'connections.json').read_text())
for g in manifest['figures']:
 ns={n['id'] for n in g['nodes']}
 for e in g['edges']:
  if e['source'] not in ns or e['target'] not in ns:errors.append(('endpoint',g['name'],e['id']))
  if not e.get('detail',e.get('label','')):errors.append(('missing signal label',g['name'],e['id']))
 for v in ['A','B']:
  p=ROOT/v/(g['name']+'.svg');root=ET.parse(p).getroot();w,h=map(float,root.attrib['viewBox'].split()[2:]);layout=json.loads(p.with_suffix('.layout.json').read_text())
  ids={el.get('id') for el in root}
  for e in g['edges']:
   expected=[e['id']] if v=='A' else [e['id']+'-source',e['id']+'-target']
   if not all(id in ids for id in expected):errors.append(('missing graphical edge',v,g['name'],e['id']))
  for t in layout['labels']:
   if t['x']<-1 or t['y']<0 or t['x']+t['w']>w+1 or t['y']+t['h']>h+1:geometry.append(('text outside canvas',v,g['name'],t['text']))
  if v=='A':
   for el in root.findall('s:polyline',NS):
    if not re.fullmatch(r'E\d+',el.get('id','')):continue
    pts=[tuple(map(float,a.split(','))) for a in el.attrib['points'].split()]
    for t in layout['labels']:
     if re.fullmatch(r'E\d+',t['text']):continue
     # Approximate ink box, not entire line-height; reviewed alongside PNGs.
     x1,y1,x2,y2=t['x']+1,t['y']+3,t['x']+t['w']-1,t['y']+t['h']-4
     for a,b in zip(pts,pts[1:]):
      hit=(a[0]==b[0] and x1<a[0]<x2 and max(min(a[1],b[1]),y1)<min(max(a[1],b[1]),y2)) or (a[1]==b[1] and y1<a[1]<y2 and max(min(a[0],b[0]),x1)<min(max(a[0],b[0]),x2))
      if hit:geometry.append(('wire crosses text',v,g['name'],el.get('id'),t['text']));break
  for ext in ['png','pdf']:
   if not p.with_suffix('.'+ext).exists():errors.append(('missing rendered artifact',str(p),ext))
for p in sorted((REPO/'docs_codex/rtl_modules').glob('*.md'))+[ROOT/'README.md']:
 s=p.read_text()
 if 'FIGURE_PLACEHOLDER' in s:errors.append(('placeholder',p.name))
 for target in re.findall(r'\]\(([^)]+)\)',s):
  if target.startswith(('http:','https:','#','mailto:')):continue
  path,_,frag=unquote(target).partition('#');q=(p.parent/path).resolve();links+=1
  if not q.exists():errors.append(('broken doc link',p.name,target));continue
  if frag.startswith('L') and frag[1:].isdigit():
   line_refs+=1
   if int(frag[1:])>len(q.read_text().splitlines()):errors.append(('line beyond file',p.name,target))
  if not q.is_relative_to(REPO):continue
  if any(part in ['hw','.bender','sw','target'] for part in q.relative_to(REPO).parts[:1]) or q.relative_to(REPO).parts[0]=='docs':
   sources[str(q.relative_to(REPO))]=hashlib.sha256(q.read_bytes()).hexdigest()
# Check the standalone comparison page links (including local anchors).
root_html=(ROOT/'index.html').read_text();anchors=set(re.findall(r'\bid="([^"]+)"',root_html))
for href in re.findall(r'(?:href|src)="([^"]+)"',root_html):
 path,_,frag=href.partition('#');q=(ROOT/path).resolve() if path else ROOT/'index.html';links+=1
 if not q.exists():errors.append(('broken gallery link',href))
 if not path and frag and frag not in anchors:errors.append(('missing gallery anchor',href))
# Preserve the original source evidence; do not refresh old hashes.
old=json.loads((ROOT.parent/'2026-10-12_rtl_modules/source_manifest.json').read_text())
for path,digest in old['source_sha256'].items():
 q=REPO/path
 if q.exists() and hashlib.sha256(q.read_bytes()).hexdigest()!=digest:errors.append(('source drift since prior batch',path))
exports=[]
try:
 from pptx import Presentation
 for v in ['A','B']:
  p=ROOT/v/'editable/diagrams.pptx'
  if not p.exists():errors.append(('PPTX not exported yet',v));continue
  deck=Presentation(p);svgs=sorted((ROOT/v).glob('*.svg'))
  if len(deck.slides)!=len(svgs):errors.append(('slide count',v))
  for slide,src in zip(deck.slides,svgs):
   texts=Counter(el.text or '' for el in ET.parse(src).getroot().findall('s:text',NS));ppt=Counter(sh.text for sh in slide.shapes if sh.has_text_frame and sh.text)
   if texts!=ppt:errors.append(('PPT text mismatch',v,src.name))
  with zipfile.ZipFile(p) as z:
   if any(name.startswith('ppt/media/') for name in z.namelist()):errors.append(('PPT raster image',v))
  exports.append(dict(variant=v,slides=len(deck.slides),objects=sum(len(s.shapes) for s in deck.slides)))
except ImportError:errors.append(('pptx module unavailable; run in skill venv',))
coverage=json.loads((ROOT/'review_coverage.json').read_text())['chapters']
if len(coverage)!=21 or not all(c['reviewed'] for c in coverage):errors.append(('chapter review coverage',len(coverage)))
for c in coverage:
 if hashlib.sha256((REPO/c['document']).read_bytes()).hexdigest()!=c['sha256']:errors.append(('chapter changed since review',c['document']))
visual_images={}
for name in ['visual_review_00_07.json','visual_review_08_16.json','visual_review_17_20.json']:
 d=json.loads((ROOT/name).read_text())
 for item in d['images']:
  p=REPO/item['file'] if item['file'].startswith('docs_codex/') else ROOT/item['file']
  if not item.get('checked') or item.get('issues'):errors.append(('visual review open',item['file']))
  if hashlib.sha256(p.read_bytes()).hexdigest()!=item['sha256']:errors.append(('visual review stale',item['file']))
  visual_images[str(p.relative_to(ROOT))]=item['sha256']
if len(visual_images)!=48:errors.append(('visual coverage',len(visual_images)))
pdf_checks=json.loads((ROOT/'pdf_checks.json').read_text())
if pdf_checks['errors'] or len(pdf_checks['files'])!=48:errors.append(('PDF checks',pdf_checks['errors']))
for item in json.loads((ROOT/'render_manifest.json').read_text()):
 if hashlib.sha256((ROOT/item['file']).read_bytes()).hexdigest()!=item['sha256']:errors.append(('render stale',item['file']))
implementation=subprocess.check_output(['git','diff','--name-only','--','hw','sw','.bender','target'],cwd=REPO,text=True).splitlines()
if implementation:errors.append(('implementation modified',implementation))
report=dict(figures=len(manifest['figures']),variants=2,connections=sum(len(g['edges']) for g in manifest['figures']),chapters=len(list((REPO/'docs_codex/rtl_modules').glob('[0-9]*.md'))),links=links,line_refs=line_refs,source_files=len(sources),errors=errors,geometry_findings=geometry,exports=exports,visually_checked_images=len(visual_images),vector_pdfs=len(pdf_checks['files']),implementation_diff=implementation,rtl_simulated=False,board_tested=False,office_tested=False)
(ROOT/'source_manifest.json').write_text(json.dumps(dict(head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=REPO,text=True).strip(),source_sha256=sources),ensure_ascii=False,indent=2))
(ROOT/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
raise SystemExit(bool(errors or geometry))
