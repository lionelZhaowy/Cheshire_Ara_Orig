"""Read-only current A validation; never refresh historical hashes or reports."""
from pathlib import Path
from collections import Counter
from urllib.parse import unquote
import json,re,hashlib,xml.etree.ElementTree as ET,subprocess,zipfile,argparse,unicodedata
import pymupdf as fitz
from pptx import Presentation
ROOT=Path(__file__).resolve().parent;REPO=ROOT.parents[2];NS={'s':'http://www.w3.org/2000/svg'}
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--output',type=Path,help='Optional new JSON report; existing files are never overwritten')
args=parser.parse_args()
if args.output and args.output.exists():parser.error('Report already exists; choose a fresh output path')
errors=[];geometry=[];links=0;sources={};line_refs=0
manifest=json.loads((ROOT/'connections.json').read_text())
for g in manifest['figures']:
 ns={n['id'] for n in g['nodes']}
 for e in g['edges']:
  if e['source'] not in ns or e['target'] not in ns:errors.append(('endpoint',g['name'],e['id']))
  if not e.get('detail',e.get('label','')):errors.append(('missing signal label',g['name'],e['id']))
 for v in ['A']:
  p=ROOT/v/(g['name']+'.svg');root=ET.parse(p).getroot();w,h=map(float,root.attrib['viewBox'].split()[2:]);layout=json.loads(p.with_suffix('.layout.json').read_text())
  ids={el.get('id') for el in root}
  for e in g['edges']:
   expected=[e['id']] if v=='A' else [e['id']+'-source',e['id']+'-target']
   if not all(id in ids for id in expected):errors.append(('missing graphical edge',v,g['name'],e['id']))
  for t in layout['labels']:
   if t['x']<-1 or t['y']<0 or t['x']+t['w']>w+1 or t['y']+t['h']>h+1:geometry.append(('text outside canvas',v,g['name'],t['text']))
  if v=='A':
   for el in root.findall('s:polyline',NS):
    if not re.fullmatch(r'E\d+',el.get('id','')):continue
    pts=[tuple(map(float,a.split(','))) for a in el.attrib['points'].split()]
    for t in layout['labels']:
     if re.fullmatch(r'E\d+',t['text']):continue
     # Approximate ink box, not entire line-height; reviewed alongside PNGs.
     x1,y1,x2,y2=t['x']+1,t['y']+3,t['x']+t['w']-1,t['y']+t['h']-4
     for a,b in zip(pts,pts[1:]):
      hit=(a[0]==b[0] and x1<a[0]<x2 and max(min(a[1],b[1]),y1)<min(max(a[1],b[1]),y2)) or (a[1]==b[1] and y1<a[1]<y2 and max(min(a[0],b[0]),x1)<min(max(a[0],b[0]),x2))
      if hit:geometry.append(('wire crosses text',v,g['name'],el.get('id'),t['text']));break
  for ext in ['png','pdf']:
   if not p.with_suffix('.'+ext).exists():errors.append(('missing rendered artifact',str(p),ext))

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def check(ok,*why):
 if not ok:errors.append(why)
def md_ids(text):
 text=re.sub(r"```.*?```", "", text, flags=re.S)
 ids=set(re.findall(r'<(?:a|span)\s+id="([^"]+)"',text));seen={}
 for title in re.findall(r'^#{1,6} (.+)$',text,re.M):
  title=re.sub(r'<[^>]+>','',title).lower().replace('`','')
  slug=''.join(c for c in title if c in ' _-' or unicodedata.category(c)[0] in 'LNM').replace(' ','-')
  n=seen.get(slug,0);seen[slug]=n+1;ids.add(slug+(f'-{n}' if n else ''))
 return ids

def link(p,target):
 global links,line_refs
 if target.startswith(('http:','https:','mailto:','//')):return
 path,_,frag=unquote(target).partition('#');q=(p.parent/path).resolve() if path else p
 links+=1
 check(q.exists(),'missing link',str(p),target)
 if not q.exists():return
 if frag.startswith('L') and frag[1:].isdigit():
  line_refs+=1;check(0<int(frag[1:])<=len(q.read_text().splitlines()),'source line',target)
 elif frag and q.suffix=='.html':check(frag in re.findall(r'\bid="([^"]+)"',q.read_text()),'HTML anchor',target)
 elif frag and q.suffix=='.md':check(frag in md_ids(q.read_text()),'Markdown anchor',target)

chapters=sorted((REPO/'docs_codex/rtl_modules').glob('[0-9]*.md'))
check(len(chapters)==21,'chapter count',len(chapters))
check(len(manifest['figures'])==24,'figure count')
check(sum(len(g['edges']) for g in manifest['figures'])==213,'connection count')
for p in chapters+[REPO/'docs_codex/rtl_modules/README.md',ROOT/'README.md']:
 text=p.read_text();check('FIGURE_PLACEHOLDER' not in text,'placeholder',p.name)
 for target in re.findall(r'\]\(([^)]+)\)',text):link(p,target)
 check('rtl_modules_native_c/' not in text and 'rtl_module_comparison/B/' not in text,'retired variant in active navigation',p.name)
for p in [ROOT/'index.html']:
 for target in re.findall(r'(?:href|src)="([^"]+)"',p.read_text()):link(p,target)

source_manifest=json.loads((ROOT/'source_manifest.json').read_text())
for path,digest in source_manifest['source_sha256'].items():
 q=REPO/path;check(q.exists() and sha(q)==digest,'source hash mismatch',path)
for g in manifest['figures']:
 doc=(REPO/'docs_codex/rtl_modules'/g['doc']).read_text()
 check(f'/A/{g["name"]}.png' in doc,'figure not referenced by chapter',g['name'])
 match=re.search(r'<!-- CONNECTIONS:'+re.escape(g['name'])+r' -->(.*?)<!-- /CONNECTIONS -->',doc,re.S)
 check(bool(match),'connection table missing',g['name'])
 if match:check(re.findall(r'^\| (E\d+) \|',match[1],re.M)==[e['id'] for e in g['edges']],'connection table coverage',g['name'])

# Original records are immutable. Inspect only their A subset and label reuse explicitly.
visual_images={}
for name in ['visual_review_00_07.json','visual_review_08_16.json','visual_review_17_20.json']:
 for item in json.loads((ROOT/name).read_text())['images']:
  p=REPO/item['file'] if item['file'].startswith('docs_codex/') else ROOT/item['file']
  if not p.is_relative_to(ROOT/'A'):continue
  check(item.get('checked') and not item.get('issues'),'historical visual issue',item['file'])
  check(sha(p)==item['sha256'],'A PNG changed since visual review',item['file'])
  visual_images[str(p.relative_to(ROOT))]=item['sha256']
check(len(visual_images)==24,'historical A visual coverage')
rendered=0
for item in json.loads((ROOT/'render_manifest.json').read_text()):
 if not item['file'].startswith('A/'):continue
 check(sha(ROOT/item['file'])==item['sha256'],'A SVG changed since rendering',item['file']);rendered+=1
check(rendered==24,'rendered A coverage')
pdfs=[]
for src in sorted((ROOT/'A').glob('*.pdf')):
 with fitz.open(src) as doc:
  fonts=[f for page in doc for f in page.get_fonts()]
  check(all(not page.get_images() for page in doc),'PDF raster image',src.name)
  check(bool(fonts) and all(doc.extract_font(f[0])[3] for f in fonts),'PDF font not embedded',src.name)
  pdfs.append(dict(file=src.name,sha256=sha(src)))
check(len(pdfs)==24,'PDF count')
deckfile=ROOT/'A/editable/diagrams.pptx';deck=Presentation(deckfile)
check(len(deck.slides)==24,'PPT slide count')
for slide,src in zip(deck.slides,sorted((ROOT/'A').glob('*.svg'))):
 svg=Counter(el.text or '' for el in ET.parse(src).getroot().findall('s:text',NS))
 ppt=Counter(sh.text for sh in slide.shapes if sh.has_text_frame and sh.text)
 check(svg==ppt,'PPT text mismatch',src.name)
with zipfile.ZipFile(deckfile) as z:check(not any(n.startswith('ppt/media/') for n in z.namelist()),'PPT media')
report=dict(check='current_A',head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=REPO,text=True).strip(),chapters=len(chapters),figures=len(manifest['figures']),connections=sum(len(g['edges']) for g in manifest['figures']),source_hashes=len(source_manifest['source_sha256']),links=links,source_line_refs=line_refs,unchanged_historical_visual_hashes=len(visual_images),new_visual_review=False,vector_pdfs=len(pdfs),ppt_slides=len(deck.slides),ppt_native_objects=sum(len(s.shapes) for s in deck.slides),errors=errors,geometry_findings=geometry,rtl_simulated=False,board_tested=False,office_tested=False,browser_interaction_tested=False)
serialized=json.dumps(report,ensure_ascii=False,indent=2)+'\n'
if args.output:
 with args.output.open('x') as f:f.write(serialized)
print(serialized)
raise SystemExit(bool(errors or geometry))
