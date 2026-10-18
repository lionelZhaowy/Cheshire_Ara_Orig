#!/usr/bin/env python3
"""Check the delivered flat collection, native object properties and edit/save.

This is an OOXML/python-pptx check, not an Office display or UI editing test.
"""
import sys
import argparse
import hashlib
from html.parser import HTMLParser
import json
from pathlib import Path
import tempfile
import zipfile
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn


sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from diagram_paths import resolve_asset

def counts(prs):
    result = dict(slides=len(prs.slides), objects=0, arrow_objects=0, arrowheads=0, dashed_objects=0)
    for slide in prs.slides:
        result['objects'] += len(slide.shapes)
        for shape in slide.shapes:
            heads = [e for e in shape._element.iter() if e.tag in (qn('a:headEnd'), qn('a:tailEnd')) and e.get('type') == 'triangle']
            result['arrow_objects'] += bool(heads)
            result['arrowheads'] += len(heads)
            result['dashed_objects'] += shape._element.find('.//'+qn('a:custDash')) is not None
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--repo', type=Path, default=Path.cwd())
    parser.add_argument('--directory', type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    directory = args.directory
    assert all(p.is_file() for p in directory.iterdir()), 'Output must be flat'
    report = json.loads((directory/'export_report.json').read_text())
    plan = json.loads((directory/'reconstruction_plan.json').read_text())
    assert len(report['figures']) == len(plan) == 45
    assert len(list(directory.glob('*.svg'))) == 45
    assert len(list(directory.glob('*.pptx'))) == 46
    assert len(list(directory.glob('*.png'))) == 50
    for entry, row in zip(plan, report['figures']):
        source = resolve_asset(args.repo, entry['source'])
        assert hashlib.sha256(source.read_bytes()).hexdigest() == entry['sha256']
        stem = row['name']
        with Image.open(directory/(stem+'.png')) as im:
            assert im.size == tuple(map(round, entry['view'][2:]))
            im.verify()
        deck = directory/(stem+'.pptx')
        actual = counts(Presentation(deck))
        omitted = row['omitted_canvas_background_ids']
        assert omitted == ['rect-001'], stem
        slide = Presentation(deck).slides[0]
        assert not any(s.name in omitted for s in slide.shapes), stem
        background_color = row['slide_background_color']
        if background_color:
            assert str(slide.background.fill.fore_color.rgb).lower() == background_color[1:]
        else:
            assert slide._element.find('.//'+qn('p:bg')) is None
        expected = dict(slides=1, objects=row['objects'], arrow_objects=row['native_arrow_objects'],
                        arrowheads=len(entry['arrows']), dashed_objects=len(entry['dashes']))
        assert actual == expected, (stem, actual, expected)
        with zipfile.ZipFile(deck) as package:
            assert package.testzip() is None
            assert not any(n.startswith('ppt/media/') for n in package.namelist())
    combined = counts(Presentation(directory/'all_diagrams.pptx'))
    assert combined == dict(slides=45, objects=3109, arrow_objects=306, arrowheads=324, dashed_objects=4)
    edits = []
    with tempfile.TemporaryDirectory(prefix='native-reconstruction-edit-') as work:
        for stem in ('example', 'ddr_interconnect', 'cheshire_software_hardware', 'learning_mechanism-xbar'):
            prs = Presentation(directory/(stem+'.pptx'))
            before = counts(prs)
            edited = 0
            for shape in prs.slides[0].shapes:
                if shape._element.find('.//'+qn('a:tailEnd')) is not None:
                    shape.line.width += 1270
                    edited += 1
                if shape._element.find('.//'+qn('a:custDash')) is not None:
                    shape.width += 12700
                    edited += 1
            assert edited > 0
            output = Path(work)/(stem+'.pptx')
            prs.save(output)
            assert counts(Presentation(output)) == before
            edits.append(dict(figure=stem, modified_objects=edited, semantic_counts_preserved=True))
    class Links(HTMLParser):
        def handle_starttag(self, tag, attrs):
            for key, value in attrs:
                if key in ('href', 'src'):
                    assert (directory/value).is_file(), value
    Links().feed((directory/'index.html').read_text())
    prior = resolve_asset(args.repo, 'docs_codex/diagrams/2026-10-11_native_pptx_validation/old_skill_hashes.json')
    skill_hashes = json.loads(prior.read_text())
    for filename, expected in skill_hashes.items():
        assert hashlib.sha256(Path(filename).read_bytes()).hexdigest() == expected
    result = dict(status='PASS', combined=combined, original_sources_unchanged=45,
                  omitted_canvas_backgrounds=sum(len(r['omitted_canvas_background_ids']) for r in report['figures']),
                  old_skill_source_files_unchanged=len(skill_hashes),
                  all_outputs_flat=True, previews_and_index_links_valid=True,
                  edit_save_roundtrips=edits, office_display_editing_tested=False)
    if args.report:
        args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
