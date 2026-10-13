"""Check source anchors, doc links, geometry and native exports (not RTL behavior)."""
from pathlib import Path
from collections import Counter
import argparse
import hashlib
import json
import re
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from PIL import ImageFont

OUT=Path(__file__).resolve().parent
REPO=OUT.parents[2]
GUIDE=REPO/'docs_codex/rtl_modules'
FONT='/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'
NS={'s':'http://www.w3.org/2000/svg'}
OLD=OUT.parent/'2026-10-12_regbus_ip_integration'


def masters():
    return [OLD/'01_regbus_whole.svg',OLD/'02_ip_register_contract.svg']+sorted(OUT.glob('*.svg'))


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--source-only',action='store_true')
    args=parser.parse_args()
    errors=[]
    figures=[]
    for svg in sorted(OUT.glob('*.svg')):
        root=ET.parse(svg).getroot()
        width,height=map(float,root.attrib['viewBox'].split()[2:])
        layout=json.loads(svg.with_suffix('.layout.json').read_text())
        textboxes=[]
        for label in layout['labels']:
            font=ImageFont.truetype(FONT,label['size'])
            length=font.getlength(label['text'])
            x=label['x']-(length/2 if label['anchor']=='middle' else 0)
            a,b,c,d=font.getbbox(label['text'],anchor='ls')
            box=(x+a,label['y']+b,x+c,label['y']+d)
            if not (box[0]>=0 and box[1]>=0 and box[2]<=width and box[3]<=height):
                errors.append((svg.name,'text out of canvas',label['text'],box))
            textboxes.append((box,label['text']))
        for i,(a,t) in enumerate(textboxes):
            for b,u in textboxes[i+1:]:
                if min(a[2],b[2])-max(a[0],b[0])>1 and min(a[3],b[3])-max(a[1],b[1])>1:
                    errors.append((svg.name,'text overlap',t,u))
            for edge in layout['edges']:
                for p,q in zip(edge['points'],edge['points'][1:]):
                    x1,y1=p; x2,y2=q
                    if x1==x2:
                        hit=a[0]+1<x1<a[2]-1 and min(max(y1,y2),a[3])-max(min(y1,y2),a[1])>2
                    else:
                        hit=a[1]+1<y1<a[3]-1 and min(max(x1,x2),a[2])-max(min(x1,x2),a[0])>2
                    if hit: errors.append((svg.name,'edge crosses text',t,p,q))
        figures.append(dict(figure=svg.name,text_objects=len(textboxes),bounds_and_intersection_checked=True))
    manifest=json.loads((OUT/'source_manifest.json').read_text())
    for path,digest in manifest['source_sha256'].items():
        if hashlib.sha256((REPO/path).read_bytes()).hexdigest()!=digest:
            errors.append(('source changed since analysis',path))
    for topic,rows in manifest['anchors'].items():
        for row in rows:
            lines=(REPO/row['path']).read_text().splitlines()
            if row['anchor'] not in lines[row['line']-1]: errors.append(('anchor mismatch',topic,row))
    docs=sorted(GUIDE.glob('*.md'))+[OUT/'README.md']
    links=0
    for page in docs:
        if not page.exists():
            if not args.source_only: errors.append(('missing doc',str(page)))
            continue
        for link in re.findall(r'\]\(([^)]+)\)',page.read_text()):
            if link.startswith(('http','#','mailto:')): continue
            path=link.split('#')[0]
            if (page.parent/path).resolve()==(OUT/'validation.json').resolve(): continue
            if args.source_only and ('/export/' in path or path.startswith('export/')): continue
            if not (page.parent/path).exists(): errors.append(('broken link',str(page.relative_to(REPO)),link))
            links+=1
    exports=[]
    if not args.source_only:
        from pptx import Presentation
        deck=Presentation(OUT/'export/diagrams.pptx')
        ms=masters()
        if len(deck.slides)!=len(ms): errors.append(('slide count',len(deck.slides),len(ms)))
        for svg,slide in zip(ms,deck.slides):
            root=ET.parse(svg).getroot()
            svg_text=Counter(''.join(t.itertext()) for t in root.findall('s:text',NS))
            ppt_text=Counter(s.text for s in slide.shapes if s.has_text_frame and s.text)
            if svg_text!=ppt_text: errors.append(('PPT text mismatch',svg.name))
            for ext in ['png','pdf','editable.svg','gray.png']:
                if not (OUT/'export'/f'{svg.stem}.{ext}').exists(): errors.append(('missing export',svg.stem,ext))
            exports.append(dict(figure=svg.name,text_objects=sum(svg_text.values()),ppt_objects=len(slide.shapes),text_retained=svg_text==ppt_text))
        with zipfile.ZipFile(OUT/'export/diagrams.pptx') as z:
            if any(p.startswith('ppt/media/') for p in z.namelist()): errors.append(('PPT raster media found',))
    changed=subprocess.check_output(['git','diff','--name-only','--','hw','.bender','sw'],cwd=REPO,text=True).splitlines()
    if changed: errors.append(('tracked implementation paths differ',changed))
    report=dict(source_only=args.source_only,errors=errors,figures=figures,exports=exports,
                source_files=len(manifest['source_sha256']),anchors=sum(map(len,manifest['anchors'].values())),
                module_chapters=len(manifest['anchors']),checked_links=links,
                implementation_diff=changed,rtl_simulated=False,board_tested=False,office_tested=False)
    dest=OUT/('source_checks.json' if args.source_only else 'validation.json')
    dest.write_text(json.dumps(report,ensure_ascii=False,indent=2))
    print(json.dumps(dict(errors=errors,figures=len(figures),exports=len(exports),anchors=report['anchors'],links=links),ensure_ascii=False,indent=2))
    raise SystemExit(bool(errors))


if __name__=='__main__': main()
