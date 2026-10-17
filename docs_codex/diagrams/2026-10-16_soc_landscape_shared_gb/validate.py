#!/usr/bin/env python3
"""Validate the localized native PPTX edit; does not validate hardware or Office."""
import hashlib
import json
from pathlib import Path
import posixpath
import subprocess
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
from edit_deck import NS, REMOVED, CHANGED, SCALE

D=Path(__file__).resolve().parent
ROOT=D.parents[2]
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()


def objects(slide):
    return {int(e.find('.//p:cNvPr',NS).get('id')):e for e in slide.find('p:cSld/p:spTree',NS)
            if E.QName(e).localname in ['sp','cxnSp']}


def summary(prs):
    result=[]
    for slide in prs.slides:
        texts=[s.text for s in slide.shapes if s.has_text_frame and s.text]
        arrows=sum(bool(s._element.xpath('.//a:headEnd[@type="triangle"] | .//a:tailEnd[@type="triangle"]')) for s in slide.shapes)
        dashes=sum(bool(s._element.xpath('.//a:custDash')) for s in slide.shapes)
        result.append({'shapes':len(slide.shapes),'text_objects':len(texts),'texts':texts,'native_arrows':arrows,'native_dashed_shapes':dashes})
    return result


def main():
    report_path=D/'evidence/validation.json';assert not report_path.exists()
    inputs=json.loads((D/'evidence/input.json').read_text())
    protected={p:sha(ROOT/p)==v for p,v in inputs['input_sha256'].items() if '/handoffs/' not in p}
    assert all(protected.values()),protected
    handoff='docs_codex/handoffs/L01_global_buffer_cpu_access.md'
    assert sha(D/'evidence/before/L01_global_buffer_cpu_access.md')==inputs['input_sha256'][handoff]
    input_path=ROOT/'docs_codex/diagrams/soc_landscape.pptx'
    output_path=D/'soc_landscape_shared_gb.pptx'
    edit=json.loads((D/'evidence/build-02/edit_report.json').read_text())
    assert sha(output_path)==edit['output_sha256']
    with ZipFile(input_path) as src,ZipFile(output_path) as out:
        assert out.testzip() is None
        unchanged=[n for n in src.namelist() if n not in edit['changed_existing_package_parts']]
        assert all(src.read(n)==out.read(n) for n in unchanged)
        assert not any(n.startswith('ppt/media/') for n in out.namelist())
        assert src.read('ppt/slides/slide1.xml')==out.read('ppt/slides/slide1.xml')
        before=objects(E.fromstring(src.read('ppt/slides/slide1.xml')))
        after=objects(E.fromstring(out.read('ppt/slides/slide2.xml')))
        assert set(before)-set(after)==REMOVED
        new=set(after)-set(before);assert new==set(range(1200,1206))
        altered={id for id in before.keys()&after.keys() if E.tostring(before[id])!=E.tostring(after[id])}
        assert altered==CHANGED,(altered,CHANGED)
        for n in [n for n in out.namelist() if n.endswith('.rels')]:
            r=E.fromstring(out.read(n));ids=[e.get('Id') for e in r];assert len(ids)==len(set(ids))
            base=posixpath.dirname(posixpath.dirname(n)) if n!='_rels/.rels' else ''
            for rel in r:
                if rel.get('TargetMode')=='External':continue
                target=posixpath.normpath(posixpath.join(base,rel.get('Target'))).lstrip('/')
                assert target in out.namelist(),(n,target)
        for index in [1,2]:
            sl=E.fromstring(out.read(f'ppt/slides/slide{index}.xml'))
            ids=sl.xpath('//p:cNvPr/@id',namespaces=NS);assert len(ids)==len(set(ids))
        # Unchanged, distinct CPU/Ara inputs S0/S2; the new port is an output.
        for id in [880,881,889,890,894,895]:assert E.tostring(before[id])==E.tostring(after[id])
        assert after[1201].find('p:spPr/a:ln/a:tailEnd',NS).get('type')=='triangle'
        assert after[1201].find('p:spPr/a:ln/a:headEnd',NS) is None
        p=after[1201].find('p:spPr/a:custGeom/a:pathLst/a:path',NS)
        xf=after[1201].find('p:spPr/a:xfrm',NS);off=xf.find('a:off',NS);ext=xf.find('a:ext',NS)
        points=[(int(off.get('x'))/SCALE+int(c[0].get('x'))/int(p.get('w'))*int(ext.get('cx'))/SCALE,
                 int(off.get('y'))/SCALE+int(c[0].get('y'))/int(p.get('h'))*int(ext.get('cy'))/SCALE) for c in p]
        assert all(abs(a-b)<.002 for a,b in zip(points[0],(1040,878.13996)))
        assert all(abs(a-b)<.002 for a,b in zip(points[-1],(1391.25197,615)))
    prs=Presentation(output_path);assert len(prs.slides)==2
    stats=summary(prs)
    # Re-save only to a separate evidence file; the delivered deck retains all
    # original slide-1 package bytes, which a general writer might normalize.
    roundtrip=D/'evidence/roundtrip-02.pptx'
    if not roundtrip.exists():prs.save(roundtrip)
    assert summary(Presentation(roundtrip))==stats
    layout=json.loads((D/'evidence/preview-03/preview_layout.json').read_text())
    second=layout['slides'][1];modified_texts={778,787,797,798,816,883,886,947,1203,1204,1205}
    overlaps=[]
    for t in second['texts']:
        if int(t['shape_id']) not in modified_texts:continue
        x,y,w,h=t['box'];assert 0<=x and 0<=y and x+w<=3600 and y+h<=2025
        for u in second['texts']:
            if u is t:continue
            xx,yy,ww,hh=u['box']
            if min(x+w,xx+ww)-max(x,xx)>1 and min(y+h,yy+hh)-max(y,yy)>1:
                overlaps.append((t['text'],u['text']))
    assert not overlaps,overlaps
    # New continuous GB and filter paths must avoid all text boxes.
    segments=[((1040,878.13996),(1040,615)),((1040,615),(1391.25197,615)),((606.25197,541.13996),(826.25197,541.13996))]
    for (x1,y1),(x2,y2) in segments:
        for t in second['texts']:
            x,y,w,h=t['box'];pad=3.84
            # Flat line caps do not extend beyond the endpoint along the path.
            cross=(x1==x2 and x-pad<x1<x+w+pad and max(min(y1,y2),y)<min(max(y1,y2),y+h)) or (y1==y2 and y-pad<y1<y+h+pad and max(min(x1,x2),x)<min(max(x1,x2),x+w))
            assert not cross,t
    assert (D/'slide-01.svg').read_bytes()==(D/'evidence/input-preview/slide-01.svg').read_bytes()
    assert (D/'slide-01.png').read_bytes()==(D/'evidence/input-preview/slide-01.png').read_bytes()
    diff=subprocess.run(['git','diff','--check'],cwd=ROOT,capture_output=True,text=True)
    assert diff.returncode==0,diff.stdout+diff.stderr
    report={'status':'pass','date':'2026-10-16','original_file_unchanged':True,
            'original_slide_xml_and_dependencies_unchanged':True,'unchanged_existing_package_parts':len(unchanged),
            'page2_unchanged_shapes':len(before)-len(REMOVED)-len(CHANGED),
            'page2_deleted_shape_ids':sorted(REMOVED),'page2_changed_shape_ids':sorted(CHANGED),
            'page2_new_native_shapes':6,'slides':[{k:v for k,v in s.items() if k!='texts'} for s in stats],
            'save_reopen':'pass','package_relationships':'pass','new_path_endpoints':'pass',
            'new_paths_avoid_text':'pass','modified_text_overlap':'none in local font approximation',
            'preview_type':'DrawingML subset to SVG; not an Office screenshot',
            'office_application_tested':False,'hardware_tested':False,'protected':protected,
            'artifact_sha256':{p.name:sha(p) for p in [output_path,*D.glob('slide-*.*')]}}
    report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
