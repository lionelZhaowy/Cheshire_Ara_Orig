#!/usr/bin/env python3
"""Check delivered SVG/native PPTX structure, text geometry, and source identity."""
from pathlib import Path
import hashlib
import json
import xml.etree.ElementTree as ET
from zipfile import ZipFile
from PIL import ImageFont
from pptx import Presentation

root=Path(__file__).resolve().parent
svg=root/'regbus_structure.svg'
ppt=root/'regbus_structure.pptx'
tree=ET.parse(svg).getroot()
ns={'s':'http://www.w3.org/2000/svg'}
texts=tree.findall('s:text',ns)
expected=[''.join(t.itertext()) for t in texts]
prs=Presentation(ppt)
assert len(prs.slides)==1
actual=[s.text for s in prs.slides[0].shapes if s.has_text_frame and s.text]
assert actual==expected, 'PPT text differs from SVG'
assert len(prs.slides[0].shapes)==len(list(tree))-2, 'Object count mismatch'
boxes=[]
font_file='/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'
for t in texts:
    font=ImageFont.truetype(font_file,round(float(t.attrib['font-size'])))
    val=''.join(t.itertext())
    x,y=float(t.attrib['x']),float(t.attrib['y'])
    width=font.getlength(val)
    anchor=t.attrib.get('text-anchor','start')
    if anchor=='middle': x-=width/2
    elif anchor=='end': x-=width
    l,top,r,bottom=font.getbbox(val,anchor='ls')
    box=(x+l,y+top,x+r,y+bottom)
    assert box[0]>=0 and box[1]>=0 and box[2]<=2400 and box[3]<=1350, t.attrib['id']
    boxes.append((t.attrib['id'],box))
overlaps=[]
for i,(a,A) in enumerate(boxes):
    for b,B in boxes[i+1:]:
        if min(A[2],B[2])-max(A[0],B[0])>1 and min(A[3],B[3])-max(A[1],B[1])>1:
            overlaps.append([a,b])
assert not overlaps, overlaps
with ZipFile(ppt) as z:
    assert z.testzip() is None
    assert not any(n.startswith('ppt/media/') for n in z.namelist()), 'Raster media present'
    doc=ET.fromstring(z.read('ppt/slides/slide1.xml'))
    a={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    lines=doc.findall('.//a:ln',a)
    arrows=sum(any(n.tag.endswith(('headEnd','tailEnd')) and n.attrib.get('type')=='triangle' for n in l) for l in lines)
    dashes=len(doc.findall('.//a:custDash',a))
    assert arrows==28 and dashes==15,(arrows,dashes)
repo=root.parents[2]
sources=json.loads((root/'source_hashes.json').read_text())
assert all(hashlib.sha256((repo/p).read_bytes()).hexdigest()==h for p,h in sources.items()), 'Source snapshot changed'
result={'status':'PASS','canvas':[2400,1350],'slide_count':1,
        'ppt_objects':len(prs.slides[0].shapes),'text_objects':len(texts),
        'native_arrow_objects':arrows,'native_dashed_objects':dashes,
        'raster_media':0,'text_overlap_count':0,'text_canvas_overflow_count':0,
        'font_measurement':font_file,'source_hashes_match':True,
        'office_display_edit_tested':False,
        'svg_sha256':hashlib.sha256(svg.read_bytes()).hexdigest(),
        'pptx_sha256':hashlib.sha256(ppt.read_bytes()).hexdigest()}
(root/'validation.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps(result,indent=2))
