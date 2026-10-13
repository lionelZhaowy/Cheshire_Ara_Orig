"""Check deliverable text retention, bounds and source fingerprints (not RTL behavior)."""
from pathlib import Path
from collections import Counter
import hashlib
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from PIL import ImageFont
from pptx import Presentation

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[2]
FONT = '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'
NS = {'s': 'http://www.w3.org/2000/svg'}
masters = sorted(ROOT.glob('0*.svg'))
deck = Presentation(ROOT/'export/diagrams.pptx')
assert len(deck.slides) == len(masters) == 4
checks = []
for source, slide in zip(masters, deck.slides):
    xml = ET.parse(source).getroot()
    width, height = map(float, xml.attrib['viewBox'].split()[2:])
    labels = json.loads(source.with_suffix('.labels.json').read_text())
    boxes = []
    for label in labels:
        font = ImageFont.truetype(FONT, label['size'])
        length = font.getlength(label['text'])
        x = label['x'] - (length/2 if label['anchor']=='middle' else 0)
        a,b,c,d = font.getbbox(label['text'],anchor='ls')
        box = (x+a,label['y']+b,x+c,label['y']+d)
        assert box[0] >= 0 and box[1] >= 0 and box[2] <= width and box[3] <= height, (source.name,label,box)
        boxes.append((box,label['text']))
    overlaps=[]
    for i,(a,t) in enumerate(boxes):
        for b,u in boxes[i+1:]:
            if min(a[2],b[2])-max(a[0],b[0]) > 1 and min(a[3],b[3])-max(a[1],b[1]) > 1:
                overlaps.append([t,u])
    assert not overlaps,(source.name,overlaps)
    svg_text=Counter(''.join(t.itertext()) for t in xml.findall('s:text',NS))
    ppt_text=Counter(shape.text for shape in slide.shapes if shape.has_text_frame and shape.text)
    assert svg_text==ppt_text,(source.name,svg_text-ppt_text,ppt_text-svg_text)
    checks.append(dict(figure=source.name,text_objects=sum(svg_text.values()),objects=len(slide.shapes),
                       text_bounds_ok=True,text_overlaps=overlaps,ppt_text_retained=True))
with zipfile.ZipFile(ROOT/'export/diagrams.pptx') as archive:
    assert not any(n.startswith('ppt/media/') for n in archive.namelist())
for link in re.findall(r'\]\(([^)]+)\)',(ROOT/'README.md').read_text()):
    if not link.startswith('http') and link != 'validation.json':
        assert (ROOT/link).exists(),link
paths=[
    'hw/cheshire_soc.sv', 'hw/cheshire_pkg.sv',
    '.bender/git/checkouts/register_interface-902ad5bfde7bb98c/src/axi_to_reg_v2.sv',
    '.bender/git/checkouts/axi-ecdc900686449c15/src/axi_to_detailed_mem.sv',
]
paths += ['.bender/git/checkouts/common_cells-7f7ae0f5e6bf7fb5/src/'+name+'.sv' for name in
          ['stream_fork','stream_fork_dynamic','stream_join','stream_to_mem','mem_to_banks_detailed']]
report=dict(figures=checks,source_sha256={p:hashlib.sha256((REPO/p).read_bytes()).hexdigest() for p in paths},
            ppt_has_raster_media=False,office_tested=False,rtl_simulated=False,
            note='Geometric/text checks do not establish electrical correctness; see README source analysis.')
(ROOT/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(checks,ensure_ascii=False))
