#!/usr/bin/env python3
"""Verify SVG/PPTX text, native geometry, source hashes and layout metrics."""
from pathlib import Path
from collections import Counter
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import hashlib,json
from PIL import ImageFont
from pptx import Presentation

root=Path(__file__).resolve().parent
fontfile='/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'
ns={'s':'http://www.w3.org/2000/svg','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
results=[]
for stem in ['regbus_bridge_rtl','regbus_decode_rtl']:
    src=root/(stem+'.svg'); deck=root/(stem+'.pptx')
    tree=ET.parse(src).getroot()
    width,height=map(float,tree.attrib['viewBox'].split()[2:])
    texts=tree.findall('s:text',ns)
    expected=[''.join(t.itertext()) for t in texts]
    prs=Presentation(deck)
    assert len(prs.slides)==1
    actual=[s.text for s in prs.slides[0].shapes if s.has_text_frame and s.text]
    assert actual==expected,stem+' text'
    shapes=len(prs.slides[0].shapes)
    assert shapes==len(list(tree))-2,(stem,shapes,len(list(tree)))
    boxes=[]
    for t,s in zip(texts,expected):
        font=ImageFont.truetype(fontfile,round(float(t.attrib['font-size'])))
        x,y=float(t.attrib['x']),float(t.attrib['y'])
        if t.attrib.get('text-anchor')=='middle':x-=font.getlength(s)/2
        l,top,r,b=font.getbbox(s,anchor='ls')
        box=x+l,y+top,x+r,y+b
        assert box[0]>=0 and box[1]>=0 and box[2]<=width and box[3]<=height,(stem,s,box)
        boxes.append((s,box))
    overlaps=[]
    for i,(a,A) in enumerate(boxes):
        for b,B in boxes[i+1:]:
            if min(A[2],B[2])-max(A[0],B[0])>1 and min(A[3],B[3])-max(A[1],B[1])>1:overlaps.append([a,b])
    assert not overlaps,(stem,overlaps)
    with ZipFile(deck) as z:
        assert z.testzip() is None
        assert not any(n.startswith('ppt/media/') for n in z.namelist())
        xml=ET.fromstring(z.read('ppt/slides/slide1.xml'))
        arrows=sum(any(e.tag.endswith(('headEnd','tailEnd')) and e.attrib.get('type')=='triangle' for e in l) for l in xml.findall('.//a:ln',ns))
        expected_arrows=sum('marker-end' in e.attrib or 'marker-start' in e.attrib for e in tree)
        assert arrows==expected_arrows
    results.append({'figure':stem,'canvas':[width,height],'native_shapes':shapes,'texts':len(texts),'native_arrows':arrows,'text_overlaps':0,'text_canvas_overflow':0,'raster_media':0,'svg_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'pptx_sha256':hashlib.sha256(deck.read_bytes()).hexdigest()})
combined=Presentation(root/'regbus_rtl_analysis.pptx')
assert len(combined.slides)==2
for i,stem in enumerate(['regbus_bridge_rtl','regbus_decode_rtl']):
    solo=Presentation(root/(stem+'.pptx'))
    assert [s.text for s in combined.slides[i].shapes if s.has_text_frame and s.text]==[s.text for s in solo.slides[0].shapes if s.has_text_frame and s.text]
with ZipFile(root/'regbus_rtl_analysis.pptx') as z:assert not any(n.startswith('ppt/media/') for n in z.namelist())
repo=root.parents[2]
hashes=json.loads((root/'source_hashes.json').read_text())
assert all(hashlib.sha256((repo/p).read_bytes()).hexdigest()==h for p,h in hashes.items())
report={'status':'PASS','figures':results,'combined_slides':2,'source_hashes_match':True,'font_measurement':fontfile,'office_display_edit_tested':False,'validation_scope':'native objects, text and source snapshots; no RTL simulation'}
(root/'validation.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps(report,indent=2))
