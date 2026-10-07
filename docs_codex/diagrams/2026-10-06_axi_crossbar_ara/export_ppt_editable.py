#!/usr/bin/env python3
"""Export this diagram as simple SVG and native, independently editable PPT shapes.

Run with python-pptx and Pillow available. The SVG export uses only stdlib.
Does not outline text, rasterize objects, or modify the original diagram.
"""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET

SVG = "http://www.w3.org/2000/svg"
ET.register_namespace("", SVG)
STYLES = {
    "": (21, "#183044", "normal"),
    "title": (34, "#183044", "bold"),
    "h": (25, "#183044", "bold"),
    "sub": (19, "#50667a", "normal"),
    "small": (17, "#50667a", "normal"),
}


def color(value):
    if re.fullmatch(r"#[0-9a-fA-F]{3}", value):
        return "#" + "".join(c * 2 for c in value[1:])
    return value


def export_svg(source):
    old = ET.parse(source).getroot()
    new = ET.Element(f"{{{SVG}}}svg", {
        "version": "1.1", "width": old.get("width"),
        "height": old.get("height"), "viewBox": old.get("viewBox"),
    })
    ids = Counter()

    def add(tag, attrs, text=None):
        ids[tag] += 1
        item = ET.SubElement(new, f"{{{SVG}}}{tag}", {
            "id": f"{tag}-{ids[tag]:03d}", **attrs,
        })
        item.text = text

    for node in old:
        tag = node.tag.split("}")[-1]
        a = dict(node.attrib)
        if tag in {"title", "desc"}:
            add(tag, {}, node.text)
        elif tag in {"style", "defs"}:
            continue
        elif tag == "text":
            size, fill, weight = STYLES[a.get("class", "")]
            add("text", {
                "x": "305" if node.text == "入端局部适配" else a["x"],
                "y": a["y"], "font-family": "Microsoft YaHei",
                "font-size": str(size), "font-weight": weight,
                "font-style": "normal", "text-anchor": "start",
                "fill": fill, "stroke": "none",
                "{http://www.w3.org/XML/1998/namespace}space": "preserve",
            }, node.text)
        elif tag == "rect":
            # Plain rectangles avoid the rounded-rectangle conversion seen in PPT.
            add("rect", {
                "x": a.get("x", "0"), "y": a.get("y", "0"),
                "width": a["width"], "height": a["height"],
                "fill": color(a["fill"]), "stroke": color(a.get("stroke", "none")),
                "stroke-width": a.get("stroke-width", "0"),
            })
        elif tag == "path":
            stroke = color(a["stroke"])
            if a.get("opacity"):
                # Composite the grey mesh onto its known background; no SVG alpha.
                alpha = float(a["opacity"])
                foreground = bytes.fromhex(stroke[1:])
                background = bytes.fromhex("eef4f9")
                stroke = "#" + "".join(
                    f"{round(f * alpha + b * (1 - alpha)):02x}"
                    for f, b in zip(foreground, background)
                )
            straight = re.fullmatch(r"M([\d.]+) ([\d.]+) H([\d.]+)", a["d"])
            attrs = {
                "fill": "none", "stroke": stroke,
                "stroke-width": a["stroke-width"],
                "stroke-linecap": "butt", "stroke-linejoin": "miter",
            }
            if straight:
                x1, y, x2 = straight.groups()
                add("line", {"x1": x1, "y1": y, "x2": x2, "y2": y, **attrs})
                end_x, end_y = float(x2), float(y)
            else:
                add("path", {"d": a["d"], **attrs})
                values = list(map(float, re.findall(r"[\d.]+", a["d"])))
                end_x, end_y = values[-2:]
            if a.get("marker-end"):
                # All marked paths in this diagram end horizontally to the right.
                name = re.search(r"#(\w+)", a["marker-end"]).group(1)
                length, half, fill = {
                    "arrow": (8, 4, "#526b80"),
                    "blue": (10, 5, "#2774c7"),
                    "orange": (10, 5, "#c66a20"),
                }[name]
                add("polygon", {
                    "points": f"{end_x-length:g},{end_y-half:g} {end_x:g},{end_y:g} {end_x-length:g},{end_y+half:g}",
                    "fill": fill, "stroke": "none",
                })
        else:
            raise ValueError(f"Unexpected SVG element: {tag}")
    old_text = [n.text for n in old if n.tag == f"{{{SVG}}}text"]
    new_text = [n.text for n in new if n.tag == f"{{{SVG}}}text"]
    assert old_text == new_text, "Text changed during SVG export"
    assert len({n.get("id") for n in new}) == len(new)
    ET.indent(new, space="  ")
    return new


def export_pptx_one(svg, destination):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
    from PIL import ImageFont

    prs = Presentation()
    width, height = float(svg.get("width")), float(svg.get("height"))
    prs.slide_width = Inches(14)
    scale = prs.slide_width / width
    prs.slide_height = round(height * scale)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    prs.core_properties.title = "Cheshire + Ara AXI Crossbar — editable shapes"
    prs.core_properties.subject = "Native shapes and text; no SVG conversion required"

    def emu(value):
        return round(float(value) * scale)

    def paint(shape, a):
        if hasattr(shape, "fill"):
            if a.get("fill", "none") == "none":
                shape.fill.background()
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(a["fill"][1:])
        if a.get("stroke", "none") == "none":
            shape.line.fill.background()
        else:
            shape.line.color.rgb = RGBColor.from_string(a["stroke"][1:])
            shape.line.width = emu(a["stroke-width"])

    def custom_path(points, cubic=False):
        min_x = min(x for x, _ in points)
        min_y = min(y for _, y in points)
        w = max(1, emu(max(x for x, _ in points) - min_x))
        h = max(1, emu(max(y for _, y in points) - min_y))
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(min_x), emu(min_y), w, h)
        sppr = shape._element.spPr
        sppr.remove(sppr.prstGeom)
        geom = OxmlElement("a:custGeom")
        for name in ["avLst", "gdLst", "ahLst", "cxnLst"]:
            geom.append(OxmlElement(f"a:{name}"))
        bounds = OxmlElement("a:rect")
        for key, value in {"l": "0", "t": "0", "r": "r", "b": "b"}.items():
            bounds.set(key, value)
        geom.append(bounds)
        paths = OxmlElement("a:pathLst")
        curve = OxmlElement("a:path")
        for key, value in {"w": str(w), "h": str(h), "fill": "none" if cubic else "norm"}.items():
            curve.set(key, value)

        def point(parent, xy):
            pt = OxmlElement("a:pt")
            pt.set("x", str(emu(xy[0] - min_x)))
            pt.set("y", str(emu(xy[1] - min_y)))
            parent.append(pt)

        move = OxmlElement("a:moveTo")
        point(move, points[0])
        curve.append(move)
        if cubic:
            command = OxmlElement("a:cubicBezTo")
            for xy in points[1:]:
                point(command, xy)
            curve.append(command)
        else:
            for xy in points[1:]:
                command = OxmlElement("a:lnTo")
                point(command, xy)
                curve.append(command)
            curve.append(OxmlElement("a:close"))
        paths.append(curve)
        geom.append(paths)
        sppr.insert(1, geom)
        return shape

    for node in svg:
        tag = node.tag.split("}")[-1]
        a = node.attrib
        if tag in {"title", "desc"}:
            continue
        if tag == "rect":
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *[emu(a[k]) for k in ["x", "y", "width", "height"]])
        elif tag == "line":
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *[emu(a[k]) for k in ["x1", "y1", "x2", "y2"]])
        elif tag in {"path", "polygon"}:
            values = list(map(float, re.findall(r"[\d.]+", a.get("d", a.get("points")))))
            points = list(zip(values[::2], values[1::2]))
            shape = custom_path(points, cubic=(tag == "path"))
        elif tag == "text":
            size = float(a["font-size"])
            bold = a["font-weight"] == "bold"
            fontfile = Path("/mnt/c/Windows/Fonts/msyhbd.ttc" if bold else "/mnt/c/Windows/Fonts/msyh.ttc")
            if fontfile.exists():
                font = ImageFont.truetype(str(fontfile), round(size * 10))
                advance = font.getlength(node.text) / 10
            else:
                advance = sum(size if ord(c) > 255 else size * 0.65 for c in node.text)
            shape = slide.shapes.add_textbox(
                emu(a["x"]), emu(float(a["y"]) - size * 1.08),
                emu(advance + size), emu(size * 1.6),
            )
            tf = shape.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.word_wrap = False
            tf.vertical_anchor = MSO_ANCHOR.TOP
            paragraph = tf.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.line_spacing = 1.0
            run = paragraph.add_run()
            run.text = node.text
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(size * scale / 12700)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(a["fill"][1:])
            for script in ["ea", "cs"]:
                typeface = OxmlElement(f"a:{script}")
                typeface.set("typeface", "Microsoft YaHei")
                run._r.get_or_add_rPr().append(typeface)
        else:
            raise ValueError(tag)
        shape.name = a["id"] + (": " + node.text[:40] if tag == "text" else "")
        if tag != "text":
            paint(shape, a)
        # add_shape carries a theme style whose effect reference adds a shadow.
        # Explicit effects and colors are required for a faithful flat diagram.
        from pptx.oxml.ns import qn
        inherited_style = shape._element.find(qn("p:style"))
        if inherited_style is not None:
            shape._element.remove(inherited_style)
        shape._element.spPr.append(OxmlElement("a:effectLst"))

    expected = [n.text for n in svg if n.tag == f"{{{SVG}}}text"]
    # Exclude empty text bodies carried by PowerPoint autoshapes.
    assert [s.text for s in slide.shapes if s.has_text_frame and s.text] == expected
    prs.save(destination)
    reopened = Presentation(destination)
    assert [s.text for s in reopened.slides[0].shapes if s.has_text_frame and s.text] == expected
    return len(slide.shapes)


def export_pptx(svgs, destination, output_dir):
    """Combine native slides without copying any media or raster fallback."""
    from copy import deepcopy
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    combined = Presentation()
    combined.slide_width = Inches(14)
    tallest = max(float(root.get("height")) for root in svgs)
    width = float(svgs[0].get("width"))
    combined.slide_height = round(tallest * combined.slide_width / width)
    counts = []
    for index, root in enumerate(svgs, 1):
        single = output_dir / f"native_slide_{index}.pptx"
        counts.append(export_pptx_one(root, single))
        source = Presentation(single)
        slide = combined.slides.add_slide(combined.slide_layouts[6])
        slide.background.fill.solid()
        from pptx.dml.color import RGBColor
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F7FAFC")
        y_offset = round((combined.slide_height - source.slide_height) / 2)
        for shape in source.slides[0].shapes:
            clone = deepcopy(shape._element)
            transform = clone.find(".//" + qn("a:xfrm"))
            off = transform.find(qn("a:off"))
            off.set("y", str(int(off.get("y")) + y_offset))
            slide.shapes._spTree.insert_element_before(clone, "p:extLst")
    combined.core_properties.title = "Cheshire Ara / DDR multiport — editable diagrams"
    combined.save(destination)
    reopened = Presentation(destination)
    for root, slide in zip(svgs, reopened.slides):
        expected = [n.text for n in root if n.tag == f"{{{SVG}}}text"]
        assert [s.text for s in slide.shapes if s.has_text_frame and s.text] == expected
    return counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out-dir", type=Path, required=True)
    parser.add_argument("--svg-only", action="store_true")
    args = parser.parse_args()
    sources = [Path(__file__).with_name(name) for name in ["crossbar_ara.svg", "ddr_qos_proposal.svg"]]
    args.out_dir.mkdir(parents=True, exist_ok=True)
    if any(args.out_dir.iterdir()):
        raise FileExistsError("Output exists; choose a fresh --out-dir")
    svgs = [export_svg(source) for source in sources]
    diagrams = []
    for source, root in zip(sources, svgs):
        target = args.out_dir / (source.stem + "_ppt.svg")
        ET.ElementTree(root).write(target, encoding="utf-8", xml_declaration=True)
        diagrams.append({
            "source": source.name,
            "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "svg": target.name,
            "svg_element_counts": dict(Counter(n.tag.split("}")[-1] for n in root)),
        })
    report = {
        "diagrams": diagrams,
        "text_preserved": True,
        "text_outlined": False,
        "font": "Microsoft YaHei",
        "rounded_rectangles": False,
        "css_or_markers": False,
        "native_pptx_objects": None,
        "powerpoint_application_verification": "See office_verification.json after running verify_powerpoint.ps1",
    }
    if not args.svg_only:
        report["native_pptx_objects"] = export_pptx(svgs, args.out_dir / "cheshire_ara_ddr_editable.pptx", args.out_dir)
    (args.out_dir / "export_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
