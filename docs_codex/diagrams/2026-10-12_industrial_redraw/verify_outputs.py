#!/usr/bin/env python3
"""Coverage, font metrics, native SVG/PPTX object and source preservation checks."""
import argparse,hashlib,json,xml.etree.ElementTree as ET,zipfile
from pathlib import Path
from PIL import Image,ImageFont
from pptx import Presentation
from pptx.oxml.ns import qn
p=argparse.ArgumentParser(description=__doc__)
p.add_argument('--directory',type=Path,required=True);p.add_argument('--repo',type=Path,default=Path(__file__).resolve().parents[3]);p.add_argument('--report',type=Path)
a=p.parse_args();inventory=json.loads((a.directory/'source_inventory.json').read_text());reports=[];issues=[]
assert len(inventory['figures'])==9 and len(inventory['excluded_learning'])==37
combined=Presentation(a.directory/'all_industrial.pptx');assert len(combined.slides)==9
for idx,f in enumerate(inventory['figures']):
 name=f['name'];stem=a.directory/name;source=a.repo/f['source']
 assert hashlib.sha256(source.read_bytes()).hexdigest()==f['source_sha256']
 root=ET.parse(stem.with_suffix('.svg')).getroot();layout=json.loads(stem.with_suffix('.layout.json').read_text());export=json.loads(stem.with_suffix('.export.json').read_text())
 assert hashlib.sha256(stem.with_suffix('.svg').read_bytes()).hexdigest()==export['svg_sha256']
 ids=[e.get('id') for e in root.iter() if e.get('id')];assert len(ids)==len(set(ids)),name
 required={
  'crossbar_ara':[f'{prefix}{i}' for prefix,n in [('source',7),('demux',7),('mux',5),('target',5)] for i in range(n)]+[f'reach-{i}-{j}' for i in range(7) for j in range(5)],
  'ddr_qos_proposal':['cpu-ara','xbar','llc','isp','npu','port0','port1','port2','arbiter','core','phy'],
  'example':['input','process','output','feedback'],
  'initialization_hw_sw':['cpu','bist','rom','llc-regs','devices','spm','rom-asm','rom-c','crt0','poll-config'],
  'program_execution_flow':['reset','early','rom-main','gpt-load','handoff','payload-start','runtime','app','exit','running','rom-return','observer','ret-loop'],
  'cheshire_software_hardware':['applications','printf','gpt','storage-hal','debug-hal','uart-dif','serial-dif','timer-dma','mmio','interfaces','cva6','ara','dma','future-ip','axi','memory','peripherals','dram','media','sensor','build','runtime','os','platform'],
  'ddr_interconnect':['cpu','ara','filter','split','width','npu','router','gb','mover','lvds','cdc','isp','select','frame','raw','isp-gb','xbar','atomics','llc','ddr-port0','ddr-port1','ddr-port2','qos','scheduler','phy-ddr'],
  'llc_cache_spm':[f'cell-{r}-{c}' for r in range(4) for c in range(8)]+[f'partition-{r}-{c}' for r in range(3) for c in range(8)]+['tag','set','block','byte','quiesce','cfg','commit','fsm','resume'],
  'soc_full_architecture':['cva6','ara-core','debug-sba','idma-master','slink-rx','vga-master','usb-master','ext-master','debug-target','reg-bridge','memory-target','idma-config','slink-tx','ext-target','bootrom','clint','soc-regs','plic','llc-regs','uart','i2c','spi','gpio','slink-regs','vga-regs','usb-regs','err-vga','err-dma','err-cpu','irq-router','axi-rt-config','clic','reg-error','reg-ext','ddr-S0','ddr-S1','ddr-S2']+[f'config-C{i}' for i in range(1,10)]
 }[name]
 assert set(required).issubset(ids),(name,set(required)-set(ids))
 boxes=[];local=[]
 for t in layout['texts']:
  ft=ImageFont.truetype('/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc',int(t['size']))
  l,top,r,b=ft.getbbox(t['text'],anchor='ls');length=ft.getlength(t['text']);x=t['x']-{'middle':length/2,'end':length}.get(t['anchor'],0)
  box=(x+l,t['y']+top,x+r,t['y']+b);boxes.append((box,t))
  bx,by,bw,bh=layout['blocks'][t['owner']] if t['owner'] else [0,0,*layout['canvas']]
  if box[0]<bx or box[1]<by or box[2]>bx+bw or box[3]>by+bh:local.append(['text_outside_owner',t['text'],box,[bx,by,bw,bh]])
 for i,(x,t) in enumerate(boxes):
  for y,u in boxes[i+1:]:
   if min(x[2],y[2])-max(x[0],y[0])>1 and min(x[3],y[3])-max(x[1],y[1])>1:local.append(['text_overlap',t['text'],u['text']])
 prs=Presentation(stem.with_suffix('.pptx'));assert len(prs.slides)==1
 expected=[e.text for e in root.iter() if e.tag.endswith('}text')]
 for slide in [prs.slides[0],combined.slides[idx]]:
  shapes=slide.shapes;assert len(shapes)==export['objects'];assert [s.text for s in shapes if s.has_text_frame and s.text]==expected
  arrows=sum(any(e.tag in [qn('a:headEnd'),qn('a:tailEnd')] and e.get('type')=='triangle' for e in s._element.iter()) for s in shapes)
  dashes=sum(s._element.find('.//'+qn('a:custDash')) is not None for s in shapes)
  assert arrows==export['native_arrow_objects'] and dashes==export['native_dashed_objects']
 with zipfile.ZipFile(stem.with_suffix('.pptx')) as z:assert z.testzip() is None and not any(n.startswith('ppt/media/') for n in z.namelist())
 with Image.open(stem.with_suffix('.png')) as im:assert im.size==tuple(layout['canvas']);im.verify()
 if local:issues.append(dict(figure=name,issues=local))
 reports.append(dict(name=name,required_objects_verified=len(required),objects=export['objects'],text_objects=len(expected),native_arrow_objects=arrows,native_dashed_objects=dashes,text_issues=local))
with zipfile.ZipFile(a.directory/'all_industrial.pptx') as z:assert z.testzip() is None and not any(n.startswith('ppt/media/') for n in z.namelist())
result=dict(status='FAIL' if issues else 'PASS',figures=reports,native_only=True,office_display_editing_tested=False,font_metrics='WenQuanYi Zen Hei; Office fonts may differ',issues=issues)
if a.report:a.report.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(dict(status=result['status'],issues=issues,figures=len(reports),objects=sum(r['objects'] for r in reports)),ensure_ascii=False,indent=2))
raise SystemExit(1 if issues else 0)
