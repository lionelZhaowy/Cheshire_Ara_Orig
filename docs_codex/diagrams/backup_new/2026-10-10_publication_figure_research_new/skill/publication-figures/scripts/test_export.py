"""Behavior checks for physical sizing, editable geometry, fonts and failure cases."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from pptx import Presentation
from svg_to_pptx import read_svg


class ExportTests(unittest.TestCase):
    def test_physical_size_and_editable_dash(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); svg=root/'figure.svg'; out=root/'out'
            svg.write_text('''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 100">
<rect id="domain" x="5" y="5" width="190" height="90" fill="none" stroke="#222222" stroke-width="1" stroke-dasharray="5 3"/>
<text id="label" x="20" y="50" font-size="12" font-family="Liberation Sans">Cache / SPM</text>
</svg>''')
            cmd=[sys.executable,str(Path(__file__).with_name('svg_to_pptx.py')),'--input',str(svg),'--out-dir',str(out),'--width-mm','100','--render']
            p=subprocess.run(cmd,capture_output=True,text=True); self.assertEqual(p.returncode,0,p.stderr)
            prs=Presentation(out/'diagrams.pptx')
            self.assertEqual(prs.slide_width,3600000)
            self.assertEqual(prs.slide_height,1800000)
            self.assertEqual(prs.slides[0].shapes[1].text,'Cache / SPM')
            size=prs.slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font.size.pt
            self.assertAlmostEqual(size,12*100*72/25.4/200,delta=.02)
            with zipfile.ZipFile(out/'diagrams.pptx') as z:
                slide=z.read('ppt/slides/slide1.xml').decode()
                self.assertIn('custDash',slide)
                self.assertFalse(any(n.startswith('ppt/media/') for n in z.namelist()))
            report=json.loads((out/'export_report.json').read_text())
            self.assertTrue(all(f['embedded'] for f in report['figures'][0]['render']['fonts']))
            self.assertFalse(report['office_application_tested'])
            self.assertNotEqual(subprocess.run(cmd,capture_output=True).returncode,0)
            small=root/'small'
            cmd[cmd.index(str(out))]=str(small); cmd[cmd.index('100')]='40'
            self.assertNotEqual(subprocess.run(cmd,capture_output=True).returncode,0)
            self.assertFalse(small.exists())

    def test_unsupported_image_is_not_flattened(self):
        with tempfile.TemporaryDirectory() as tmp:
            svg=Path(tmp)/'image.svg'
            svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 100"><image href="private.png" width="200" height="100"/></svg>')
            with self.assertRaisesRegex(ValueError,'Unsupported element image'):
                read_svg(svg)


if __name__=='__main__': unittest.main()
