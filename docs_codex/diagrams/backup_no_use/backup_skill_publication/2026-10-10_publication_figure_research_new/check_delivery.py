#!/usr/bin/env python3
"""Verify the delivered source/export pairs and preserved originals."""
import hashlib
from html.parser import HTMLParser
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import zipfile
import pymupdf
from pptx import Presentation
from PIL import Image

here=Path(__file__).resolve().parent
repo=here.parents[2]
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
old=json.loads((here/'preservation.json').read_text())
changed=[p for p,h in old.items() if not (Path(p) if Path(p).is_absolute() else repo/p).is_file() or sha(Path(p) if Path(p).is_absolute() else repo/p)!=h]
assert not changed, changed
data=json.loads((here/'manifest.json').read_text())
results=[]
for group,figures in data.items():
    folder=here.parent/(group+'_new')
    report=json.loads((folder/'export/export_report.json').read_text())
    prs=Presentation(folder/'export/diagrams.pptx')
    assert len(prs.slides)==len(figures)
    assert abs(prs.slide_width/36000-182)<.001
    with zipfile.ZipFile(folder/'export/diagrams.pptx') as z:
        assert z.testzip() is None and not any(n.startswith('ppt/media/') for n in z.namelist())
    for f,slide,entry in zip(figures,prs.slides,report['figures']):
        name=f['name']; source=folder/(name+'.svg')
        assert sha(source)==entry['sha256']
        root=ET.parse(source).getroot()
        expected=[n.text for n in root if n.tag.endswith('}text')]
        actual=[s.text for s in slide.shapes if s.has_text_frame and s.text]
        assert actual==expected
        assert len(slide.shapes)==len(root)
        sizes=[r.font.size.pt for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs for r in p.runs]
        assert min(sizes)>=9
        with pymupdf.open(folder/f'export/{name}.pdf') as pdf:
            page=pdf[0]
            assert abs(page.rect.width*25.4/72-182)<.05
            assert not page.get_images()
            assert all(pdf.extract_font(row[0])[3] for row in page.get_fonts(full=True))
        with Image.open(folder/f'export/{name}.png') as png:
            assert png.width==2150
        results.append({'group':group,'name':name,'objects':len(slide.shapes),'texts':len(expected),'min_font_pt':min(sizes),'source_sha256':sha(source)})

class Links(HTMLParser):
    def handle_starttag(self,tag,attrs):
        for k,v in attrs:
            if k in ('href','src') and v and not v.startswith(('#','http','data:')):
                assert (here/v.split('#')[0]).exists(),v
Links().feed((here/'index.html').read_text())
for folder in [here,*[here.parent/(name+'_new') for name in data]]:
    for p in folder.glob('*.md'):
        for url in re.findall(r'\]\(([^)]+)\)',p.read_text()):
            if not url.startswith(('http','#')) and (p.parent/url).resolve() != here/'validation.json':
                assert (p.parent/url.split('#')[0]).exists(),(p,url)
skill=here/'skill/publication-figures'; installed=Path.home()/'.codex/skills/publication-figures'
skill_files=[p for p in skill.rglob('*') if p.is_file() and '__pycache__' not in p.parts]
assert all(sha(p)==sha(installed/p.relative_to(skill)) for p in skill_files)
summary={'figure_pages':len(results),'unique_figure_sources':len({r['source_sha256'] for r in results}),
         'pptx_files':len(data),'pdf_fonts_embedded':True,'pptx_texts_and_objects_preserved':True,
         'old_files_unchanged':len(old),'installed_skill_matches_snapshot':True,
         'office_application_tested':False,'figures':results}
(here/'validation.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:v for k,v in summary.items() if k!='figures'},ensure_ascii=False,indent=2))
