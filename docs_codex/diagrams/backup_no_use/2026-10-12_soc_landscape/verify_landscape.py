#!/usr/bin/env python3
"""Check this diagram's native objects, coverage and font-metric text bounds.

This verifies SVG/PPTX structure, not Office rendering or hardware behavior.
"""
import argparse
import hashlib
import json
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.oxml.ns import qn


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--directory', type=Path, default=Path(__file__).resolve().parent)
    p.add_argument('--report', type=Path)
    p.add_argument('--stem', choices=['soc_rearranged','soc_reference_style'], default='soc_rearranged')
    a = p.parse_args()
    stem = a.directory/a.stem
    root = ET.parse(stem.with_suffix('.svg')).getroot()
    layout = json.loads(stem.with_suffix('.layout.json').read_text())
    export = json.loads(stem.with_suffix('.export.json').read_text())
    digest = hashlib.sha256(stem.with_suffix('.svg').read_bytes()).hexdigest()
    assert digest == export['svg_sha256']
    ids = [e.get('id') for e in root.iter() if e.get('id')]
    assert len(ids) == len(set(ids)), 'Duplicate SVG object ID'
    required = {
        'masters': ['cpu','debug-master','ara','idma-master','slink-master','vga-master','usb-master'],
        'slaves': ['debug-target','reg-bridge','memory-target','idma-target','slink-target'],
        'regbus_peripherals': ['bootrom','clint','soc-regs','plic','llc-regs','uart','i2c','spi','gpio',
                              'slink-regs','vga-regs','usb-regs','err-vga','err-dma','err-cpu',
                              'irq-router','axi-rt-regs','clic'],
        'regbus_error_extension': ['reg-error','reg-ext'],
        'axi_extensions': ['axi-ext-master','axi-ext-target'],
        'ddr_ports': ['ddr-S0','ddr-S1','ddr-S2'],
        'new_controls': ['config-C'+str(i) for i in range(1,10)],
        'shared_memory_image': ['npu-compute','gb-adapter','global-buffer','npu-transfer',
                                'lvds','pixel-cdc','isp','frame-select','frame-dma','platform-clock','phy']}
    if a.stem == 'soc_reference_style':
        aliases = {'cpu':'cva6','ara':'ara-core','debug-master':'debug-sba','slink-master':'slink-rx',
                   'idma-target':'idma-config','slink-target':'slink-tx','axi-ext-master':'ext-master',
                   'axi-ext-target':'ext-target','axi-rt-regs':'axi-rt-config','gb-adapter':'gb-router',
                   'platform-clock':'clock-reset','phy':'ddr-phy'}
        required = {k:[aliases.get(n,n) for n in v] for k,v in required.items()}
    for group, names in required.items():
        assert set(names).issubset(ids), (group, set(names)-set(ids))
    for i in range(9):
        assert sum(t['text']=='C'+str(i+1) for t in layout['texts']) == 2
    text_boxes, issues = [], []
    for t in layout['texts']:
        f = ImageFont.truetype('/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc', t['size'])
        l, top, r, b = f.getbbox(t['text'], anchor='ls')
        advance = f.getlength(t['text'])
        x = t['x'] - {'middle': advance/2, 'end': advance}.get(t['anchor'], 0)
        box = (x+l,t['y']+top,x+r,t['y']+b)
        text_boxes.append((box,t))
        bounds = [0,0,*layout['canvas']] if t['owner'] is None else layout['blocks'][t['owner']]
        bx,by,bw,bh = bounds
        if box[0]<bx or box[1]<by or box[2]>bx+bw or box[3]>by+bh:
            issues.append(['text_outside_owner',t['text'],box,bounds])
    for i,(x,t) in enumerate(text_boxes):
        for y,u in text_boxes[i+1:]:
            if min(x[2],y[2])-max(x[0],y[0])>1 and min(x[3],y[3])-max(x[1],y[1])>1:
                issues.append(['text_overlap',t['text'],u['text']])
    assert not issues, issues
    pptx = stem.with_suffix('.pptx')
    prs = Presentation(pptx)
    assert len(prs.slides)==1
    shapes = prs.slides[0].shapes
    assert len(shapes)==export['objects']
    actual_text = [s.text for s in shapes if s.has_text_frame and s.text]
    expected_text = [e.text for e in root.iter() if e.tag.endswith('}text')]
    assert actual_text==expected_text
    arrows = sum(any(e.tag in [qn('a:headEnd'),qn('a:tailEnd')] and e.get('type')=='triangle'
                     for e in s._element.iter()) for s in shapes)
    dashes = sum(s._element.find('.//'+qn('a:custDash')) is not None for s in shapes)
    assert arrows==export['native_arrow_objects'] and dashes==export['native_dashed_objects']
    with zipfile.ZipFile(pptx) as z:
        assert z.testzip() is None
        assert not any(n.startswith('ppt/media/') for n in z.namelist())
    with Image.open(stem.with_suffix('.png')) as im:
        assert im.size==tuple(layout['canvas'])
        im.verify()
    result = dict(status='PASS', canvas=layout['canvas'], coverage={k:len(v) for k,v in required.items()},
                  objects=len(shapes),text_objects=len(actual_text),native_arrows=arrows,native_dashed_objects=dashes,
                  text_bounds_issues=issues,font_metrics='WenQuanYi Zen Hei; Office font may differ',
                  pptx_native_only=True, office_display_editing_tested=False,
                  hashes={ext:hashlib.sha256(stem.with_suffix(ext).read_bytes()).hexdigest()
                          for ext in ['.svg','.png','.pptx']})
    if a.report:
        a.report.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
