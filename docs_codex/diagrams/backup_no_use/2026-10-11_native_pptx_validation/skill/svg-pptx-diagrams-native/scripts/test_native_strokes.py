"""Test one-object arrows/dashes, edit/save retention and safe rejection."""
from pathlib import Path
import tempfile
import unittest
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.oxml.ns import qn
from svg_to_pptx import read_svg, create_deck

ROOT=Path(__file__).resolve().parent.parent
SVG='{http://www.w3.org/2000/svg}'


class NativeStrokes(unittest.TestCase):
    def test_one_object_each_and_edit_roundtrip(self):
        figure=read_svg(ROOT/'assets/native_example.svg')
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'native.pptx'
            report=create_deck([figure],path,'source')
            self.assertEqual(report[0]['objects'],25)
            self.assertEqual(report[0]['native_arrow_objects'],6)
            self.assertEqual(report[0]['native_dashed_objects'],2)
            deck=Presentation(path);shapes={s.name:s for s in deck.slides[0].shapes}
            arrow=shapes['input-to-compute']
            self.assertEqual(arrow._element.tag,qn('p:cxnSp'))
            self.assertIsNone(arrow._element.find('.//'+qn('a:headEnd')))
            self.assertEqual(arrow._element.find('.//'+qn('a:tailEnd')).get('type'),'triangle')
            reverse=shapes['reverse-status']
            self.assertGreater(reverse.begin_x,reverse.end_x)
            both=shapes['control-bidirectional']
            self.assertIsNotNone(both._element.find('.//'+qn('a:headEnd')))
            self.assertIsNotNone(both._element.find('.//'+qn('a:tailEnd')))
            curve=shapes['feedback']
            self.assertEqual(len(curve._element.findall('.//'+qn('a:cubicBezTo'))),1)
            elbow=shapes['configuration-elbow']
            self.assertEqual(len(elbow._element.findall('.//'+qn('a:lnTo'))),4)
            dash=shapes['processing-domain']._element.find('.//'+qn('a:ds'))
            self.assertEqual(dash.attrib,{'d':'500000','sp':'300000'})
            # Editing endpoints and resizing retains decorations, with no separate head to move.
            arrow.end_x+=914400
            shapes['processing-domain'].width+=914400
            curve.left+=91440
            saved_x=arrow.end_x; saved_width=shapes['processing-domain'].width
            deck.save(path)
            reopened=Presentation(path); shapes={s.name:s for s in reopened.slides[0].shapes}
            self.assertEqual(len(shapes),25)
            self.assertEqual(shapes['input-to-compute'].end_x,saved_x)
            self.assertEqual(shapes['processing-domain'].width,saved_width)
            self.assertIsNotNone(shapes['input-to-compute']._element.find('.//'+qn('a:tailEnd')))
            self.assertIsNotNone(shapes['feedback']._element.find('.//'+qn('a:tailEnd')))
            self.assertIsNotNone(shapes['processing-domain']._element.find('.//'+qn('a:custDash')))

    def test_svg_roundtrip_retains_marker_semantics(self):
        before=read_svg(ROOT/'assets/native_example.svg')
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'normalized.svg'
            ET.ElementTree(before['svg']).write(path,encoding='utf-8')
            after=read_svg(path)
            self.assertEqual([i['native_stroke'] for i in before['items']],[i['native_stroke'] for i in after['items']])
            self.assertEqual([i['attrs']['id'] for i in before['items']],[i['attrs']['id'] for i in after['items']])
            # Odd dash arrays repeat as required, instead of creating extra line objects.
            root=ET.parse(path).getroot(); domain=next(n for n in root if n.get('id')=='processing-domain')
            domain.set('stroke-dasharray','10 6 2');ET.ElementTree(root).write(path)
            result=read_svg(path)
            self.assertEqual(next(i['native_stroke']['dash'] for i in result['items'] if i['attrs']['id']=='processing-domain'),[10,6,2,10,6,2])

    def test_invalid_markers_and_dash_fail(self):
        changes=[
            ('markerWidth="3"','markerWidth="6"'),
            ('points="0,0 3,1.5 0,3"','points="0,0 3,3 0,3"'),
            ('url(#blue-arrow)','url(#missing)'),
            ('orient="auto-start-reverse"','orient="auto"'),
            ('stroke-dasharray="10 6"','stroke-dasharray="10 0"'),
            ('marker-end="url(#blue-arrow)"','marker-mid="url(#blue-arrow)"'),
        ]
        original=(ROOT/'assets/native_example.svg').read_text()
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'bad.svg'
            for old,new in changes:
                with self.subTest(replacement=new):
                    path.write_text(original.replace(old,new))
                    with self.assertRaises(ValueError):read_svg(path)


if __name__=='__main__':unittest.main()
