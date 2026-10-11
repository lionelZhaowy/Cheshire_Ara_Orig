"""Background removal must not discard white modules or visible frames."""
import copy
from pathlib import Path
import tempfile
import unittest
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.oxml.ns import qn
from svg_to_pptx import read_svg, create_deck


class CanvasBackground(unittest.TestCase):
    def export(self, first, rest='', view='0 0 100 80'):
        self.work = tempfile.TemporaryDirectory()
        self.addCleanup(self.work.cleanup)
        root = Path(self.work.name)
        svg = root/'test.svg'
        svg.write_text(f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{view}">{first}{rest}</svg>')
        figure = read_svg(svg)
        before = copy.deepcopy(figure)
        ppt = root/'test.pptx'
        report = create_deck([figure], ppt, '16:9')[0]
        self.assertEqual(figure['items'], before['items'], 'PPT filtering must not mutate SVG source items')
        self.assertEqual(ET.tostring(figure['svg']), ET.tostring(before['svg']))
        return Presentation(ppt).slides[0], report

    def test_white_background_and_same_colour_outline_removed_only_from_ppt(self):
        for stroke in ('none', '#fff'):
            with self.subTest(stroke=stroke):
                slide, report = self.export(
                    f'<g fill="#fff" stroke="{stroke}"><rect id="canvas" x="10" y="20" width="100" height="80"/></g>',
                    '<rect id="module" x="20" y="30" width="30" height="20" fill="#fff"/>',
                    '10 20 100 80')
                self.assertEqual([s.name for s in slide.shapes], ['module'])
                self.assertEqual(report['omitted_canvas_background_ids'], ['canvas'])
                self.assertIsNone(slide._element.find('.//'+qn('p:bg')))

    def test_coloured_canvas_is_a_slide_background_not_a_shape(self):
        slide, report = self.export('<rect id="canvas" width="100" height="80" fill="#f7fafc"/>')
        self.assertEqual(len(slide.shapes), 0)
        self.assertEqual(str(slide.background.fill.fore_color.rgb), 'F7FAFC')
        self.assertEqual(report['slide_background_color'], '#f7fafc')

    def test_interior_frames_rounded_and_late_rectangles_are_retained(self):
        variants = [
            '<rect width="90" height="80" fill="#fff"/>',
            '<rect width="100" height="80" fill="#fff" stroke="#123456"/>',
            '<rect width="100" height="80" fill="#fff" stroke="#fff" stroke-dasharray="4 2"/>',
            '<rect width="100" height="80" fill="#fff" rx="5"/>',
            '<rect width="100" height="80" fill="none" stroke="#fff"/>',
        ]
        for first in variants:
            with self.subTest(first=first):
                slide, report = self.export(first)
                self.assertEqual(len(slide.shapes), 1)
                self.assertEqual(report['omitted_canvas_background_ids'], [])
        slide, report = self.export('<text x="10" y="20">label</text>',
                                   '<rect width="100" height="80" fill="#fff"/>')
        self.assertEqual(len(slide.shapes), 2)
        self.assertEqual(report['omitted_canvas_background_ids'], [])


if __name__ == '__main__':
    unittest.main()
