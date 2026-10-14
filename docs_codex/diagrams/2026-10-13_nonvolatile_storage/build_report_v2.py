#!/usr/bin/env python3
"""Dense technical briefing revision. Keeps the v1 generator and all v1 assets."""
import argparse
import hashlib
import json
from pathlib import Path
import sys
import zipfile

from pptx import Presentation
import build as base

ROOT=Path(__file__).resolve().parent
C=base.C
KINGSTON='https://media.kingston.com/pdfs/emmc/itemp-emmc_flyer_en.pdf?_=v1'
PHY='https://www.synopsys.com/designware-ip/interface-ip/mobile-storage/sd-emmc-phy.html'
JEDEC='https://e2e.ti.com/cfs-file/__key/communityserver-discussions-components-files/791/JESD84_2D00_B51-_2800_1_2900_.pdf'
W8='https://www.winbond.com/hq/new-online-purchasing-guide/?__locale=en&pLine=%2Fproduct%2Fcode-storage-flash%2Fqspi-nand%2F&pNo=W25N08MW'

def link(d,pts):
    points=' '.join(f'{x},{y}' for x,y in pts)
    d.add(f'<polyline id="ID" points="{points}" fill="none" stroke="{C["blue"]}" stroke-width="2.5" marker-start="url(#arrow-blue)" marker-end="url(#arrow-blue)"/>')

def make_diagrams(out):
    d=base.Diagram()
    d.box(50,38,430,147,'#E8F1FA')
    d.text(70,66,'eMMC 器件：NAND + 管理控制器',23,bold=True,anchor='start')
    d.box(68,84,130,80,'#D6E7F5');d.text(133,115,'NAND',24,bold=True);d.text(133,146,'参数存储',22)
    d.box(235,84,225,80,'#FFFFFF');d.text(347,115,'介质管理 / ECC',24,bold=True);d.text(347,146,'坏块、磨损、地址映射',21)
    link(d,[(198,125),(235,125)])
    d.box(685,84,220,80,'#FFF1DD');d.text(795,114,'I/O PAD + 收发时序',22,bold=True);d.text(795,145,'电平、驱动、采样',22)
    link(d,[(460,125),(685,125)])
    d.text(580,99,'CLK / CMD / DAT',20)
    d.text(580,154,'单端数字接口',20)
    d.box(1035,84,260,80,'#E0EFE9',dash=True);d.text(1165,113,'SoC eMMC 主机',25,bold=True);d.text(1165,145,'命令 / CRC / FIFO / DMA',21)
    link(d,[(905,125),(1035,125)]);d.text(970,99,'片内接口',20)
    d.box(1085,3,160,44,'#FFFFFF',C['teal']);d.text(1165,32,'CPU 配置',22,bold=True)
    d.arrow([(1165,47),(1165,84)],'teal',True)
    d.box(1425,84,125,80,'#EEE8F6');d.text(1487,116,'DDR',27,bold=True);d.text(1487,146,'运行参数',21)
    d.arrow([(1295,125),(1425,125)]);d.text(1360,99,'AXI 写入',20)
    d.text(800,216,'双箭头合并总线交互；读参数方向为左 → 右。主机控制器与 eMMC 内部管理控制器职责不同。',22,color=C['muted'])
    (out/'emmc_v2.svg').write_text(base.svg_root(''.join(d.items),235))
    d=base.Diagram()
    d.box(50,38,465,147,'#E8F1FA');d.text(70,66,'片外 SPI NAND：按页读取，器件内纠错',23,bold=True,anchor='start')
    for x,w,a,b in [(70,115,'NAND','参数阵列'),(225,110,'ECC','纠错状态'),(375,120,'页缓存','串行输出')]:
        d.box(x,84,w,80,'#FFFFFF');d.text(x+w/2,115,a,24,bold=True);d.text(x+w/2,146,b,21)
    d.arrow([(185,125),(225,125)]);d.arrow([(335,125),(375,125)])
    d.box(680,84,180,80,'#FFF1DD');d.text(770,114,'I/O PAD',24,bold=True);d.text(770,145,'1.8 V / 采样',22)
    link(d,[(495,125),(680,125)]);d.text(590,99,'SCK / CS# / IO',18)
    d.box(975,84,180,80,'#E0EFE9');d.text(1065,114,'SPI Host',26,bold=True);d.text(1065,145,'现有 RTL / FIFO',21)
    link(d,[(860,125),(975,125)])
    d.box(1245,84,150,80,'#EEE8F6');d.text(1320,114,'CPU 驱动',24,bold=True);d.text(1320,145,'读出并搬运',21)
    d.arrow([(1155,125),(1245,125)]);d.text(1200,99,'MMIO',19)
    d.arrow([(1320,84),(1320,17),(1065,17),(1065,84)],'teal',True);d.text(1190,8,'配置 / 命令',19,color=C['teal'])
    d.box(1470,84,80,80,'#FFF1DD');d.text(1510,129,'DDR',24,bold=True);d.arrow([(1395,125),(1470,125)])
    d.text(800,216,'候选读取路径；当前 SPI Host 无独立 DMA 口。片内 ECC 不等于 eMMC 式完整介质管理。',22,color=C['muted'])
    (out/'spi_nand_v2.svg').write_text(base.svg_root(''.join(d.items),235))

def header(p,title,subtitle,band):
    p.rect(50,34,7,50,C['teal'])
    p.text(74,29,1480,title,40,bold=True)
    p.text(74,88,1480,subtitle,23,color=C['muted'])
    p.rect(50,126,1500,44,'#DFEAF2')
    p.text(66,133,1470,band,25,color=C['navy'],bold=True)

def section(p,x,y,w,title):
    p.rect(x,y,w,34,C['navy']);p.text(x+12,y+4,w-24,title,24,bold=True,color='#FFFFFF')

def row(p,x,y,w,label,lines,labelw=112,height=58):
    labelw=max(labelw,140)
    p.rect(x,y,w,height,'#FFFFFF',C['line'])
    p.text(x+10,y+8,labelw-16,label,22,bold=True,color=C['blue'])
    for i,t in enumerate(lines):p.text(x+labelw,y+6+25*i,w-labelw-10,t,22)

def footer(p,items):
    p.rect(50,839,1500,1,C['line'])
    x=50
    for title,url,w in items:
        p.text(x,851,w,title,18,color=C['blue'],url=url);x+=w+20
    p.text(50,878,1450,'公开资料选型示例；频率为建议评估值，非当前工程已达标结果。  |  2026-10-13 · 修订 2',16,color=C['muted'])
    p.text(1490,878,60,f'{p.index}/2',16,color=C['muted'])

def main():
    a=argparse.ArgumentParser();a.add_argument('--out-dir',type=Path,default=ROOT/'revision_02')
    a.add_argument('--skill-dir',type=Path,default=Path.home()/'.codex/skills/svg-pptx-diagrams-native')
    args=a.parse_args();out=args.out_dir
    if out.exists() and any(out.iterdir()):raise SystemExit('Output is not empty; refusing overwrite.')
    out.mkdir(parents=True,exist_ok=True)
    make_diagrams(out)
    sys.path.insert(0,str(args.skill_dir/'scripts'));import svg_to_pptx as native
    nr=native.create_deck([native.read_svg(out/f) for f in ('emmc_v2.svg','spi_nand_v2.svg')],out/'diagrams.pptx','source',base.FONTFILE,16)
    diagram_deck=Presentation(out/'diagrams.pptx')
    prs=Presentation();prs.slide_width=base.emu(1600);prs.slide_height=base.emu(900)
    p=base.Page(prs,1)
    header(p,'eMMC 方案：8-bit SDR / 50 MHz，优先评估','目标：约 500 MiB 校正参数断电保存，开机加载到 DDR；允许 30 秒，NOR 继续承担启动。',
           '建议配置：4 GB 工业级 eMMC 5.1  |  1-bit、≤400 kHz 初始化  →  8-bit High Speed SDR、50 MHz')
    p.embed(diagram_deck.slides[0],out/'emmc_v2.svg',180)
    section(p,50,422,735,'接口与协议：命令、响应、数据分开传输')
    row(p,50,456,735,'信号 / 引脚',['CLK：主机输出；CMD：双向命令与响应。','DAT[7:0]：双向数据；RST_n：按器件配置使用。'])
    row(p,50,514,735,'数据规则',['按逻辑块读写；主机检查响应、CRC 与超时。','DAT0 可指示忙；无 AXI 式逐拍 VALID / READY。'])
    row(p,50,572,735,'初始化',['识别器件 → 读取能力 → CMD6 配置总线与模式。','先低速初始化，再切 8-bit / High Speed 并提高 CLK。'])
    row(p,50,630,735,'连续读取',['CMD18 多块读 → 接收数据 → DMA 写 DDR。','命令响应 ≠ 搬运完成；检查错误后校验参数包。'])
    row(p,50,688,735,'模式边界',['当前建议 SDR：每周期每根数据线传 1 bit。','DS 用于 HS400；本建议模式不使用 DS 采样。'])
    section(p,810,422,740,'器件选型与 ASIC 实现：控制器、I/O、时序分层')
    row(p,810,456,740,'候选器件',['Kingston E04GS14DXI：4 GB、MLC、-40～85°C。','8 GB 备选 EMMC08G-WV28；均为公开选型示例。'],120)
    row(p,810,514,740,'频率 / 带宽',['50 MHz × 8 bit = 50 MB/s ≈ 47.7 MiB/s 线速。','SDR 规格上限 52 MHz；50 MHz 是建议工作点。'],120)
    row(p,810,572,740,'电气 / PHY',['单端数字接口；必须配置匹配 VCCQ 的 I/O PAD。','评估 1.8 V I/O；VCC、VCCQ 与电源时序按料号。'],120)
    row(p,810,630,740,'实现边界',['SDR 可评估数字收发 + I/O；仍须完成时序收敛。','HS200/HS400 常需调谐、DLL/延迟线等 PHY 功能。'],120)
    row(p,810,688,740,'工程改动',['新增主机 / DMA、驱动及 I/O；需确认 IP 交付边界。','确认 PAD、采样、PCB 裕量后，实测完整加载时间。'],120)
    p.rect(50,760,1500,64,'#E0EFE9')
    p.text(66,767,1470,'预算：500 MiB ÷ 30 s = 16.7 MiB/s 下限；建议有效搬运 ≥25 MiB/s，20 s 搬运 + 约 10 s 初始化 / 校验。',23,bold=True,color=C['teal'])
    p.text(66,797,1470,'容量：可建议 A/B 两份参数（约 1000 MiB），保留增长余量；更新须完成落盘、校验与版本切换，并验证掉电恢复。',22)
    footer(p,[('器件：Kingston 工业级 eMMC',KINGSTON,380),('协议：JEDEC 5.1（TI 托管）',JEDEC,410),('PHY：Synopsys 实现示例',PHY,370),('Micron eMMC',base.MICRON,210)])
    p.slide.notes_slide.notes_text_frame.text='\n'.join([
        '用户需求：约500 MiB盲元补偿及非均匀性校正参数；开机加载DDR；30秒可接受。暂将预算定义为存储初始化至加载并校验完成。',
        '本轮建议50MHz而非此前讨论的52MHz：50MHz处在High Speed SDR规定的0–52MHz范围内，作为便于评估的工作点，不是已冻结时钟或IP承诺。初始化选择不超过400kHz；器件保持初始兼容模式，读取能力并完成CMD6模式/总线切换后才升频。初始化命令细节及电气配置应按具体料号实现。',
        'Kingston工业级产品表确认E04GS14DXI为4GB MLC、-40至85摄氏度、eMMC5.1 HS400，EMMC08G-WV28为8GB MLC同温度档。支持高速能力不要求在HS400运行；完整料号数据手册、电压、供货生命周期、保持年限及耐久仍需供应商确认。VCCQ 1.8V是评估方向，不把未取得的具体料号供电细节写为已确认。',
        'eMMC主机控制器为数字RTL；ASIC I/O单元提供电平/驱动/输入接收/ESD。数字总线不等于只需RTL，也不等于所有模式强制购买数模混合PHY。SDR仍需检查输入采样相位、输出建立保持、电气负载、上拉、驱动强度、封装PCB及PVT；不可承诺任意普通GPIO单元可达50MHz。HS200调谐与HS400 strobe常需要延迟/DLL，具体模拟或数字实现由IP决定。',
        'DMA完成与命令响应不同；需检查主机错误并确保DDR可见性，再验参数包。当前Cheshire没有已集成eMMC接口。新增主机IP的控制口、AXI、IO/PHY、仿真模型和固件需分别确认交付范围。',
        'B/s用十进制MB；MiB/s用2^20。50MB/s=47.6837MiB/s，500MiB的纯线速极限约10.49秒；25MiB/s目标要求至少约52.4%的线速有效利用率，非实测吞吐。',
        'A/B是建议非用户冻结需求；采用4/8GB不代替掉电保护、持久化、版本记录冗余与验证。',
        '官方资料，2026-10-13查阅：',KINGSTON,JEDEC,PHY,base.MICRON,base.FAQ,
        '源码依据：hw/cheshire_soc.sv，hw/bootrom/cheshire_bootrom.c，docs/um/sw.md。'])
    p1=p

    p=base.Page(prs,2)
    header(p,'SPI NAND 方案：Quad SDR / 100 MHz，条件性备选','优势：可评估复用现有 SPI Host；代价：新增 NAND 管理软件，并验证采样时序与 CPU 搬运吞吐。',
           '建议配置：W25N08MW 系列，8 Gbit（1 GiB）、1.8 V  |  首次联调 10 MHz → 条件满足后评估 Quad SDR 100 MHz')
    p.embed(diagram_deck.slides[1],out/'spi_nand_v2.svg',180)
    section(p,50,422,735,'接口与协议：SPI 命令 + NAND 页访问')
    row(p,50,456,735,'信号 / 模式',['CS#、SCK：主机输出；IO[3:0]：按事务改变方向。','Quad 为四线数据阶段，不等于所有命令均为四线。'])
    row(p,50,514,735,'一次页读取',['发页地址 / 读页命令 → 等待就绪并检查 ECC。','再发读缓存命令，按列地址 / dummy 周期接收数据。'])
    row(p,50,572,735,'软件职责',['识别器件、配置 ECC / Quad、管理坏块与预留块。','检查不可纠错错误；CPU 排空 FIFO 并写入 DDR。'])
    row(p,50,630,735,'候选容量',['W25N-MW 系列含 2 / 4 / 8 Gbit，最高示例 1 GiB。','500 MiB 单份有空间；双副本仅余 24 MiB，过紧。'])
    section(p,810,422,740,'频率、I/O 与系统吞吐：复用存在约束')
    row(p,810,456,740,'工作频率',['10 MHz 为首次联调建议；100 MHz 为评估目标。','系列标称 STR 最高 166 MHz，不等于系统可达。'],120)
    row(p,810,514,740,'时钟前提',['当前 SCK = clk_i / [2 × (CLKDIV + 1)]。','100 MHz SCK 要求 clk_i ≥200 MHz 且分频匹配。'],120)
    row(p,810,572,740,'物理实现',['单端 1.8 V PAD；校核采样相位、驱动和板级负载。','无高速差分 SerDes；高频仍可能需专用 I/O 时序。'],120)
    row(p,810,630,740,'主要风险',['现有控制器无独立 DMA；Regbus / CPU 可能限速。','不能只换 Flash：需 NAND 驱动与错误 / 掉电验证。'],120)
    xs=[50,370,665,980,1550]
    data=[('方案 / 建议数据模式','纯线速（不含开销）','500 MiB 线速下限','评估结论'),
          ('eMMC · 8-bit SDR 50 MHz','50 MB/s ≈ 47.7 MiB/s','约 10.5 秒','优先评估：新增主机 / DMA 与 I/O'),
          ('SPI NAND · Quad SDR 100 MHz','50 MB/s ≈ 47.7 MiB/s','约 10.5 秒','备选：时钟、页读取、软件搬运须达标')]
    for j,rowdata in enumerate(data):
        y=701+j*37;p.rect(50,y,1500,37,C['navy'] if j==0 else ('#E9F1EC' if j==1 else '#FFFFFF'))
        for k,t in enumerate(rowdata):p.text(xs[k]+9,y+6,xs[k+1]-xs[k]-18,t,21 if j else 22,color='#FFFFFF' if not j else C['ink'],bold=not j)
    p.text(50,816,1500,'验收口径：完整 500 MiB 初始化 + 加载 + 校验 ≤30 秒；SPI Quad 50 MHz 仅 23.8 MiB/s 线速，低于 25 MiB/s 设计目标。',20,color=C['amber'])
    footer(p,[('器件：Winbond W25N-MW',base.WINBOND,350),('8 Gbit：W25N08MW 型号入口',W8,400),('NAND / 管理职责：Micron',base.NAND,400)])
    p.slide.notes_slide.notes_text_frame.text='\n'.join([
        'W25N08MW系列是8Gbit器件候选；最终封装和温度后缀、ECC位数、供货及数据手册尚待确认。W25N-MW公开系列页列1.8V、2/4/8Gb、STR166MHz/DTR83MHz、页组织与BBM LUT；本方案只建议STR/SDR，不默认当前SPI支持DTR。',
        '10MHz仅用于初次电气/驱动联调，不是SPI NAND协议强制初始化频率；提高到100MHz须先确认器件命令时序、控制器移位和采样实现、PAD、封装PCB及系统频率。当前clk_i与公共SoC时钟相连，不能只改一个divider就保证100MHz；独立高速SPI域需要额外CDC和复位适配，未授权实施。',
        'SPI读NAND一般先Array to Cache，等Busy清除并检查ECC，再Read Cache输出。具体opcode、dummy数量和IO方向按料号；不将NOR命令套到NAND。图是功能关系，不规定器件内部ECC流水线位置。连续读可以优化页边界损失，但必须按器件模式与ECC规则验证。',
        '本地SPI Host的FSM clkdiv每半周期计数，SCK=clk_i/(2*(CLKDIV+1))。因此在整比分频下200MHz输入、divider=0才可产生100MHz SCK；至少200MHz只是必要条件，不是时序或吞吐保证。公共Regbus与CPU也有争用，FIFO满导致停顿会降低有效带宽。',
        '100MHz Quad STR=50MB/s≈47.6837MiB/s；500MiB/47.6837≈10.48576秒，只是纯线速极限。页读取等待、软件FIFO访问与DDR写入、初始化和校验都会增加时间。Quad50MHz=25MB/s≈23.8419MiB/s，不能达到25MiB/s建议目标，但并不单独证明一定超过30秒；实际30秒门槛需测量。',
        '复用已有SPI减少RTL范围，但软件坏块、ECC错误处理与更新可靠性工作增加。设备内ECC不能保证SPI链路无误码，参数包仍需校验。8Gbit=1024MiB，两份500MiB后仅24MiB，不计坏块与预留，单颗不作为稳妥双副本推荐。',
        '公开资料，2026-10-13查阅：',base.WINBOND,W8,base.MW_ECC,base.NAND,
        '源码依据：hw/cheshire_soc.sv:1445；.bender/git/checkouts/opentitan_peripherals-7b624fb57f78de9a/src/spi_host/rtl/spi_host_fsm.sv；同依赖的sw/device/lib/dif/dif_spi_host.c；sw/lib/hal/spi_sdcard.c。未更改或动态验收。'])
    metrics=[p1.save(out),p.save(out)]
    final=out/'nonvolatile_storage_options_v2.pptx';prs.save(final)
    rr=Presentation(final);checks=[]
    for n,s in enumerate(rr.slides,1):
        ids=[x.get('id') for x in s._element.xpath('.//p:cNvPr')];assert len(ids)==len(set(ids))
        urls=[r.hyperlink.address for sh in s.shapes if sh.has_text_frame for para in sh.text_frame.paragraphs for r in para.runs if r.hyperlink.address]
        assert len(urls)==(4 if n==1 else 3)
        for sh in s.shapes:
            assert sh.left>=0 and sh.top>=0
            assert sh.left+sh.width<=rr.slide_width+base.emu(2),(n,sh.name)
            assert sh.top+sh.height<=rr.slide_height+base.emu(2),(n,sh.name)
        checks.append(dict(page=n,objects=len(s.shapes),links=urls,arrows=len(s._element.xpath('.//a:tailEnd')),notes_chars=len(s.notes_slide.notes_text_frame.text)))
    with zipfile.ZipFile(final) as z:assert not any(n.startswith('ppt/media/') for n in z.namelist())
    source=json.loads((ROOT/'output/source_manifest.json').read_text())
    f='.bender/git/checkouts/opentitan_peripherals-7b624fb57f78de9a/src/spi_host/rtl/spi_host_fsm.sv'
    source[f]=hashlib.sha256((base.REPO/f).read_bytes()).hexdigest()
    for f,sha in source.items():assert hashlib.sha256((base.REPO/f).read_bytes()).hexdigest()==sha,f
    (out/'source_manifest.json').write_text(json.dumps(source,indent=2)+'\n')
    report={'revision':2,'slides':checks,'text_metrics':metrics,'native_diagram_report':nr,'office_verified':False,
            'preview':'CairoSVG from shared layout; not Office rendering','files':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in out.iterdir() if p.is_file()}}
    (out/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(checks,ensure_ascii=False));print(final)

if __name__=='__main__':main()
