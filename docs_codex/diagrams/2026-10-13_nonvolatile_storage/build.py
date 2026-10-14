#!/usr/bin/env python3
"""Build original native diagrams and a two-slide deck; no production files touched."""
import argparse
import copy
import hashlib
import html
import importlib.util
import json
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
import zipfile

import cairosvg
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[2]
FONT = 'Microsoft YaHei'
FONTFILE = Path('/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc')
C = dict(ink='#193047', muted='#53677B', blue='#246DA0', teal='#218078',
         bg='#F7F9FC', line='#CAD5DF', navy='#163B56', amber='#9B6828')
MICRON = 'https://www.micron.com/products/storage/managed-nand/emmc'
WINBOND = 'https://www.winbond.com/hq/product/code-storage-flash/qspi-nand/w25n-mw/index.html?__locale=en'
FAQ = 'https://www.micron.com/sales-support/sales/faqs'
NAND = 'https://www.micron.com/products/storage/nand-flash/choosing-the-right-nand'
MW_ECC = 'https://www.winbond.com/hq/product/code-storage-flash/qspi-nand/w25n-mw/?__locale=ja&partNo=W25N02MWTBIG'
SCALE = 9144  # 1600 px -> 16 inches

def svg_root(body, height=320):
    return f'<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="{height}" viewBox="0 0 1600 {height}">{body}</svg>'

class Diagram:
    def __init__(self):
        self.items = []
        self.n = 0
        self.items.append('<defs>' + ''.join(
            f'<marker id="arrow-{k}" markerUnits="strokeWidth" markerWidth="3" markerHeight="3" viewBox="0 0 3 3" refX="3" refY="1.5" orient="auto-start-reverse"><polygon points="0,0 3,1.5 0,3" fill="{C[k]}"/></marker>'
            for k in ('blue', 'teal')) + '</defs>')
    def add(self, s):
        self.n += 1
        self.items.append(s.replace('ID', f'd{self.n}'))
    def box(self, x,y,w,h,fill,stroke=C['ink'],dash=False):
        self.add(f'<rect id="ID" x="{x}" y="{y}" width="{w}" height="{h}" rx="6" fill="{fill}" stroke="{stroke}" stroke-width="1.6"'+(' stroke-dasharray="8 5"' if dash else '')+'/>')
    def text(self,x,y,s,size=27,color=C['ink'],bold=False,anchor='middle'):
        self.add(f'<text id="ID" x="{x}" y="{y}" font-family="{FONT}" font-size="{size}" font-weight="{"bold" if bold else "normal"}" text-anchor="{anchor}" fill="{color}">{html.escape(s)}</text>')
    def arrow(self, pts, kind='blue',dash=False):
        p=' '.join(f'{x},{y}' for x,y in pts)
        self.add(f'<polyline id="ID" points="{p}" fill="none" stroke="{C[kind]}" stroke-width="2.8" marker-end="url(#arrow-{kind})"'+(' stroke-dasharray="7 5"' if dash else '')+'/>')
    def save(self,path):
        path.write_text(svg_root(''.join(self.items)),encoding='utf8')

def diagrams(out):
    d=Diagram()
    d.box(60,85,520,205,'#EAF3FA')
    d.text(82,116,'片外 eMMC · 断电保存参数',26,bold=True,anchor='start')
    d.box(85,143,175,118,'#D6E7F5')
    d.text(172,181,'NAND 阵列',27,bold=True)
    d.text(172,224,'校正参数',25)
    d.box(315,143,240,118,'#FFFFFF')
    d.text(435,174,'内部管理控制器',26,bold=True)
    d.text(435,210,'纠错 / 坏块管理',24)
    d.text(435,242,'磨损均衡 / 地址映射',22)
    d.arrow([(260,202),(315,202)])
    d.box(715,143,255,118,'#E4F2ED',dash=True)
    d.text(842,177,'eMMC 主机',28,bold=True)
    d.text(842,213,'控制器 + DMA',27)
    d.text(842,246,'拟新增 IP 与驱动',22,color=C['teal'])
    d.box(750,9,185,53,'#FFFFFF',C['teal'])
    d.text(842,45,'CPU 配置',25,bold=True)
    d.arrow([(842,62),(842,143)],'teal',True)
    d.text(926,101,'控制口',22,color=C['teal'])
    d.arrow([(555,202),(715,202)])
    d.text(643,179,'eMMC 总线',22)
    d.box(1100,143,180,118,'#FFF0D9')
    d.text(1190,183,'DDR',32,bold=True)
    d.text(1190,223,'参数工作副本',23)
    d.arrow([(970,202),(1100,202)])
    d.text(1035,179,'AXI 搬运',22)
    d.box(1390,143,150,118,'#EEE8F6')
    d.text(1465,184,'ISP',32,bold=True)
    d.text(1465,225,'运行时读取',23)
    d.arrow([(1280,202),(1390,202)])
    d.text(1335,179,'参数',22)
    d.text(800,314,'箭头表示参数读取的数据方向；CPU 下发读命令，DMA 完成后软件检查状态与内容。',23,color=C['muted'])
    d.save(out/'emmc.svg')

    d=Diagram()
    d.box(60,58,600,202,'#EAF3FA')
    d.text(82,93,'片外 SPI NAND · 断电保存参数',26,bold=True,anchor='start')
    for x,w,t1,t2 in [(85,155,'NAND 阵列','按页读取'),(290,125,'ECC','器件内纠错'),(465,165,'页缓存','串行输出')]:
        d.box(x,122,w,107,'#FFFFFF')
        d.text(x+w/2,164,t1,26,bold=True)
        d.text(x+w/2,203,t2,23)
    d.arrow([(240,177),(290,177)])
    d.arrow([(415,177),(465,177)])
    d.box(755,122,200,107,'#E4F2ED')
    d.text(855,164,'SPI Host',29,bold=True)
    d.text(855,203,'接收 FIFO',24)
    d.arrow([(630,177),(755,177)])
    d.text(704,152,'Quad',23)
    d.box(1055,122,230,107,'#F0ECF7')
    d.text(1170,164,'CPU / 驱动',28,bold=True)
    d.text(1170,203,'读 FIFO、写 DDR',23)
    d.arrow([(955,177),(1055,177)])
    d.text(1005,152,'MMIO',22)
    d.arrow([(1170,122),(1170,38),(855,38),(855,122)],'teal',True)
    d.text(1015,25,'配置 / 读命令',22,color=C['teal'])
    d.box(1400,122,140,107,'#FFF0D9')
    d.text(1470,165,'DDR',30,bold=True)
    d.text(1470,204,'参数副本',23)
    d.arrow([(1285,177),(1400,177)])
    d.text(1342,152,'存储写入',21)
    d.text(800,293,'候选路径：复用 SPI Host，新增 NAND 驱动；未假设现有 SPI 已具备 DMA 搬运。',24,color=C['muted'])
    d.save(out/'spi_nand.svg')

def rgb(c): return RGBColor.from_string(c.lstrip('#'))
def emu(v): return round(v*SCALE)

class Page:
    def __init__(self, prs, index):
        self.slide=prs.slides.add_slide(prs.slide_layouts[6])
        self.svg=[]
        self.metrics=[]
        self.index=index
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb=rgb(C['bg'])
        self.svg.append('<rect width="1600" height="900" fill="'+C['bg']+'"/>')
    def rect(self,x,y,w,h,fill,line=None):
        s=self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,emu(x),emu(y),emu(w),emu(h))
        s.fill.solid();s.fill.fore_color.rgb=rgb(fill)
        if line: s.line.color.rgb=rgb(line);s.line.width=emu(1)
        else:s.line.fill.background()
        self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{fill}" stroke="{line or "none"}"/>')
    def text(self,x,y,w,s,size=30,color=None,bold=False,url=None):
        color=color or C['ink']
        box=self.slide.shapes.add_textbox(emu(x),emu(y),emu(w),emu(size*1.4))
        tf=box.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.TOP
        p=tf.paragraphs[0];p.space_before=p.space_after=0;p.line_spacing=1.0
        r=p.add_run();r.text=s;r.font.name=FONT;r.font.size=emu(size)
        r.font.bold=bold;r.font.color.rgb=rgb(color)
        for k in ('ea','cs'):
            node=OxmlElement(f'a:{k}');node.set('typeface',FONT);r._r.get_or_add_rPr().append(node)
        if url:r.hyperlink.address=url
        el=f'<text x="{x}" y="{y+size*.94}" font-family="{FONT}" font-size="{size}" font-weight="{"bold" if bold else "normal"}" fill="{color}">{html.escape(s)}</text>'
        if url:el=f'<a href="{html.escape(url,quote=True)}">{el}</a>'
        self.svg.append(el)
        advance=ImageFont.truetype(str(FONTFILE),round(size*10)).getlength(s)/10
        self.metrics.append(dict(text=s,width=round(advance,2),allocated=w,size_pt=size*.72))
        assert advance<=w, (s,advance,w)
    def embed(self,source_slide,svg_file,y):
        for source in source_slide.shapes:
            el=copy.deepcopy(source.element)
            self.slide.shapes._spTree.insert_element_before(el,'p:extLst')
            shape=self.slide.shapes[-1]
            shape.top+=emu(y)
            for n in el.xpath('.//p:cNvPr'):
                n.set('id',str(len(self.slide.shapes)+1))
        root=ET.fromstring(svg_file.read_text())
        self.svg.append(f'<g transform="translate(0 {y})">'+''.join(ET.tostring(c,encoding='unicode') for c in root)+'</g>')
    def common(self,title,sub):
        self.rect(60,43,8,58,C['teal'])
        self.text(88,38,1450,title,46,bold=True)
        self.text(88,108,1450,sub,28,color=C['muted'])
    def footer(self,label,url):
        self.rect(60,849,1480,1,C['line'])
        self.text(60,862,1040,label,19,color=C['blue'],url=url)
        self.text(1200,862,340,f'方案评估 · 2026.10.13   {self.index}/2',18,color=C['muted'])
    def save(self,out):
        path=out/f'slide_{self.index:02}.svg'
        path.write_text(svg_root(''.join(self.svg),900),encoding='utf8')
        cairosvg.svg2png(url=str(path),write_to=str(out/f'slide_{self.index:02}.png'),output_width=1920,output_height=1080)
        return self.metrics

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--out-dir',type=Path,default=ROOT/'output')
    ap.add_argument('--skill-dir',type=Path,default=Path.home()/'.codex/skills/svg-pptx-diagrams-native')
    args=ap.parse_args();out=args.out_dir
    if out.exists() and any(out.iterdir()):raise SystemExit('Refusing to overwrite nonempty output; choose a new --out-dir.')
    out.mkdir(parents=True,exist_ok=True)
    diagrams(out)
    sys.path.insert(0,str(args.skill_dir/'scripts'))
    import svg_to_pptx as native
    figs=[native.read_svg(out/n) for n in ('emmc.svg','spi_nand.svg')]
    native_report=native.create_deck(figs,out/'diagrams.pptx','source',FONTFILE,16)
    native_prs=Presentation(out/'diagrams.pptx')
    prs=Presentation();prs.slide_width=emu(1600);prs.slide_height=emu(900)
    p1=Page(prs,1)
    p1.common('eMMC：内置管理控制器的大容量参数存储','推荐优先评估 4 GB / 8 GB 工业级器件；保留现有 NOR 承担启动。')
    p1.rect(60,165,1480,96,'#E7EFF5')
    for x,a,b in [(86,'500 MiB','盲元补偿 / 非均匀性校正参数'),(615,'30 秒','初始化、加载与校验的预算'),(1130,'DDR 常驻','运行时 ISP 读取参数')]:
        p1.text(x,177,470,a,34,bold=True,color=C['blue'])
        p1.text(x,223,470,b,23,color=C['muted'])
    p1.embed(native_prs.slides[0],out/'emmc.svg',277)
    for x,w,title,lines in [
        (60,460,'器件内部完成管理',['纠错、坏块管理与磨损均衡','主机按逻辑数据块读写','断电后参数仍保留']),
        (565,470,'工程需新增主机通路',['新增控制器 IP、PAD 与驱动','CPU 配置，DMA 批量写 DDR','完成后检查错误并校验参数']),
        (1080,460,'容量与可靠更新建议',['4 GB / 8 GB 为选型档位','可预留 A/B 两份参数','写备用区 → 校验 → 切换版本'])]:
        p1.rect(x,619,w,4,C['teal'])
        p1.text(x,637,w,title,29,bold=True)
        for j,line in enumerate(lines):p1.text(x,685+j*40,w,line,27)
    p1.footer('官方资料：Micron eMMC  ·  点击打开产品与文档',MICRON)
    p1.slide.notes_slide.notes_text_frame.text=(
        '需求：用户确认约500 MiB校正参数在开机完整加载至DDR，允许30秒加载。此页将30秒暂按存储初始化至加载校验完成预算，不代表从上电到完整系统就绪。\n'
        '结构图为建议的参数读取路径，蓝色实线为数据方向、绿色虚线为CPU配置；不是完整双向协议图。管理控制器内部功能为逻辑概念，不规定器件物理微架构。DDR至ISP含互连，未冻结实际端口。\n'
        '当前Cheshire没有已接好的eMMC通路。新增主机控制器建议带DMA，控制寄存器可经适配进入Regbus，数据经AXI；IP、电压、PAD、时钟、ID及缓存维护均待集成规格确定。\n'
        'A/B是本方案建议：两份约1000 MiB，尚非用户冻结需求。更新须确认数据已落入非易失介质，版本记录需冗余/校验，断电注错验证不能由A/B名义替代。4/8 GB不是采购承诺。\n'
        f'官方资料（2026-10-13查阅）：\n{MICRON}\n{FAQ}\n{NAND}\n'
        '当前源码：hw/cheshire_soc.sv:1445 SPI Host；hw/bootrom/cheshire_bootrom.c:76 内置启动介质选择；docs/um/sw.md:68 两阶段启动说明。保留NOR启动，后续加载程序在DDR ready后读取参数。')

    p2=Page(prs,2)
    p2.common('SPI NAND：复用 SPI Host 的扩容候选','W25N-MW 示例：最高 8 Gbit = 1 GiB；NAND 驱动需新增，NOR 驱动不能直接复用。')
    p2.embed(native_prs.slides[1],out/'spi_nand.svg',158)
    cols=[60,325,910,1540]
    rows=[('比较项','eMMC','SPI NAND'),
          ('容量选择','建议 4 GB / 8 GB，方便扩展','示例 1 GiB；双份 500 MiB 余量紧'),
          ('硬件与软件','新增主机 / DMA；器件管理介质','评估复用 SPI；驱动处理页、ECC、坏块'),
          ('性能验收','测量主机 DMA 至 DDR 有效吞吐','核对 Quad、时钟、FIFO 与软件搬运')]
    for j,row in enumerate(rows):
        y=478+j*47
        p2.rect(60,y,1480,47,C['navy'] if j==0 else ('#FFFFFF' if j%2 else '#EDF2F6'))
        for k,cell in enumerate(row):p2.text(cols[k]+15,y+8,cols[k+1]-cols[k]-30,cell,25 if j else 26,color='#FFFFFF' if j==0 else C['ink'],bold=(j==0 or k==0))
    p2.rect(60,686,1480,90,'#E3F0EA')
    p2.text(82,697,1410,'带宽下限：500 ÷ 30 = 16.7 MiB/s     建议有效搬运目标：≥25 MiB/s',29,bold=True,color=C['teal'])
    p2.text(82,740,1410,'25 MiB/s 对应约 20 秒搬运，另留约 10 秒初始化与校验；器件标称速度不等于系统吞吐。',24)
    p2.text(60,795,1480,'选择建议：优先评估 eMMC；SPI NAND 在 Quad 读取与端到端性能达标后，可作为备选。',28,bold=True)
    p2.footer('官方资料：Winbond W25N-MW  ·  点击打开规格与文档',WINBOND)
    p2.slide.notes_slide.notes_text_frame.text=(
        '蓝线为建议的参数读取方向；CPU经控制口下发SPI命令、读取接收FIFO并写DDR。当前SPI Host接公共Regbus，没有将DMA直接接到SPI FIFO的现有保证。图为存储读取逻辑示意，未规定ECC与页缓存的物理实现次序。\n'
        'W25N-MW官网列出2/4/8Gb、1.8V、SPI/Dual/Quad、页组织与BBM LUT。具体料号的ECC能力以数据手册为准；本页不冻结ECC位数。具备器件内ECC不意味着具有eMMC式完整介质管理。\n'
        '容量：500 MiB=524288000 bytes；8 Gbit按该NAND二进制阵列容量为1 GiB=1024 MiB；两份参数1000 MiB只余24 MiB，尚需考虑坏块、备用块、元数据与增长，不建议把它视为稳妥双副本配置。\n'
        '吞吐下限500/30=16.6667 MiB/s（约17.48 MB/s），只计数据传输；25 MiB/s为建议有效搬运目标，不是器件或工程实测。整体验收应从初始化开始计时至加载并校验结束<=30秒。若初始化校验超时须调整预算。\n'
        'Quad STR 50 MHz的纯线速仅25 MB/s≈23.84 MiB/s，未计页读、命令、FIFO与软件开销，不能声称达到25 MiB/s目标。器件标称83 MB/s也不能直接套到当前SPI Host。\n'
        f'官方资料（2026-10-13查阅）：\n{WINBOND}\n{MW_ECC}\n{NAND}\n'
        '当前源码：hw/cheshire_soc.sv:1445，sw/lib/hal/spi_sdcard.c:192及239，sw/lib/hal/spi_s25fs512s.c；本地DIF的SPI时钟分频及读FIFO逻辑。现有SD-SPI驱动25MHz单线是SD软件路径条件，不是SPI NAND的通用最高频率。\n'
        '待适配内容：复位与ID识别、读页和读缓存命令、ECC状态判定、不可纠错错误、坏块映射/预留和更新掉电恢复；本次未实现或测试。')
    preview_metrics=[p1.save(out),p2.save(out)]
    final=out/'nonvolatile_storage_options.pptx';prs.save(final)
    reread=Presentation(final)
    checks=[]
    for i,s in enumerate(reread.slides,1):
        ids=[int(n.get('id')) for n in s._element.xpath('.//p:cNvPr')]
        assert len(ids)==len(set(ids)),ids
        urls=[r.hyperlink.address for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink.address]
        assert urls==[[MICRON],[WINBOND]][i-1],urls
        assert len(s.notes_slide.notes_text_frame.text)>500
        for shape in s.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=reread.slide_width+emu(2),(i,shape.name)
            assert shape.top+shape.height<=reread.slide_height+emu(2),(i,shape.name)
        checks.append(dict(slide=i,shapes=len(s.shapes),hyperlinks=urls,notes_chars=len(s.notes_slide.notes_text_frame.text),
                           native_arrowheads=len(s._element.xpath('.//a:tailEnd'))))
    with zipfile.ZipFile(final) as z:
        assert not any(n.startswith('ppt/media/') for n in z.namelist())
    assert len(reread.slides)==2
    sources=['hw/cheshire_soc.sv','hw/bootrom/cheshire_bootrom.c','sw/lib/hal/spi_sdcard.c','sw/lib/hal/spi_s25fs512s.c',
             '.bender/git/checkouts/opentitan_peripherals-7b624fb57f78de9a/sw/device/lib/dif/dif_spi_host.c','docs/um/sw.md']
    manifest={n:hashlib.sha256((REPO/n).read_bytes()).hexdigest() for n in sources}
    (out/'source_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    report=dict(slides=checks,diagram_export=native_report,text_width_checks=preview_metrics,
                editable_native_objects=True,whole_slide_raster_images=False,
                preview_renderer='CairoSVG from shared SVG/PPT layout; not an Office render',office_verified=False,
                input_head='bf1b4198dfd74cc78291c152b8b682d523e785c5',
                sha256={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in out.iterdir() if p.is_file()})
    (out/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(checks,ensure_ascii=False))
    print(final)

if __name__=='__main__':main()
