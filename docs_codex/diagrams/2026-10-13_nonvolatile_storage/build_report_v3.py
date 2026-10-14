#!/usr/bin/env python3
"""3.3 V PAD revision. Preserves v1/v2 generators and artifacts."""
import argparse
import hashlib
import json
from pathlib import Path
import sys
import zipfile

from pptx import Presentation
import build as base

ROOT=Path(__file__).resolve().parent
C=base.C
ISSI='https://www.mouser.com/datasheet/2/198/ISSI_6_8_22_IS21_22TF08G__Rev_A_0426_2022-3132334.pdf'
SWISSBIT='https://www.swissbit.com/files/public/Documents/Flyer/Product_Profile_Swissbit_eMMC.pdf'
LV='https://www.winbond.com/hq/product/code-storage-flash/qspi-nand/w25n-lv/?__locale=en'
LV4='https://winbond.com/hq/product/code-storage-flash/qspi-nand/w25n-lv/?__locale=en&partNo=W25N04LVCEIG'
PHY='https://www.synopsys.com/designware-ip/interface-ip/mobile-storage/sd-emmc-phy.html'
JEDEC='https://e2e.ti.com/cfs-file/__key/communityserver-discussions-components-files/791/JESD84_2D00_B51-_2800_1_2900_.pdf'


def link(d,pts):
    points=' '.join(f'{x},{y}' for x,y in pts)
    d.add(f'<polyline id="ID" points="{points}" fill="none" stroke="{C["blue"]}" stroke-width="2.5" marker-start="url(#arrow-blue)" marker-end="url(#arrow-blue)"/>')

def make_diagrams(out):
    d=base.Diagram()
    d.box(50,38,430,147,'#E8F1FA')
    d.text(70,66,'eMMC 器件：NAND + 管理控制器',23,bold=True,anchor='start')
    d.box(68,84,130,80,'#D6E7F5');d.text(133,115,'NAND',24,bold=True);d.text(133,146,'参数存储',22)
    d.box(235,84,225,80,'#FFFFFF');d.text(347,115,'介质管理 / ECC',24,bold=True);d.text(347,146,'坏块、磨损、地址映射',21)
    link(d,[(198,125),(235,125)])
    d.box(685,84,220,80,'#FFF1DD');d.text(795,114,'I/O PAD + 收发时序',22,bold=True);d.text(795,145,'3.3 V / 25OD33',22)
    link(d,[(460,125),(685,125)])
    d.text(580,99,'CLK / CMD / DAT',20)
    d.text(580,154,'3.3 V 单端数字',20)
    d.box(1035,84,260,80,'#E0EFE9',dash=True);d.text(1165,113,'SoC eMMC 主机',25,bold=True);d.text(1165,145,'命令 / CRC / FIFO / DMA',21)
    link(d,[(905,125),(1035,125)]);d.text(970,99,'片内接口',20)
    d.box(1085,3,160,44,'#FFFFFF',C['teal']);d.text(1165,32,'CPU 配置',22,bold=True)
    d.arrow([(1165,47),(1165,84)],'teal',True)
    d.box(1425,84,125,80,'#EEE8F6');d.text(1487,116,'DDR',27,bold=True);d.text(1487,146,'供 ISP 使用',20)
    d.arrow([(1295,125),(1425,125)]);d.text(1360,99,'AXI 写入',20)
    d.text(800,216,'双箭头合并总线交互；读参数方向为左 → 右。主机控制器与 eMMC 内部管理控制器职责不同。',22,color=C['muted'])
    (out/'emmc_v3.svg').write_text(base.svg_root(''.join(d.items),235))
    d=base.Diagram()
    d.box(50,38,465,147,'#E8F1FA');d.text(70,66,'片外 SPI NAND：按页读取，器件内纠错',23,bold=True,anchor='start')
    for x,w,a,b in [(70,115,'NAND','参数阵列'),(225,110,'ECC','纠错状态'),(375,120,'页缓存','串行输出')]:
        d.box(x,84,w,80,'#FFFFFF');d.text(x+w/2,115,a,24,bold=True);d.text(x+w/2,146,b,21)
    d.arrow([(185,125),(225,125)]);d.arrow([(335,125),(375,125)])
    d.box(680,84,180,80,'#FFF1DD');d.text(770,114,'I/O PAD',24,bold=True);d.text(770,145,'3.3 V / 采样',22)
    link(d,[(495,125),(680,125)]);d.text(590,99,'SCK / CS# / IO',18)
    d.box(975,84,180,80,'#E0EFE9');d.text(1065,114,'SPI Host',26,bold=True);d.text(1065,145,'现有 RTL / FIFO',21)
    link(d,[(860,125),(975,125)])
    d.box(1245,84,150,80,'#EEE8F6');d.text(1320,114,'CPU 驱动',24,bold=True);d.text(1320,145,'读出并搬运',21)
    d.arrow([(1155,125),(1245,125)]);d.text(1200,99,'MMIO',19)
    d.arrow([(1320,84),(1320,17),(1065,17),(1065,84)],'teal',True);d.text(1190,8,'配置 / 命令',19,color=C['teal'])
    d.box(1470,84,80,80,'#FFF1DD');d.text(1510,129,'DDR',24,bold=True);d.arrow([(1395,125),(1470,125)])
    d.text(800,216,'候选读取路径；当前 SPI Host 无独立 DMA 口。片内 ECC 不等于 eMMC 式完整介质管理。',22,color=C['muted'])
    (out/'spi_nand_v3.svg').write_text(base.svg_root(''.join(d.items),235))

def header(p,title,subtitle,band):
    p.rect(50,34,7,50,C['teal'])
    p.text(74,29,1480,title,40,bold=True)
    p.text(74,88,1480,subtitle,23,color=C['muted'])
    p.rect(50,126,1500,44,'#DFEAF2')
    p.text(66,133,1470,band,25,color=C['navy'],bold=True)

def section(p,x,y,w,title):
    p.rect(x,y,w,34,C['navy']);p.text(x+12,y+4,w-24,title,24,bold=True,color='#FFFFFF')

def row(p,x,y,w,label,lines,labelw=112,height=58):
    labelw=max(labelw,140)
    p.rect(x,y,w,height,'#FFFFFF',C['line'])
    p.text(x+10,y+8,labelw-16,label,22,bold=True,color=C['blue'])
    for i,t in enumerate(lines):p.text(x+labelw,y+6+25*i,w-labelw-10,t,22)

def footer(p,items):
    p.rect(50,839,1500,1,C['line'])
    x=50
    for title,url,w in items:
        p.text(x,851,w,title,18,color=C['blue'],url=url);x+=w+20
    p.text(50,878,1450,'公开资料选型示例；频率为建议评估值，非当前工程已达标结果。  |  2026-10-13 · 修订 3',16,color=C['muted'])
    p.text(1490,878,60,f'{p.index}/2',16,color=C['muted'])

def main():
    a=argparse.ArgumentParser();a.add_argument('--out-dir',type=Path,default=ROOT/'revision_03')
    a.add_argument('--skill-dir',type=Path,default=Path.home()/'.codex/skills/svg-pptx-diagrams-native')
    args=a.parse_args();out=args.out_dir
    if out.exists() and any(out.iterdir()):raise SystemExit('Output is not empty; refusing overwrite.')
    out.mkdir(parents=True,exist_ok=True)
    make_diagrams(out)
    sys.path.insert(0,str(args.skill_dir/'scripts'));import svg_to_pptx as native
    nr=native.create_deck([native.read_svg(out/f) for f in ('emmc_v3.svg','spi_nand_v3.svg')],out/'diagrams.pptx','source',base.FONTFILE,16)
    diagram_deck=Presentation(out/'diagrams.pptx')
    prs=Presentation();prs.slide_width=base.emu(1600);prs.slide_height=base.emu(900)
    p=base.Page(prs,1)
    header(p,'eMMC：3.3 V I/O，8-bit SDR / 50 MHz，优先评估','平台输入：1.2 V 内核、3.3 V PAD（25OD33）、WB 封装；500 MiB 参数开机装入 DDR，允许 30 秒。',
           '建议：8 GB 工业级 eMMC，VCC / VCCQ 均取 3.3 V  |  ≤400 kHz 初始化 → 8-bit High Speed SDR / 50 MHz')
    p.embed(diagram_deck.slides[0],out/'emmc_v3.svg',180)
    section(p,50,422,735,'接口与协议：命令、响应、数据分开传输')
    row(p,50,456,735,'信号 / 引脚',['CLK：主机输出；CMD：双向命令与响应。','DAT[7:0]：双向数据；RST_n：按器件配置使用。'])
    row(p,50,514,735,'数据规则',['按逻辑块读写；主机检查响应、CRC 与超时。','DAT0 可指示忙；无 AXI 式逐拍 VALID / READY。'])
    row(p,50,572,735,'初始化',['识别器件 → 读取能力 → CMD6 配置总线与模式。','先低速初始化，再切 8-bit / High Speed 并提高 CLK。'])
    row(p,50,630,735,'连续读取',['CMD18 多块读 → 接收数据 → DMA 写 DDR。','命令响应 ≠ 搬运完成；检查错误后校验参数包。'])
    row(p,50,688,735,'模式边界',['当前建议 SDR：每周期每根数据线传 1 bit。','3.3 V 不采用 HS200 / HS400；本方案不使用 DS。'])
    section(p,810,422,740,'器件选型与 ASIC 实现：控制器、I/O、时序分层')
    row(p,810,456,740,'候选器件',['ISSI IS21TF08G-JCLI：8 GB、eMMC 5.1。','工业级 -40～85°C，153-ball FBGA；非采购冻结。'],120)
    row(p,810,514,740,'频率 / 带宽',['50 MHz × 8 bit = 50 MB/s ≈ 47.7 MiB/s 线速。','SDR 规格上限 52 MHz；50 MHz 是建议工作点。'],120)
    row(p,810,572,740,'电气 / PHY',['VCC：存储供电 2.7–3.6 V；VCCQ：I/O 取 3.3 V。','该工业级支持 3.3 V；不要误选仅 1.8 V 的 A2 款。'],120)
    row(p,810,630,740,'实现边界',['SDR 可用数字收发 + 合适 PAD，非必需混合 PHY。','仍需核对输入阈值、采样、驱动、WB / PCB 时序。'],120)
    row(p,810,688,740,'工程改动',['新增主机 / DMA、驱动及 PAD 适配；保留 NOR 启动。','同为 3.3 V 可评估直连，须验证电气与上电兼容。'],120)
    p.rect(50,760,1500,64,'#E0EFE9')
    p.text(66,767,1470,'预算：500 MiB ÷ 30 s = 16.7 MiB/s 下限；建议有效搬运 ≥25 MiB/s，20 s 搬运 + 约 10 s 初始化 / 校验。',23,bold=True,color=C['teal'])
    p.text(66,797,1470,'容量：可建议 A/B 两份参数（约 1000 MiB），保留增长余量；更新须完成落盘、校验与版本切换，并验证掉电恢复。',22)
    footer(p,[('ISSI 原厂手册（Mouser 托管）',ISSI,405),('JEDEC 5.1（TI 托管）',JEDEC,325),('备选系列：Swissbit EM-30',SWISSBIT,380),('PHY 实现示例',PHY,220)])
    p.slide.notes_slide.notes_text_frame.text='\n'.join(['需求：500 MiB参数开机完整装入DDR；暂将30秒定义为存储初始化至加载并校验完成。运行时ISP从DDR读取，图仅画加载路径，NOR继续启动。', '2026-10-13用户截图确认：内核1.2V；I/O供电3.3V；PAD为2.5V器件的3.3V Over-Drive方案25OD33；WB封装。2.5V是器件描述，不应据此把外部I/O设为2.5V，亦不推断可用1.8V。未复制受限PDK内容，未读取PAD库。', 'ISSI原厂IS21/22TF08G Rev.A 2022-04-26，第2页明确VCC=2.7–3.6V、VCCQ=1.8V/3.3V，A2车规例外仅1.8V；第36页明确IS21TF08G-JCLI为8GB工业级153FBGA。建议两电源域各按手册接3.3V。VCC是内部存储供电，VCCQ是总线I/O供电；不能只看VCC。完整电气容差、供电时序、最新规格/供货生命周期仍待厂商确认。表中HS200/HS400性能测于VCCQ1.8V，不能移用为3.3V实测值。', '备选Swissbit EM-30系列公开产品页包含4–512GB并支持VCCQ 2.7–3.6V；仅为备选系列，不冻结具体后缀。', 'eMMC采用JEDEC命令/响应/数据，CLK由主机输出，CMD双向，DAT[7:0]双向；初始化1-bit且CLK不超过400kHz；读能力、CMD6完成总线/模式切换后升到50MHz。High Speed SDR上限52MHz，3.3V可用；HS200/HS400仅低压I/O，本方案不使用DS。不是SPI，也不能直接复用SPI Host。', '50MHz x 8bit/8 =50MB/s=47.6837MiB/s；500MiB纯线速10.48576s，不含命令/等待/DMA开销。设计目标25MiB/s即20s，余约10s初始化/校验。具体频率为建议，须确认主机分频能力和有效吞吐。', 'SDR为单端数字接口，可评估数字逻辑+工艺PAD；不自动要求数模混合PHY或高速差分SerDes。3.3V匹配后可评估不加外部电平转换，但仍须VIH/VIL与VOH/VOL、驱动/上拉、上电掉电及钳位兼容、输入采样/输出建立保持、WB寄生/PCB负载/PVT验证。1.2V内核通过合法PAD适配，不能把3.3V直接接内核。不能承诺现有25OD33任意cell达到50MHz。', '当前工程未集成eMMC主机，拟新增控制器/DMA和驱动；IP数字控制器、PAD/PHY及时序模型交付边界需确认。DMA完成、DDR可见、数据校验须分别检查；A/B双副本为建议，容量充裕不等于更新中掉电安全。', '公开来源：ISSI原厂手册（Mouser托管）、JEDEC规范（TI托管）、Swissbit产品介绍、Synopsys示例。链接如下：'] + [ISSI,JEDEC,SWISSBIT,PHY,base.FAQ,'源码：hw/cheshire_soc.sv；hw/bootrom/cheshire_bootrom.c；docs/um/sw.md。'])
    p1=p

    p=base.Page(prs,2)
    header(p,'SPI NAND：改选 3 V 系列，Quad SDR / 100 MHz 待验证','适配 3.3 V PAD：改评估 W25N-LV；原 W25N-MW 为 1.8 V 方案，不再作为直接连接候选。',
           '候选：W25N08LV，8 Gbit（1 GiB）、3 V 档  |  10 MHz 联调 → 50 MHz 验证 → 条件满足后评估 100 MHz')
    p.embed(diagram_deck.slides[1],out/'spi_nand_v3.svg',180)
    section(p,50,422,735,'接口与协议：SPI 命令 + NAND 页访问')
    row(p,50,456,735,'信号 / 模式',['CS#、SCK：主机输出；IO[3:0]：按事务改变方向。','Quad 为四线数据阶段，不等于所有命令均为四线。'])
    row(p,50,514,735,'一次页读取',['发页地址 / 读页命令 → 等待就绪并检查 ECC。','再发读缓存命令，按列地址 / dummy 周期接收数据。'])
    row(p,50,572,735,'软件职责',['识别器件、配置 ECC / Quad、管理坏块与预留块。','检查不可纠错错误；CPU 排空 FIFO 并写入 DDR。'])
    row(p,50,630,735,'候选容量',['8 Gbit = 1 GiB；页为 4096 B + 256 B 备用区。','双份 500 MiB 仅余 24 MiB，坏块预留空间过紧。'])
    section(p,810,422,740,'频率、I/O 与系统吞吐：复用存在约束')
    row(p,810,456,740,'工作频率',['10 / 50 MHz 逐步联调，Quad 100 MHz 为目标。','LV 系列标称 STR 最高 133 MHz，非系统保证。'],120)
    row(p,810,514,740,'时钟前提',['当前 SCK = clk_i / [2 × (CLKDIV + 1)]。','100 MHz SCK 要求 clk_i ≥200 MHz 且分频匹配。'],120)
    row(p,810,572,740,'物理实现',['LV 系列确认 3 V 档；8 Gbit 后缀电气表待确认。','按 3.3 V PAD / WB 评估；100 MHz 须校核裕量。'],120)
    row(p,810,630,740,'主要风险',['现有控制器无独立 DMA；Regbus / CPU 可能限速。','不能只换 Flash：需 NAND 驱动与错误 / 掉电验证。'],120)
    xs=[50,370,665,980,1550]
    data=[('方案 / 建议数据模式','纯线速（不含开销）','500 MiB 线速下限','评估结论'),
          ('eMMC · 8-bit SDR 50 MHz','50 MB/s ≈ 47.7 MiB/s','约 10.5 秒','优先评估：新增主机 / DMA 与 I/O'),
          ('SPI NAND · Quad SDR 100 MHz','50 MB/s ≈ 47.7 MiB/s','约 10.5 秒','备选：时钟、页读取、软件搬运须达标')]
    for j,rowdata in enumerate(data):
        y=701+j*37;p.rect(50,y,1500,37,C['navy'] if j==0 else ('#E9F1EC' if j==1 else '#FFFFFF'))
        for k,t in enumerate(rowdata):p.text(xs[k]+9,y+6,xs[k+1]-xs[k]-18,t,21 if j else 22,color='#FFFFFF' if not j else C['ink'],bold=not j)
    p.text(50,816,1500,'验收口径：完整 500 MiB 初始化 + 加载 + 校验 ≤30 秒；SPI Quad 50 MHz 仅 23.8 MiB/s 线速，低于 25 MiB/s 设计目标。',20,color=C['amber'])
    footer(p,[('Winbond W25N-LV：3 V / 8 Gbit',LV,435),('同系列 4 Gbit：2.7–3.6 V 详情',LV4,430),('NAND 管理职责：Micron',base.NAND,400)])
    p.slide.notes_slide.notes_text_frame.text='\n'.join(['本轮替换原1.8V W25N08MW候选。Winbond W25N-LV官方系列页明确3V、2/4/8Gbit、STR133MHz/DTR66MHz；4/8Gbit页4096+256B。8Gbit即1GiB，不是8GB。本方案只评估Quad STR/SDR，不默认现有Host支持DTR。', '证据分级：W25N08LV是官网已列的3V家族8Gbit方向，尚未取得具体8Gbit订货后缀的完整DC/AC电气表。官网同系列W25N04LVCEIG明确2.7–3.6V，但不能把4Gbit全套参数直接当成8Gbit手册。8Gbit接3.3V是候选设计方向，需供应商确认具体后缀供电容差、VIH/VIL、时序/封装/温度/供货后才冻结。', '如果需要已有明确电压范围的替代评估，可研究两颗W25N04LVCEIG（每颗4Gbit/512MiB，2.7–3.6V）组成1GiB；会新增片选和跨器件软件管理，需与现有NOR的片选占用一起审查，不作为本PPT主方案。单颗512MiB装500MiB只余12MiB，预留不足，不建议作为稳妥单颗容量方案。', '读页命令带行地址，将array读入cache；等待Busy清除检查ECC，再按列地址/dummy读cache。内部ECC和BBM LUT不能代替驱动坏块/预留/不可纠错处理与参数包校验；不能复用NOR协议驱动。图为功能链，不限定内部ECC物理位置。', '10MHz首次联调、50MHz过渡验证均为工程建议，不是协议强制初始化频率。100MHz Quad STR目标须在3.3V PAD、WB和PCB上确认采样/驱动和时序；器件133MHz标称上限不代表当前系统可达。', '当前SPI Host SCK=clk_i/[2*(CLKDIV+1)]；100MHz SCK要求输入至少200MHz且整比分频匹配。当前公共SoC clk_i接入，不能只改divider就保证该速率；未改频率或新增CDC。现有SPI Host无独立DMA口，经Regbus/FIFO由CPU读出写DDR；这条路径是吞吐瓶颈候选，需实测。', '100MHz Quad线速50MB/s≈47.7MiB/s，500MiB纯线速约10.5s；25MiB/s目标是20s搬运+约10s初始化校验。50MHz Quad线速23.8MiB/s低于25目标，但不单独证明超过30s；端到端含页等待、软件访问、DDR写、初始化校验测量才决定。', '两份500MiB共1000MiB；1GiB只余24MiB，坏块/预留/元数据与增长余量不足，A/B可靠更新只是建议。eMMC优先；SPI NAND为电气、时钟与软件吞吐达标后的备选。', '源码：hw/cheshire_soc.sv:1445；.bender/git/checkouts/opentitan_peripherals-7b624fb57f78de9a/src/spi_host/rtl/spi_host_fsm.sv；该依赖dif_spi_host.c；sw/lib/hal/spi_sdcard.c。未改RTL/software、未实测IO或吞吐。公开链接如下：'] + [LV,LV4,base.NAND])
    metrics=[p1.save(out),p.save(out)]
    final=out/'nonvolatile_storage_options_v3.pptx';prs.save(final)
    rr=Presentation(final);checks=[]
    for n,s in enumerate(rr.slides,1):
        ids=[x.get('id') for x in s._element.xpath('.//p:cNvPr')];assert len(ids)==len(set(ids))
        urls=[r.hyperlink.address for sh in s.shapes if sh.has_text_frame for para in sh.text_frame.paragraphs for r in para.runs if r.hyperlink.address]
        assert len(urls)==(4 if n==1 else 3)
        for sh in s.shapes:
            assert sh.left>=0 and sh.top>=0
            assert sh.left+sh.width<=rr.slide_width+base.emu(2),(n,sh.name)
            assert sh.top+sh.height<=rr.slide_height+base.emu(2),(n,sh.name)
        checks.append(dict(page=n,objects=len(s.shapes),links=urls,arrows=len(s._element.xpath('.//a:tailEnd')),notes_chars=len(s.notes_slide.notes_text_frame.text)))
    with zipfile.ZipFile(final) as z:assert not any(n.startswith('ppt/media/') for n in z.namelist())
    source=json.loads((ROOT/'output/source_manifest.json').read_text())
    f='.bender/git/checkouts/opentitan_peripherals-7b624fb57f78de9a/src/spi_host/rtl/spi_host_fsm.sv'
    source[f]=hashlib.sha256((base.REPO/f).read_bytes()).hexdigest()
    for f,sha in source.items():assert hashlib.sha256((base.REPO/f).read_bytes()).hexdigest()==sha,f
    (out/'source_manifest.json').write_text(json.dumps(source,indent=2)+'\n')
    report={'revision':3,'slides':checks,'text_metrics':metrics,'native_diagram_report':nr,'office_verified':False,
            'preview':'CairoSVG from shared layout; not Office rendering','files':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in out.iterdir() if p.is_file()}}
    (out/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(checks,ensure_ascii=False));print(final)

if __name__=='__main__':main()
