#!/usr/bin/env python3
"""Verify the appended page, protected pages and native editable objects."""
import argparse
import hashlib
import json
from pathlib import Path
import posixpath
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
from build import NS, CAPTION
from append_uniform_page import CENTERS, GROUPS
from validate import objects, box, text, points, intersects, summary

D=Path(__file__).resolve().parent
ROOT=D.parents[2]
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()


def close(a,b):
    # Legacy imported shapes use coordinates ~0.005 px below the rounded labels.
    return all(abs(x-y)<.01 for x,y in zip(a,b))


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--deck',type=Path,required=True)
    ap.add_argument('--report',type=Path,required=True)
    ap.add_argument('--preview',type=Path,required=True)
    args=ap.parse_args();assert not args.report.exists()
    inputs=json.loads((D/'evidence/page3-input.json').read_text())
    protected={p:sha(ROOT/p)==v for p,v in inputs['protected'].items()}
    assert all(protected.values()),[p for p,v in protected.items() if not v]
    assert all(sha(ROOT/p)==v for p,v in inputs['backups'].items())
    src=D/'evidence/page3-before/soc_landscape_m5_gb.pptx'
    assert sha(src)==inputs['input_sha256']
    with ZipFile(src) as zi,ZipFile(args.deck) as zo:
        assert zo.testzip() is None
        changed=[n for n in zi.namelist() if zi.read(n)!=zo.read(n)]
        assert set(changed)=={'ppt/presentation.xml','ppt/_rels/presentation.xml.rels','docProps/app.xml','[Content_Types].xml'},changed
        added=set(zo.namelist())-set(zi.namelist())
        assert added=={'ppt/slides/slide3.xml','ppt/slides/_rels/slide3.xml.rels','ppt/notesSlides/notesSlide3.xml','ppt/notesSlides/_rels/notesSlide3.xml.rels'}
        assert not any(n.startswith('ppt/media/') for n in zo.namelist())
        for n in zo.namelist():
            if not n.endswith('.rels'):continue
            rels=E.fromstring(zo.read(n));ids=[r.get('Id') for r in rels]
            assert len(ids)==len(set(ids))
            base=posixpath.dirname(posixpath.dirname(n)) if n!='_rels/.rels' else ''
            for r in rels:
                if r.get('TargetMode')=='External':continue
                target=posixpath.normpath(posixpath.join(base,r.get('Target'))).lstrip('/')
                assert target in zo.namelist(),(n,target)
        before=objects(E.fromstring(zi.read('ppt/slides/slide2.xml')))
        root=E.fromstring(zo.read('ppt/slides/slide3.xml'));after=objects(root)
        assert before.keys()==after.keys()
        ids=root.xpath('//p:cNvPr/@id',namespaces=NS);assert len(ids)==len(set(ids))
        altered={i for i in before if E.tostring(before[i])!=E.tostring(after[i])}
        expected={860,911,958,1086,1301,*range(948,955)}
        for group in GROUPS:expected.update(group[:4]);expected.update(group[4])
        assert altered==expected,(altered^expected)
        # Text, original font sizes, stroke styles and arrow directions are preserved.
        for i in before:
            assert text(before[i])==text(after[i]),i
            assert [E.tostring(r) for r in before[i].findall('.//a:rPr',NS)]==[E.tostring(r) for r in after[i].findall('.//a:rPr',NS)],i
            l0=before[i].find('p:spPr/a:ln',NS);l1=after[i].find('p:spPr/a:ln',NS)
            assert (E.tostring(l0) if l0 is not None else None)==(E.tostring(l1) if l1 is not None else None),i
        assert text(after[886])==CAPTION
        names=[el.find('.//p:cNvPr',NS).get('name') for el in after.values()]
        assert names.count('global-buffer')==1
        owners={};centers=[]
        for idx,(cx,group) in enumerate(zip(CENTERS,GROUPS)):
            drop,port,label,target,labels=group
            x,y,w,h=box(after[port]);centers.append(x+w/2)
            assert close((x+w/2,y,w,h),(cx,1013.1,96,29))
            assert text(after[label])==(['M0','M1','M2','M3','M4','M5','M_EXT'][idx])
            x,y,w,h=box(after[target]);assert close((x+w/2,y,h),(cx,1093.1,155))
            assert close(points(after[drop])[0],(cx,1042.1))
            assert close(points(after[drop])[-1],(cx,1093.1))
            owners[label]=port;owners.update({i:target for i in labels})
        pitch=[b-a for a,b in zip(centers,centers[1:])]
        assert max(pitch)-min(pitch)<.001
        for g0,g1 in zip(GROUPS,GROUPS[1:]):
            x,y,w,h=box(after[g0[3]]);xx,yy,ww,hh=box(after[g1[3]])
            assert xx-x-w>20
        owners.update({949:948,950:948,952:951,953:951})
        assert close(points(after[1301])[0],(CENTERS[5]+180,1170))
        assert close(points(after[1301])[-1],points(before[860])[-1])
        assert close(points(after[860])[-1],points(before[1301])[-1])
        assert close(points(after[860])[0],points(before[860])[0])
        assert close(points(after[1086])[0],(CENTERS[2]+105,1237.1))
        assert close(points(after[1086])[-1],points(before[1086])[-1])
        routes={i:points(after[i]) for i in [860,958,1086,1301,954,*[g[0] for g in GROUPS]]}
        for i in [860,1086,1301]:
            path=after[i].find('p:spPr/a:custGeom/a:pathLst/a:path',NS)
            assert len(path.findall('a:moveTo',NS))==1
            assert len(path.findall('a:cubicBezTo',NS))==(0 if i==1086 else 1)
        # The two GB bottom-entry paths no longer cross. Keep only the two
        # intentional bridges: ISP/DDR and M5/M_EXT.
        def strict_cross(a,b,c,d):
            cross=lambda u,v:u[0]*v[1]-u[1]*v[0]
            sub=lambda u,v:(u[0]-v[0],u[1]-v[1])
            den=cross(sub(b,a),sub(d,c))
            if abs(den)<1e-8:return False
            t=cross(sub(c,a),sub(d,c))/den;u=cross(sub(c,a),sub(b,a))/den
            return 0<t<1 and 0<u<1
        assert not any(strict_cross(a,b,c,d) for a,b in zip(routes[860],routes[860][1:])
                       for c,d in zip(routes[1301],routes[1301][1:]))
        obstacles=[781,785,790,796,804,809,813,821,831,836,839,844,847,
                   862,865,868,871,874,877,884,915,921,927,932,939,945,1304]
        # The atomics->LLC path is intentionally internal to the memory target.
        for i,path in routes.items():
            for a,b in zip(path,path[1:]):
                for id in obstacles:
                    if id==927 and i in [1086,954]:continue
                    x,y,w,h=box(after[id])
                    assert not intersects(a,b,(x+.01,y+.01,w-.02,h-.02)),(i,id)
    layout=json.loads((args.preview/'preview_layout.json').read_text())['slides'][2]
    for t in layout['texts']:
        i=int(t['shape_id']);x,y,w,h=t['box']
        if i in owners:
            bx,by,bw,bh=box(after[owners[i]])
            assert bx<=x and by<=y and x+w<=bx+bw and y+h<=by+bh,(i,t['box'])
        if i in altered:
            for u in layout['texts']:
                if u is t:continue
                xx,yy,ww,hh=u['box']
                assert not (min(x+w,xx+ww)-max(x,xx)>1 and min(y+h,yy+hh)-max(y,yy)>1),(t,u)
        for line_id,path in routes.items():
            assert not any(intersects(a,b,t['box']) for a,b in zip(path,path[1:])),(line_id,t)
    for i in [1,2]:
        for suffix in ['svg','png']:
            assert (args.preview/f'slide-{i:02d}.{suffix}').read_bytes()==(D/f'slide-{i:02d}.{suffix}').read_bytes()
    # Validate that the inspected preview has the exact same slide geometries
    # as the final candidate (build 02 only corrects the new notes page number).
    with ZipFile(D/'evidence/page3-build-01/soc_landscape_m5_gb.pptx') as preview_source,ZipFile(args.deck) as final:
        assert all(preview_source.read(f'ppt/slides/slide{i}.xml')==final.read(f'ppt/slides/slide{i}.xml') for i in [1,2,3])
    prs=Presentation(args.deck);assert len(prs.slides)==3
    stats=summary(prs)
    rt=args.report.with_suffix('.roundtrip.pptx');assert not rt.exists()
    prs.save(rt);assert summary(Presentation(rt))==stats
    assert stats[1]==stats[2]
    report={'status':'pass','input_sha256':sha(src),'output_sha256':sha(args.deck),
            'pages':3,'pages1_and2_all_parts_byte_identical':True,'pages1_and2_previews_byte_identical':True,
            'changed_existing_package_parts':changed,'added_package_parts':sorted(added),
            'changed_shapes_on_page3':sorted(altered),'unchanged_shapes_on_page3':len(before)-len(altered),
            'target_centers':centers,'adjacent_spacing':pitch,'target_alignment':'pass',
            'gb_bottom_entries_swapped':True,'m5_isp_paths_no_longer_cross':True,
            'same_text_fonts_arrow_styles':True,'native_paths_continuous':True,
            'changed_text_bounds_and_overlap':'pass','modified_routes_avoid_modules_and_text':'pass',
            'native_save_reopen':'pass','package_relationships':'pass',
            'slides':[{**{k:v for k,v in s.items() if k!='texts'},'text_objects':len(s['texts'])} for s in stats],
            'protected_files':protected,'office_application_tested':False,'hardware_tested':False,
            'preview_type':'DrawingML geometry and fallback fonts, not Office screenshots'}
    args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k not in ['protected_files','changed_shapes_on_page3']},ensure_ascii=False,indent=2))


if __name__=='__main__':main()
