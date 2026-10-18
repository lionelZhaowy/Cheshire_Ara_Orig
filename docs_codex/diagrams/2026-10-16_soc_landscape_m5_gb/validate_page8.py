#!/usr/bin/env python3
"""Audit the five geometry edits against the actual manually saved input."""
import argparse,hashlib,json,posixpath
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
from validate import NS,objects,points,intersects,text,summary
from build_page6 import named
from build_page8 import EDITED,EXPECTED,box,commands

D=Path(__file__).resolve().parent;ROOT=D.parents[2]
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def semantic(e):return None if e is None else (e.tag,tuple(sorted(e.attrib.items())),e.text,tuple(semantic(c) for c in e))
def close(a,b):return len(a)==len(b) and all(abs(x-y)<.01 for x,y in zip(a,b))
def on_boundary(p,b):
    x,y=p;bx,by,w,h=b
    return bx-.01<=x<=bx+w+.01 and by-.01<=y<=by+h+.01 and min(abs(x-bx),abs(x-bx-w),abs(y-by),abs(y-by-h))<.01

def skeleton(el):
    """Remove only the bridge curvature, keeping every native endpoint/corner."""
    ps=[c[1][-1] for c in commands(el)];out=[]
    for p in ps:
        out.append(p)
        while len(out)>=3:
            a,b,c=out[-3:]
            if (abs(a[0]-b[0])<.01 and abs(b[0]-c[0])<.01) or (abs(a[1]-b[1])<.01 and abs(b[1]-c[1])<.01):out.pop(-2)
            else:break
    return out

def crossings(paths):
    result=[];names=list(paths)
    for i,name in enumerate(names):
        for other in names[i+1:]:
            if not ({name,other}&EDITED):continue
            for a,b in zip(paths[name],paths[name][1:]):
                for c,d in zip(paths[other],paths[other][1:]):
                    if abs(a[1]-b[1])<.01 and abs(c[0]-d[0])<.01:
                        x,y=c[0],a[1];hit=min(a[0],b[0])+.1<x<max(a[0],b[0])-.1 and min(c[1],d[1])+.1<y<max(c[1],d[1])-.1
                    elif abs(a[0]-b[0])<.01 and abs(c[1]-d[1])<.01:
                        x,y=a[0],c[1];hit=min(c[0],d[0])+.1<x<max(c[0],d[0])-.1 and min(a[1],b[1])+.1<y<max(a[1],b[1])-.1
                    else:continue
                    if hit:result.append(dict(pair=sorted([name,other]),at=[x,y]))
    return result

def audit(build,preview,report,deck):
    assert not report.exists();snap=json.loads((D/'evidence/page8-input.json').read_text())
    source=D/'evidence/page8-before/soc_landscape_m5_gb.pptx';assert sha(source)==EXPECTED==snap['input_sha256']
    protected={p:sha(ROOT/p)==h for p,h in snap['protected'].items()};assert all(protected.values())
    assert all(sha(ROOT/p)==h for p,h in snap['backups'].items())
    spec=json.loads((build/'edit_report.json').read_text());assert len(spec['bridges'])==6
    with ZipFile(source) as zi,ZipFile(deck) as zo:
        assert zo.testzip() is None
        changed=[n for n in zi.namelist() if zi.read(n)!=zo.read(n)]
        assert set(changed)=={'[Content_Types].xml','ppt/presentation.xml','ppt/_rels/presentation.xml.rels','docProps/app.xml'}
        added=set(zo.namelist())-set(zi.namelist())
        assert added=={'ppt/slides/slide8.xml','ppt/slides/_rels/slide8.xml.rels','ppt/notesSlides/notesSlide8.xml','ppt/notesSlides/_rels/notesSlide8.xml.rels'}
        for name in zo.namelist():
            if not name.endswith('.rels'):continue
            rels=E.fromstring(zo.read(name));ids=[r.get('Id') for r in rels];assert len(ids)==len(set(ids))
            parent=posixpath.dirname(posixpath.dirname(name)) if name!='_rels/.rels' else ''
            for r in rels:
                if r.get('TargetMode')!='External':assert posixpath.normpath(posixpath.join(parent,r.get('Target'))).lstrip('/') in zo.namelist()
        old=E.fromstring(zi.read('ppt/slides/slide7.xml'));slide=E.fromstring(zo.read('ppt/slides/slide8.xml'))
        before=objects(old);after=objects(slide);native=named(slide);oldnative=named(old)
        assert set(before)==set(after);ids=slide.xpath('//p:cNvPr/@id',namespaces=NS);assert len(ids)==len(set(ids))
        clean=deepcopy(slide);cleanobjs=objects(clean);changed_objects=[]
        for id,oldel in before.items():
            el=after[id];name=el.find('.//p:cNvPr',NS).get('name').split(': ')[0]
            assert semantic(el.find('p:spPr/a:ln',NS))==semantic(oldel.find('p:spPr/a:ln',NS)),name
            if semantic(oldel)==semantic(el):continue
            changed_objects.append(name);assert name in EDITED
            sp=cleanobjs[id].find('p:spPr',NS)
            for tag in ['xfrm','custGeom']:sp.replace(sp.find('a:'+tag,NS),deepcopy(oldel.find('p:spPr/a:'+tag,NS)))
        assert set(changed_objects)==EDITED
        assert semantic(clean)==semantic(old),'Restoring the five geometries must reproduce the complete input page'
        notes=E.fromstring(zo.read('ppt/notesSlides/notesSlide8.xml'))
        assert notes.xpath('//a:fld[@type="slidenum"]/a:t/text()',namespaces=NS)==['8']
    paths={n:skeleton(el) for n,el in native.items() if el.find('p:spPr/a:custGeom',NS) is not None and el.find('p:spPr/a:ln',NS) is not None}
    actual=crossings(paths);assert len(actual)==6,actual
    for hit in actual:
        matching=[b for b in spec['bridges'] if sorted([b['edge'],b['over']])==hit['pair'] and close(b['at'],hit['at'])]
        assert len(matching)==1,hit
    for name,r in spec['routes'].items():
        el=native[name];actual_pts=points(el)
        assert close(actual_pts[0],r['points'][0]) and close(actual_pts[-1],r['points'][-1]),name
        assert on_boundary(actual_pts[0],box(native[r['src']])) and on_boundary(actual_pts[-1],box(native[r['dst']])),name
        assert len(paths[name])==len(r['points']) and all(close(a,b) for a,b in zip(paths[name],r['points'])),name
        path=el.find('p:spPr/a:custGeom/a:pathLst/a:path',NS)
        assert len(path.findall('a:moveTo',NS))==1
        curves=[c for c in commands(el) if c[0]=='cubicBezTo'];wanted=[b for b in spec['bridges'] if b['edge']==name]
        assert len(curves)==len(wanted),(name,len(curves))
        for c,b in zip(curves,sorted(wanted,key=lambda j: (j['at'][1] if name=='isp-gb' else (-j['at'][0] if name=='emmc-physical' else j['at'][0])))):
            controls=c[1][:2];x,y=b['at']
            if b['direction']=='right':assert min(p[0] for p in controls)>x
            elif b['direction']=='down':assert min(p[1] for p in controls)>y
            else:assert max(p[1] for p in controls)<y
    # Independently enforce the expected interface styles and semantic anchors.
    for name in ['mover-f1','frame-f2','roi-f3-request','isp-gb']:
        ln=native[name].find('p:spPr/a:ln',NS)
        assert ln.find('a:custDash',NS) is not None
        assert all(ln.find('a:'+end,NS).get('type')=='triangle' for end in ['headEnd','tailEnd'])
    for name in ['llc-f0']+['front-ddr-'+str(i) for i in range(4)]:
        ln=native[name].find('p:spPr/a:ln',NS);assert ln.find('a:custDash',NS) is None
        assert all(ln.find('a:'+end,NS).get('type')=='triangle' for end in ['headEnd','tailEnd'])
    for name,port in [('mover-f1','front-F1'),('frame-f2','front-F2')]:
        b=box(native[port]);assert close(points(native[name])[-1],[b[0]+b[2]/2,b[1]])
    assert close([paths['frame-f2'][2][0]-paths['mover-f1'][1][0]],[30])
    assert paths['mover-f1'][-2][1]<paths['frame-f2'][-2][1]<paths['roi-f3-request'][-2][1]
    for i in range(1,17):
        for prefix in ['badge-C','reg-badge-C']:assert text(native[f'{prefix}{i}-text'])==f'C{i}'
    assert text(native['roi-vdma-detail'])=='二维 AXI 读'
    assert list(native).count('global-buffer')==1
    layout=json.loads((preview/'preview_layout.json').read_text())['slides'][7]
    regions={'soc-boundary','compute','gb-region','image-rx','reg-region','ddr-controller','roi-region','compression-front','shared-payload-services','memory-target'}
    obstacles={n:box(el) for n,el in native.items() if el.find('p:spPr/a:prstGeom[@prst="rect"]',NS) is not None and not text(el) and n not in regions}
    issues={'line_through_text':[],'line_through_box':[],'text_overlap':[]}
    id_names={id:el.find('.//p:cNvPr',NS).get('name').split(': ')[0] for id,el in after.items()}
    for name in EDITED:
        ps=points(native[name])
        for t in layout['texts']:
            if any(intersects(a,b,t['box'],1) for a,b in zip(ps,ps[1:])):issues['line_through_text'].append([name,id_names[int(t['shape_id'])]])
        for bn,b in obstacles.items():
            r=[b[0]+.02,b[1]+.02,b[2]-.04,b[3]-.04]
            if any(intersects(a,c,r) for a,c in zip(ps,ps[1:])):issues['line_through_box'].append([name,bn])
    for i,t in enumerate(layout['texts']):
        x,y,w,h=t['box']
        for u in layout['texts'][i+1:]:
            xx,yy,ww,hh=u['box']
            if min(x+w,xx+ww)-max(x,xx)>1 and min(y+h,yy+hh)-max(y,yy)>1:issues['text_overlap'].append([id_names[int(t['shape_id'])],id_names[int(u['shape_id'])]])
    assert not any(issues.values()),issues
    # The final legend must remain wholly to the right of the long corridors.
    assert min(box(el)[0] for n,el in native.items() if n.startswith('legend-'))>3005
    prs=Presentation(deck);assert len(prs.slides)==8;stats=summary(prs);rt=report.with_suffix('.roundtrip.pptx');assert not rt.exists();prs.save(rt)
    assert summary(Presentation(rt))==stats
    with ZipFile(rt) as z:
        reread=named(E.fromstring(z.read('ppt/slides/slide8.xml')))
        for n,el in native.items():
            assert text(reread[n])==text(el)
            assert semantic(reread[n].find('p:spPr',NS))==semantic(el.find('p:spPr',NS))
    result=dict(status='pass',date='2026-10-17',head=snap['head'],input_sha256=sha(source),output_sha256=sha(deck),pages=8,
        preserved_pages=list(range(1,8)),changed_existing_package_parts=changed,new_package_parts=sorted(added),
        modified_objects=sorted(changed_objects),unchanged_objects=len(before)-len(changed_objects),
        only_xfrm_and_custGeom_changed=True,all_style_font_text_and_other_geometry_preserved=True,
        actual_related_crossings=actual,bridges=spec['bridges'],layout=issues,protected_file_count=len(protected),
        native_save_reopen='pass',package_relationships='pass',preview_source=str(deck.relative_to(ROOT)),
        slides=[{**{k:v for k,v in s.items() if k!='texts'},'text_objects':len(s['texts'])} for s in stats],office_tested=False,hardware_tested=False)
    report.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:result[k] for k in ['status','output_sha256','modified_objects','unchanged_objects','actual_related_crossings','layout']},ensure_ascii=False,indent=2))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--build',type=Path,required=True);p.add_argument('--preview',type=Path,required=True);p.add_argument('--report',type=Path,required=True);p.add_argument('--deck',type=Path)
    a=p.parse_args();audit(a.build.resolve(),a.preview.resolve(),a.report.resolve(),(a.deck or a.build/'soc_landscape_m5_gb.pptx').resolve())
