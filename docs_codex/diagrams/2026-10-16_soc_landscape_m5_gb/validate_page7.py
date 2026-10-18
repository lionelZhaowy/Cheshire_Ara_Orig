#!/usr/bin/env python3
"""Scope, native style, configuration and connection audit for the appended page."""
import argparse,hashlib,json,posixpath
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
from validate import NS,objects,points,box,text,summary
from inspect_page7 import inspect
from build_page7 import NEW_TO_OLD,OLD_TO_NEW,EDIT_EDGES,NEW_EDGES,MIRROR_BLOCKS
from build_page6 import named

D=Path(__file__).resolve().parent;ROOT=D.parents[2]
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def semantic(e):return None if e is None else (e.tag,tuple(sorted(e.attrib.items())),e.text,tuple(semantic(c) for c in e))
def close(a,b):return len(a)==len(b) and all(abs(x-y)<.01 for x,y in zip(a,b))
def on_boundary(p,b):
    x,y=p;bx,by,bw,bh=[b[k] for k in ['x','y','w','h']]
    return bx-.01<=x<=bx+bw+.01 and by-.01<=y<=by+bh+.01 and min(abs(x-bx),abs(x-bx-bw),abs(y-by),abs(y-by-bh))<.01

def audit(build,preview,report,deck=None):
    assert not report.exists()
    snap=json.loads((D/'evidence/page7-input.json').read_text());src=D/'evidence/page7-before/soc_landscape_m5_gb.pptx'
    dst=deck or build/'soc_landscape_m5_gb.pptx';assert sha(src)==snap['input_sha256']
    protected={n:sha(ROOT/n)==v for n,v in snap['protected'].items()};assert all(protected.values())
    assert all(sha(ROOT/n)==v for n,v in snap['backups'].items())
    edits=json.loads((build/'edit_report.json').read_text())
    allowed=set(edits['moved'])|set(edits['renamed'])|set(edits['text_changed'])|EDIT_EDGES
    model=json.loads((build/'slide-07.json').read_text());base=json.loads((D/'evidence/page6-f3-build/slide-06.json').read_text())
    boxes={b['id']:b for b in model['boxes']+model['regions']};texts={t['id']:t for t in model['texts']};edges={e['id']:e for e in model['edges']}
    with ZipFile(src) as zi,ZipFile(dst) as zo:
        assert zo.testzip() is None
        changed=[n for n in zi.namelist() if zi.read(n)!=zo.read(n)]
        assert set(changed)=={'[Content_Types].xml','ppt/presentation.xml','ppt/_rels/presentation.xml.rels','docProps/app.xml'}
        added=set(zo.namelist())-set(zi.namelist())
        assert added=={'ppt/slides/slide7.xml','ppt/slides/_rels/slide7.xml.rels','ppt/notesSlides/notesSlide7.xml','ppt/notesSlides/_rels/notesSlide7.xml.rels'}
        for n in zo.namelist():
            if not n.endswith('.rels'):continue
            rels=E.fromstring(zo.read(n));ids=[r.get('Id') for r in rels];assert len(ids)==len(set(ids))
            parent=posixpath.dirname(posixpath.dirname(n)) if n!='_rels/.rels' else ''
            for r in rels:
                if r.get('TargetMode')!='External':assert posixpath.normpath(posixpath.join(parent,r.get('Target'))).lstrip('/') in zo.namelist()
        old=E.fromstring(zi.read('ppt/slides/slide6.xml'));slide=E.fromstring(zo.read('ppt/slides/slide7.xml'))
        oldobjs=objects(old);objs=objects(slide);native=named(slide);oldnative=named(old)
        assert set(oldobjs).issubset(objs) and len(objs)==len(oldobjs)+8
        ids=slide.xpath('//p:cNvPr/@id',namespaces=NS);assert len(ids)==len(set(ids))
        modified=[];untouched=0
        names={i:o.find('.//p:cNvPr',NS).get('name').split(': ')[0] for i,o in objs.items()}
        for id,before in oldobjs.items():
            after=objs[id];name=before.find('.//p:cNvPr',NS).get('name').split(': ')[0]
            if semantic(before)==semantic(after):untouched+=1;continue
            modified.append(name);assert name in allowed,name
            # Whitelist the allowed properties, then require all other native
            # content (fills, theme modifiers, borders, fonts, arrows) to match.
            cleaned=deepcopy(after)
            if name in edits['renamed'] or name in edits['text_changed']:
                cleaned.find('.//p:cNvPr',NS).set('name',before.find('.//p:cNvPr',NS).get('name'))
            if name in edits['moved'] or name in EDIT_EDGES or (name in edits['text_changed'] and 'badge-' in name):
                sp=cleaned.find('p:spPr',NS);sp.replace(sp.find('a:xfrm',NS),deepcopy(before.find('p:spPr/a:xfrm',NS)))
            if name in EDIT_EDGES:
                sp=cleaned.find('p:spPr',NS);sp.replace(sp.find('a:custGeom',NS),deepcopy(before.find('p:spPr/a:custGeom',NS)))
            if name in edits['text_changed']:
                cleaned.find('.//a:t',NS).text=before.find('.//a:t',NS).text
            assert semantic(cleaned)==semantic(before),name
        stripped=[]
        for root in [old,slide]:
            root=deepcopy(root);tree=root.find('p:cSld/p:spTree',NS)
            for el in list(tree):
                if E.QName(el).localname in ['sp','cxnSp']:tree.remove(el)
            stripped.append(semantic(root))
        assert stripped[0]==stripped[1]
        notes=E.fromstring(zo.read('ppt/notesSlides/notesSlide7.xml'))
        assert notes.xpath('//a:fld[@type="slidenum"]/a:t/text()',namespaces=NS)==['7']
    # Model positions and labels must match actual native objects.
    for name,b in boxes.items():assert close(box(native[name]),[b[k] for k in ['x','y','w','h']]),name
    for name,t in texts.items():assert text(native[name])==t['text'],name
    for e in model['edges']:
        el=native[e['id']];pts=points(el)
        assert close(pts[0],e['points'][0]) and close(pts[-1],e['points'][-1]),e['id']
        paths=el.findall('p:spPr/a:custGeom/a:pathLst/a:path',NS)
        assert len(paths)==1 and len(paths[0].findall('a:moveTo',NS))==1
        for tag,wanted in [('tailEnd',e['arrow']),('headEnd',e['both'])]:
            end=el.find('p:spPr/a:ln/a:'+tag,NS);assert (end is not None and end.get('type')=='triangle')==wanted,e['id']
        assert (el.find('p:spPr/a:ln/a:custDash',NS) is not None)==e['dash'],e['id']
        if e['id'] in EDIT_EDGES|NEW_EDGES:
            assert on_boundary(pts[0],boxes[e['src']]) and on_boundary(pts[-1],boxes[e['dst']]),e['id']
    # Check all inherited edge semantics; only seven routes change geometry.
    for olde in base['edges']:
        e=deepcopy(edges[olde['id']]);e['points']=olde['points'];assert e==olde,olde['id']
        assert semantic(native[olde['id']].find('p:spPr/a:ln',NS))==semantic(oldnative[olde['id']].find('p:spPr/a:ln',NS))
    for name in ['mover-f1','frame-f2','roi-f3-request']:
        assert edges[name]['both'] and edges[name]['dash'] and edges[name]['kind']=='data'
    for name in ['llc-f0']+['front-ddr-'+str(i) for i in range(4)]:
        assert edges[name]['both'] and not edges[name]['dash'] and edges[name]['kind']=='data'
    for name in NEW_EDGES:assert edges[name]['kind']=='pixel' and edges[name]['arrow'] and not edges[name]['both'] and not edges[name]['dash']
    # Independent expected function mapping, not inferred from the result labels.
    expected={1:['PLL / 时钟','复位'],2:['DDR','Controller'],3:['DDR','PHY'],4:['LVDS','RX'],5:['LVDS','TX'],
        6:['帧写入','DMA'],7:['帧输出','DMA'],8:['ROI 任务','/ 调度'],9:['ISP'],10:['NPU','计算'],11:['GB','搬运'],
        12:['GB','管理'],13:['eMMC Host','/ DMA'],14:['压缩公共','控制'],15:['压缩器'],16:['解压器']}
    expected_module={1:'pll',2:'ddr-controller',3:'ddr-phy',4:'lvds-rx',5:'lvds-tx',6:'frame-dma',7:'roi-vdma',8:'roi-scheduler',
        9:'isp',10:'npu',11:'gb-mover',12:'global-buffer',13:'emmc-host',14:'front-adapt',15:'compressor',16:'decompressor'}
    for i,ls in expected.items():
        assert [text(native[f'config-C{i}-title-{k}']) for k in range(len(ls))]==ls
        for prefix in ['badge-C','reg-badge-C']:assert text(native[f'{prefix}{i}-text'])==f'C{i}'
        b=boxes[f'config-C{i}'];assert close([b['x'],b['y']],[71.25+((i-1)%8)*211.75,1745.1951771653542+((i-1)//8)*134])
        assert semantic(native[f'cfg-C{i}'])==semantic(oldnative[f'cfg-C{i}'])
        # Module badge must stay in its original functional block's corner.
        old_i=NEW_TO_OLD[i];bm=boxes[f'badge-C{i}'];owner=boxes[expected_module[i]]
        assert owner['x']-.1<=bm['x']<=owner['x']+owner['w']+.1 and owner['y']-40<=bm['y']<=owner['y']+owner['h'],(i,expected_module[i])
        assert int(native[f'badge-C{i}'].find('.//p:cNvPr',NS).get('id'))==int(oldnative[f'badge-C{old_i}'].find('.//p:cNvPr',NS).get('id'))
    for prefix in ['config-C','reg-badge-C','badge-C']:
        actual=[n for n in native if n.startswith(prefix) and n[len(prefix):].isdigit()]
        assert set(actual)=={prefix+str(i) for i in range(1,17)}
    assert text(native['roi-vdma-title'])=='帧输出 DMA' and text(native['roi-vdma-detail'])=='二维 AXI 读'
    assert sum(b['id']=='global-buffer' for b in model['boxes'])==1
    # Follow the actual source/target chain, rather than infer it from box order.
    for a,b in zip(MIRROR_BLOCKS,MIRROR_BLOCKS[1:]):
        e=edges[a+'-to-'+b];assert e['src']==a and e['dst']==b and e['points'][0][0]>e['points'][-1][0]
    assert edges['sensor-rx']['src']=='image-sensor' and edges['sensor-rx']['dst']=='lvds-rx'
    assert edges['raw-bypass']['src']=='rx-cdc' and edges['raw-bypass']['dst']=='frame-select'
    assert edges['isp-gb']['src']=='isp' and edges['isp-gb']['dst']=='gb-adapter' and edges['isp-gb']['dash'] and edges['isp-gb']['both']
    assert edges['frame-f2']['src']=='frame-dma' and edges['frame-f2']['dst']=='front-F2'
    assert edges['tx-processor']['src']=='lvds-tx' and edges['tx-processor']['dst']=='output-processor'
    assert any(b['edge']=='frame-f2' and b['over']=='isp-gb' for b in model['bridges'])
    # External boxes inherit the actual eMMC style, not a guessed theme color.
    for name in ['image-sensor','output-processor']:
        for xp in ['p:spPr/a:solidFill','p:spPr/a:ln','p:style']:
            assert semantic(native[name].find(xp,NS))==semantic(native['emmc-device'].find(xp,NS))
        for j in [0,1]:assert native[name+f'-t{j}'].xpath('.//a:rPr/@sz',namespaces=NS)==native[f'emmc-device-t{j}'].xpath('.//a:rPr/@sz',namespaces=NS)
    layout=inspect(build,preview);assert not any(layout.values()),layout
    # Badges are separate shapes: check them explicitly against unrelated text.
    prev=json.loads((preview/'preview_layout.json').read_text())['slides'][6];badge_issues=[]
    for t in prev['texts']:
        name=names[int(t['shape_id'])];x,y,w,h=t['box']
        for b in [v for v in boxes.values() if v['id'].startswith(('badge-','reg-badge-'))]:
            if texts.get(name,{}).get('owner')==b['id']:continue
            if min(x+w,b['x']+b['w'])>max(x,b['x'])+.1 and min(y+h,b['y']+b['h'])>max(y,b['y'])+.1:badge_issues.append([name,b['id']])
    assert not badge_issues;layout['text_through_badge_frame']=badge_issues
    prs=Presentation(dst);assert len(prs.slides)==7;stats=summary(prs)
    rt=report.with_suffix('.roundtrip.pptx');assert not rt.exists();prs.save(rt);assert summary(Presentation(rt))==stats
    # Read back the saved arrow styles and text, not just counts.
    with ZipFile(rt) as z:
        again=named(E.fromstring(z.read('ppt/slides/slide7.xml')))
        for n,e in native.items():
            assert text(again[n])==text(e)
            assert semantic(again[n].find('p:spPr/a:ln',NS))==semantic(e.find('p:spPr/a:ln',NS))
    result=dict(status='pass',date='2026-10-17',head=snap['head'],input_sha256=sha(src),output_sha256=sha(dst),pages=7,
        preserved_pages=[1,2,3,4,5,6],changed_existing_package_parts=changed,new_package_parts=sorted(added),
        old_page6_objects_unchanged=untouched,modified_old_objects=modified,all_inherited_styles_and_fonts_preserved=True,
        configuration_mapping=expected,configuration_modules=expected_module,layout=layout,paths=len(model['edges']),
        native_save_reopen='pass',package_relationships='pass',protected_file_count=len(protected),
        slides=[{**{k:v for k,v in s.items() if k!='texts'},'text_objects':len(s['texts'])} for s in stats],
        preview_source=str(dst.relative_to(ROOT)),office_tested=False,hardware_tested=False)
    report.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:result[k] for k in ['status','output_sha256','old_page6_objects_unchanged','layout','slides']},ensure_ascii=False,indent=2))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--build',type=Path,required=True);p.add_argument('--preview',type=Path,required=True)
    p.add_argument('--report',type=Path,required=True);p.add_argument('--deck',type=Path);a=p.parse_args()
    audit(a.build.resolve(),a.preview.resolve(),a.report.resolve(),a.deck.resolve() if a.deck else None)
