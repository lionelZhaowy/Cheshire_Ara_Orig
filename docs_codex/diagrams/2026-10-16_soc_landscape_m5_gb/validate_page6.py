#!/usr/bin/env python3
"""Audit native-only slide-6 changes against the backed-up slide 5."""
import argparse,hashlib,json,posixpath
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
from validate import NS,objects,points,summary
from inspect_page6 import inspect
from build_page6 import RESTORE,GRAY,named

D=Path(__file__).resolve().parent;ROOT=D.parents[2]
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def semantic(e):return None if e is None else (e.tag,tuple(sorted(e.attrib.items())),e.text,tuple(semantic(c) for c in e))
def close(a,b):return all(abs(x-y)<.01 for x,y in zip(a,b))
def main():
    p=argparse.ArgumentParser();p.add_argument('--build',type=Path,required=True);p.add_argument('--preview',type=Path,required=True);p.add_argument('--report',type=Path,required=True);a=p.parse_args();assert not a.report.exists()
    snap=json.loads((D/'evidence/page6-input.json').read_text());src=D/'evidence/page6-before/soc_landscape_m5_gb.pptx';dst=a.build/'soc_landscape_m5_gb.pptx';assert sha(src)==snap['input_sha256']
    protected={n:sha(ROOT/n)==v for n,v in snap['protected'].items()};assert all(protected.values());assert all(sha(ROOT/n)==v for n,v in snap['backups'].items())
    with ZipFile(src) as zi,ZipFile(dst) as zo:
        assert zo.testzip() is None
        changed=[n for n in zi.namelist() if zi.read(n)!=zo.read(n)]
        assert set(changed)=={'[Content_Types].xml','ppt/presentation.xml','ppt/_rels/presentation.xml.rels','docProps/app.xml'}
        added=set(zo.namelist())-set(zi.namelist());assert added=={'ppt/slides/slide6.xml','ppt/slides/_rels/slide6.xml.rels','ppt/notesSlides/notesSlide6.xml','ppt/notesSlides/_rels/notesSlide6.xml.rels'}
        for n in zo.namelist():
            if not n.endswith('.rels'):continue
            rels=E.fromstring(zo.read(n));ids=[r.get('Id') for r in rels];assert len(ids)==len(set(ids))
            base=posixpath.dirname(posixpath.dirname(n)) if n!='_rels/.rels' else ''
            for r in rels:
                if r.get('TargetMode')!='External':assert posixpath.normpath(posixpath.join(base,r.get('Target'))).lstrip('/') in zo.namelist()
        old=E.fromstring(zi.read('ppt/slides/slide5.xml'));slide=E.fromstring(zo.read('ppt/slides/slide6.xml'));old_objs=objects(old);objs=objects(slide);native=named(slide)
        assert set(old_objs).issubset(objs);assert len(objs)==len(old_objs)+2
        ids=slide.xpath('//p:cNvPr/@id',namespaces=NS);assert len(ids)==len(set(ids))
        names={i:e.find('.//p:cNvPr',NS).get('name').split(': ')[0] for i,e in objs.items()}
        modified=[];untouched=0
        for id,before in old_objs.items():
            name=names[id];after=objs[id]
            if semantic(before)==semantic(after):untouched+=1;continue
            modified.append(name)
            assert name in RESTORE or name.startswith('legend-'),name
            if name in RESTORE:
                cleaned=deepcopy(after);ln=cleaned.find('p:spPr/a:ln',NS)
                dash=ln.find('a:custDash',NS);assert dash is not None;ln.remove(dash)
                if name in GRAY:
                    color=ln.find('a:solidFill/a:srgbClr',NS);assert color.get('val')=='93939D';color.set('val','314D66')
                assert semantic(cleaned)==semantic(before),name
            else:
                assert before.xpath('.//a:rPr/@sz',namespaces=NS)==after.xpath('.//a:rPr/@sz',namespaces=NS)
                assert semantic(before.find('p:spPr/a:ln',NS))==semantic(after.find('p:spPr/a:ln',NS)),name
        assert set(RESTORE).issubset(modified)
        # Every slide-level property outside its shape tree is also untouched.
        stripped=[]
        for root in [old,slide]:
            root=deepcopy(root);tree=root.find('p:cSld/p:spTree',NS)
            for el in list(tree):
                if E.QName(el).localname in ['sp','cxnSp']:tree.remove(el)
            stripped.append(semantic(root))
        assert stripped[0]==stripped[1]
        refs={'page3':named(E.fromstring(zi.read('ppt/slides/slide3.xml')))}
        with ZipFile(D/'evidence/page4-controls-before/soc_landscape_m5_gb.pptx') as ref:refs['page4']=named(E.fromstring(ref.read('ppt/slides/slide4.xml')))
        for name,(page,ref) in RESTORE.items():assert semantic(native[name].find('p:spPr/a:ln/a:custDash',NS))==semantic(refs[page][ref].find('p:spPr/a:ln/a:custDash',NS))
        notes=E.fromstring(zo.read('ppt/notesSlides/notesSlide6.xml'));assert notes.xpath('//a:fld[@type="slidenum"]/a:t/text()',namespaces=NS)==['6']
    model=json.loads((a.build/'slide-06.json').read_text());base=json.loads((D/'evidence/page5-build-02/slide-05.json').read_text())
    for key in ['boxes','regions','dots','bridges']:assert model[key]==base[key],key
    old_edges={e['id']:e for e in base['edges']};edges={e['id']:e for e in model['edges']};definitions={t['id']:t for t in model['texts']}
    assert set(edges)-set(old_edges)=={'legend-dashed'}
    for name,olde in old_edges.items():
        if name.startswith('legend-'):continue
        e=deepcopy(edges[name])
        if name in RESTORE:e['dash']=olde['dash']
        if name in GRAY:e['kind']=olde['kind']
        assert e==olde,name
    for e in model['edges']:
        el=native[e['id']];pts=points(el);assert close(pts[0],e['points'][0]) and close(pts[-1],e['points'][-1]),e['id']
        paths=el.findall('p:spPr/a:custGeom/a:pathLst/a:path',NS);assert len(paths)==1 and len(paths[0].findall('a:moveTo',NS))==1
        for tag,wanted in [('tailEnd',e['arrow']),('headEnd',e['both'])]:
            end=el.find('p:spPr/a:ln/a:'+tag,NS);assert (end is not None and end.get('type')=='triangle')==wanted,e['id']
        assert (el.find('p:spPr/a:ln/a:custDash',NS) is not None)==e['dash'],e['id']
    layout=inspect(a.build,a.preview);assert not any(layout.values()),layout
    preview=json.loads((a.preview/'preview_layout.json').read_text())['slides'][5];boxes={b['id']:b for b in model['boxes']+model['regions']};badge_issues=[]
    for t in preview['texts']:
        name=names[int(t['shape_id'])];x,y,w,h=t['box']
        for b in [v for v in boxes.values() if v['id'].startswith(('badge-','reg-badge-'))]:
            if definitions.get(name,{}).get('owner')==b['id']:continue
            if min(x+w,b['x']+b['w'])>max(x,b['x'])+.1 and min(y+h,b['y']+b['h'])>max(y,b['y'])+.1:badge_issues.append([name,b['id']])
    assert not badge_issues,badge_issues;layout['text_through_badge_frame']=badge_issues
    prs=Presentation(dst);assert len(prs.slides)==6;stats=summary(prs);rt=a.report.with_suffix('.roundtrip.pptx');assert not rt.exists();prs.save(rt);assert summary(Presentation(rt))==stats
    report={'status':'pass','date':'2026-10-17','input_sha256':sha(src),'output_sha256':sha(dst),'pages':6,'unchanged_pages':[1,2,3,4,5],'changed_existing_parts':changed,'new_parts':sorted(added),'theme_master_layout_relationships_unchanged':True,'slide6_unchanged_existing_objects':untouched,'slide6_modified_objects':modified,'restored_lines':RESTORE,'native_geometry_widths_arrow_directions_unchanged':True,'configuration_title_topology_unchanged':True,'layout':layout,'native_paths':len(model['edges']),'explicit_bridges':len(model['bridges']),'native_save_reopen':'pass','package_relationships':'pass','slides':[{**{k:v for k,v in s.items() if k!='texts'},'text_objects':len(s['texts'])} for s in stats],'protected_files':protected,'office_tested':False,'hardware_tested':False}
    a.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps({k:report[k] for k in ['status','output_sha256','pages','slide6_unchanged_existing_objects','slide6_modified_objects','layout','slides']},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
