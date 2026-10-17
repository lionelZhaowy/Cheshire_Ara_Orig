#!/usr/bin/env python3
"""Validate page-4 artifacts and their preservation of the original three pages."""
import argparse,hashlib,json,posixpath
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
from inspect_page4 import inspect
from validate import NS,objects,points,summary

D=Path(__file__).resolve().parent;ROOT=D.parents[2]
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()

def close(a,b):return all(abs(x-y)<.01 for x,y in zip(a,b))

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--build',type=Path,required=True);ap.add_argument('--preview',type=Path,required=True);ap.add_argument('--report',type=Path,required=True);a=ap.parse_args();assert not a.report.exists()
    inp=json.loads((D/'evidence/page4-input.json').read_text());source=D/'evidence/page4-before/soc_landscape_m5_gb.pptx';dest=a.build/'soc_landscape_m5_gb.pptx'
    assert sha(source)==inp['input_sha256']
    protected={p:sha(ROOT/p)==v for p,v in inp['protected'].items()};assert all(protected.values())
    assert all(sha(ROOT/p)==v for p,v in inp['backups'].items())
    with ZipFile(source) as zi,ZipFile(dest) as zo:
        assert zo.testzip() is None
        changed=[n for n in zi.namelist() if zi.read(n)!=zo.read(n)]
        assert set(changed)=={'[Content_Types].xml','ppt/presentation.xml','ppt/_rels/presentation.xml.rels','docProps/app.xml'},changed
        assert set(zo.namelist())-set(zi.namelist())=={'ppt/slides/slide4.xml','ppt/slides/_rels/slide4.xml.rels','ppt/notesSlides/notesSlide4.xml','ppt/notesSlides/_rels/notesSlide4.xml.rels'}
        assert not any(n.startswith('ppt/media/') for n in zo.namelist())
        for n in zo.namelist():
            if not n.endswith('.rels'):continue
            r=E.fromstring(zo.read(n));ids=[v.get('Id') for v in r];assert len(ids)==len(set(ids))
            base=posixpath.dirname(posixpath.dirname(n)) if n!='_rels/.rels' else ''
            for rel in r:
                if rel.get('TargetMode')=='External':continue
                target=posixpath.normpath(posixpath.join(base,rel.get('Target'))).lstrip('/')
                assert target in zo.namelist(),(n,target)
        slide=E.fromstring(zo.read('ppt/slides/slide4.xml'));obj=objects(slide)
        ids=slide.xpath('//p:cNvPr/@id',namespaces=NS);assert len(ids)==len(set(ids))
        native={el.find('.//p:cNvPr',NS).get('name').split(': ')[0]:el for el in obj.values()}
        assert slide.find('p:timing',NS) is None,'Copied timing must not retain stale shape references'
        notes=E.fromstring(zo.read('ppt/notesSlides/notesSlide4.xml'))
        assert notes.xpath('//a:fld[@type="slidenum"]/a:t/text()',namespaces=NS)==['4']
    model=json.loads((a.build/'slide-04.json').read_text());boxes={b['id']:b for b in model['boxes']+model['regions']};edges={e['id']:e for e in model['edges']}
    # Preserve the existing functional objects after the authorized reflow.
    retained=['cva6','ara','ara-filter','ara-width','npu','gb-adapter','global-buffer','gb-mover','lvds-rx','rx-cdc','isp','frame-select','frame-dma','pll','axi-crossbar','debug-sba','idma-master','slink-rx','vga','usb-dma','ext-master','debug-target','reg-bridge','memory-target','atomics','llc-spm','idma-config','slink-tx','gb-interface','ext-target','regbus','ddr-controller','ddr-arb','ddr-rw','ddr-command','ddr-phy','ddr-device']
    assert all(k in boxes for k in retained)
    assert len([b for b in model['boxes'] if b['id']=='global-buffer'])==1
    assert all('reg-'+str(row)+'-'+str(col) in boxes for row in range(6) for col in range(5) if (row,col)!=(5,4))
    text='\n'.join(t['text'] for t in model['texts'])
    for s in ['8 × 6 固定端口 + 扩展','2 lanes · VLSU / AXI：64 bit','CPU / Ara / eMMC DMA','≤256 ROI / 帧 · 不缩放','四口共享压缩器','四口共享解压器','普通数据旁路','DDR Controller · 4 个 AXI slave','独立 IRQ 向量']:assert s in text,s
    assert 'M_GB' not in text and '7 × 6' not in text
    xc=[];mc=[]
    for i in range(9):
        b=boxes['xbar-in-'+('S_EXT' if i==8 else 'S'+str(i))];xc.append(b['x']+b['w']/2);assert b['y']==966
    for i in range(7):
        b=boxes['xbar-out-'+('M_EXT' if i==6 else 'M'+str(i))];mc.append(b['x']+b['w']/2);assert b['y']==1056
    for centers in [xc,mc]:
        gaps=[b-a for a,b in zip(centers,centers[1:])];assert max(gaps)-min(gaps)<.001
    expected={
      'emmc-physical':('emmc-device','emmc-host'),'master-7':('emmc-host','xbar-in-S7'),
      'target-5':('xbar-out-M5','gb-interface'),'m5-gb':('gb-interface','gb-adapter'),
      'adapter-sram':('gb-adapter','global-buffer'),'isp-gb':('isp','gb-adapter'),
      'roi-vdma-control':('roi-scheduler','roi-vdma'),'vdma-pixels':('roi-vdma','roi-fifo'),
      'fifo-lvds':('roi-fifo','lvds-tx'),'lvds-external':('lvds-tx','lvds-peer'),
      'llc-f0':('llc-spm','front-F0'),'mover-f1':('gb-mover','front-F1'),
      'frame-f2':('frame-dma','front-F2'),'roi-f3-request':('roi-vdma','front-F3'),
      'f3-roi-read-data':('front-F3','roi-vdma'),'regext-new-config':('reg-3-4','roi-region'),
      'config-emmc':('new-config-trunk','emmc-host'),'config-codec':('new-config-trunk','compression-front'),
      'roi-irq':('roi-region','new-irq-sync'),'emmc-irq':('emmc-host','new-irq-trunk'),
      'codec-irq':('compression-front','new-irq-trunk'),'new-irq-plic':('new-irq-sync','reg-0-3')}
    for id,pair in expected.items():assert (edges[id]['src'],edges[id]['dst'])==pair
    for i in range(4):
        assert edges['front-input-'+str(i)]['dst']=='front-adapt'
        assert edges['front-ddr-'+str(i)]['src']=='front-route'
        assert edges['front-ddr-'+str(i)]['dst']=='ddr-S'+str(i)
    assert edges['adapt-compressor']['src']=='front-adapt' and edges['compressor-route']['dst']=='front-route'
    assert edges['read-to-decoder']['src']=='front-route' and edges['decoder-up']['dst']=='front-adapt'
    for id in ['adapt-plain-bypass','plain-bypass-route']:assert edges[id]['both']
    # Continuous native paths and exact source/target geometry; no dangling tips.
    for e in model['edges']:
        actual=points(native[e['id']]);assert close(actual[0],e['points'][0]) and close(actual[-1],e['points'][-1]),e['id']
        paths=native[e['id']].findall('p:spPr/a:custGeom/a:pathLst/a:path',NS)
        assert len(paths)==1 and len(paths[0].findall('a:moveTo',NS))==1
        ln=native[e['id']].find('p:spPr/a:ln',NS)
        for attr,required in [('tailEnd',e['arrow']),('headEnd',e['both'])]:
            end=ln.find('a:'+attr,NS);assert (end is not None and end.get('type')=='triangle')==required
        for end,pt in [('src',e['points'][0]),('dst',e['points'][-1])]:
            if e[end] not in boxes:continue
            b=boxes[e[end]];x,y,w,h=[b[k] for k in ['x','y','w','h']];px,py=pt
            valid=((abs(px-x)<.01 or abs(px-x-w)<.01) and y-.01<=py<=y+h+.01 or (abs(py-y)<.01 or abs(py-y-h)<.01) and x-.01<=px<=x+w+.01)
            assert valid,(e['id'],end,pt)
    layout=inspect(a.build,a.preview);assert not any(layout.values()),layout
    for i in [1,2,3]:
        for ext in ['svg','png']:assert (a.preview/f'slide-{i:02d}.{ext}').read_bytes()==(D/f'slide-{i:02d}.{ext}').read_bytes(),(i,ext)
    prs=Presentation(dest);assert len(prs.slides)==4;stats=summary(prs)
    rt=a.report.with_suffix('.roundtrip.pptx');assert not rt.exists();prs.save(rt);assert summary(Presentation(rt))==stats
    assert stats[3]['texts']==[t['text'] for t in model['texts']]
    report={'status':'pass','date':'2026-10-16','source_sha256':sha(source),'output_sha256':sha(dest),
      'pages':4,'pages1_to3_and_dependencies_byte_identical':True,'pages1_to3_svg_png_byte_identical':True,
      'changed_existing_package_parts':changed,'native_save_reopen':'pass','relationships':'pass',
      'retained_functional_modules':retained,'retained_register_and_configuration_targets':29,
      'top_ports':['S'+str(i) for i in range(8)]+['S_EXT'],'bottom_ports':['M'+str(i) for i in range(6)]+['M_EXT'],
      'both_port_rows_uniform':True,'critical_connections':expected,'four_independent_ddr_paths':'F0→S0 / F1→S1 / F2→S2 / F3→S3',
      'shared_encoder_decoder_bypass':'pass','all_native_paths_continuous':'pass','all_module_endpoints':'pass','layout':layout,
      'explicit_crossing_bridges':len(model['bridges']),'slides':[{**{k:v for k,v in s.items() if k!='texts'},'text_objects':len(s['texts'])} for s in stats],
      'protected_files':protected,'office_application_tested':False,'hardware_tested':False,
      'preview_type':'explicit DrawingML geometry and local font substitution, not Office screenshots'}
    a.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k not in ['protected_files','retained_functional_modules','critical_connections']},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
