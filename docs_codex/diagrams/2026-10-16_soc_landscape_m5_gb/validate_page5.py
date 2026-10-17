#!/usr/bin/env python3
"""Preservation, native geometry, semantic paths and configuration alignment audit."""
import argparse,hashlib,json,posixpath
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
from validate import NS,objects,points,summary
from inspect_page5 import inspect
from build_page5 import DY,FULL_AXI

D=Path(__file__).resolve().parent;ROOT=D.parents[2]
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def semantic(e):return None if e is None else (e.tag,tuple(sorted(e.attrib.items())),e.text,tuple(semantic(c) for c in e))
def close(a,b):return all(abs(x-y)<.01 for x,y in zip(a,b))

def main():
    p=argparse.ArgumentParser();p.add_argument('--build',type=Path,required=True);p.add_argument('--preview',type=Path,required=True);p.add_argument('--report',type=Path,required=True);a=p.parse_args();assert not a.report.exists()
    snap=json.loads((D/'evidence/page5-input.json').read_text());src=D/'evidence/page5-before/soc_landscape_m5_gb.pptx';dst=a.build/'soc_landscape_m5_gb.pptx'
    assert sha(src)==snap['input_sha256']
    protected={n:sha(ROOT/n)==v for n,v in snap['protected'].items()};assert all(protected.values())
    assert all(sha(ROOT/n)==v for n,v in snap['backups'].items())
    with ZipFile(src) as zi,ZipFile(dst) as zo:
        assert zo.testzip() is None
        changed=[n for n in zi.namelist() if zi.read(n)!=zo.read(n)]
        assert set(changed)=={'[Content_Types].xml','ppt/presentation.xml','ppt/_rels/presentation.xml.rels','docProps/app.xml'},changed
        added=set(zo.namelist())-set(zi.namelist())
        assert added=={'ppt/slides/slide5.xml','ppt/slides/_rels/slide5.xml.rels','ppt/notesSlides/notesSlide5.xml','ppt/notesSlides/_rels/notesSlide5.xml.rels'}
        for n in zo.namelist():
            if not n.endswith('.rels'):continue
            rels=E.fromstring(zo.read(n));ids=[r.get('Id') for r in rels];assert len(ids)==len(set(ids))
            base=posixpath.dirname(posixpath.dirname(n)) if n!='_rels/.rels' else ''
            for r in rels:
                if r.get('TargetMode')=='External':continue
                assert posixpath.normpath(posixpath.join(base,r.get('Target'))).lstrip('/') in zo.namelist()
        old=objects(E.fromstring(zi.read('ppt/slides/slide4.xml')));third=objects(E.fromstring(zi.read('ppt/slides/slide3.xml')))
        slide=E.fromstring(zo.read('ppt/slides/slide5.xml'));obj=objects(slide)
        assert all(semantic(obj[i])==semantic(third[i]) for i in [7,9])
        ids=slide.xpath('//p:cNvPr/@id',namespaces=NS);assert len(ids)==len(set(ids))
        assert slide.find('p:timing',NS) is None
        assert not zo.namelist().count('ppt/media/image1.png')
        notes=E.fromstring(zo.read('ppt/notesSlides/notesSlide5.xml'));assert notes.xpath('//a:fld[@type="slidenum"]/a:t/text()',namespaces=NS)==['5']
    named=lambda objs:{e.find('.//p:cNvPr',NS).get('name').split(': ')[0]:e for e in objs.values()}
    native=named(obj);old_native=named(old);names={i:e.find('.//p:cNvPr',NS).get('name').split(': ')[0] for i,e in obj.items()}
    model=json.loads((a.build/'slide-05.json').read_text());base=json.loads((D/'evidence/page4-controls-build-04/slide-04.json').read_text())
    boxes={b['id']:b for b in model['boxes']+model['regions']};edges={e['id']:e for e in model['edges']};definitions={t['id']:t for t in model['texts']}
    original_boxes={b['id']:b for b in base['boxes']+base['regions']}
    assert set(boxes)==set(original_boxes)
    for b in boxes.values():assert 0<=b['x'] and 0<=b['y'] and b['x']+b['w']<=3600 and b['y']+b['h']<=2025
    for name in set(native)&set(old_native):
        assert semantic(native[name].find('p:spPr/a:solidFill',NS))==semantic(old_native[name].find('p:spPr/a:solidFill',NS)),name
        assert native[name].xpath('.//a:rPr/@sz',namespaces=NS)==old_native[name].xpath('.//a:rPr/@sz',namespaces=NS),name
    # Translation preserves all main module positions except explicitly reflowed groups.
    for name in ['cva6','ara','ara-filter','ara-width','global-buffer','gb-adapter','frame-select','axi-crossbar','regbus','pll','roi-scheduler','roi-vdma','roi-fifo','lvds-tx','ddr-controller','ddr-phy']:
        b=boxes[name];oldb=original_boxes[name];assert close([b['x'],b['y'],b['w'],b['h']],[oldb['x'],oldb['y']+DY,oldb['w'],oldb['h']]),name
    for row in range(2):
        items=[boxes[f'config-C{i}'] for i in range(1+8*row,9+8*row)]
        assert len({(b['y'],b['h'],b['w']) for b in items})==1
        assert items[0]['x']==boxes['regbus']['x']
        assert abs(items[-1]['x']+items[-1]['w']-boxes['regbus']['x']-boxes['regbus']['w'])<.01
        assert all(u['x']+u['w']<v['x'] for u,v in zip(items,items[1:]))
        for i in range(1+8*row,9+8*row):assert edges[f'cfg-C{i}']['src']==f'config-row-{row+1}'
    assert boxes['config-C1']['y']>boxes['reg-region']['y']+boxes['reg-region']['h']
    assert abs(boxes['config-C9']['y']+boxes['config-C9']['h']-boxes['pll']['y']-boxes['pll']['h'])<.01
    assert len([b for b in model['boxes'] if b['id']=='global-buffer'])==1
    for olde in base['edges']:
        if olde['id'].startswith(('legend-','cfg-')) or olde['id'] in ['regext-config','config-trunk']:continue
        e=edges[olde['id']];assert (e['src'],e['dst'],e['kind'],e['arrow'])==(olde['src'],olde['dst'],olde['kind'],olde['arrow']),e['id']
        assert e['both']==(True if e['id'] in FULL_AXI else olde['both']),e['id']
    assert all(edges[n]['both'] and edges[n]['arrow'] for n in FULL_AXI)
    assert not any(e['kind']=='irq' for e in model['edges'])
    assert 'f3-roi-read-data' not in edges
    # Validate endpoints, a single continuous native path, and both native arrowheads.
    for e in model['edges']:
        el=native[e['id']];pts=points(el);assert close(pts[0],e['points'][0]) and close(pts[-1],e['points'][-1])
        paths=el.findall('p:spPr/a:custGeom/a:pathLst/a:path',NS);assert len(paths)==1 and len(paths[0].findall('a:moveTo',NS))==1
        for tag,wanted in [('tailEnd',e['arrow']),('headEnd',e['both'])]:
            end=el.find('p:spPr/a:ln/a:'+tag,NS);assert (end is not None and end.get('type')=='triangle')==wanted
        for role,pt in [('src',e['points'][0]),('dst',e['points'][-1])]:
            if e[role] not in boxes:continue
            b=boxes[e[role]];x,y,w,h=[b[k] for k in ['x','y','w','h']];px,py=pt
            assert ((abs(px-x)<.01 or abs(px-x-w)<.01) and y-.01<=py<=y+h+.01 or (abs(py-y)<.01 or abs(py-y-h)<.01) and x-.01<=px<=x+w+.01),(e['id'],role)
    layout=inspect(a.build,a.preview);assert not any(layout.values()),layout
    preview=json.loads((a.preview/'preview_layout.json').read_text())['slides'][4]
    badge_issues=[]
    for t in preview['texts']:
        name=names[int(t['shape_id'])];x,y,w,h=t['box']
        for b in [v for v in boxes.values() if v['id'].startswith(('badge-','reg-badge-'))]:
            if definitions.get(name,{}).get('owner')==b['id']:continue
            if min(x+w,b['x']+b['w'])>max(x,b['x'])+.1 and min(y+h,b['y']+b['h'])>max(y,b['y'])+.1:badge_issues.append([name,b['id']])
    assert not badge_issues,badge_issues;layout['text_through_badge_frame']=badge_issues
    prs=Presentation(dst);assert len(prs.slides)==5;stats=summary(prs)
    rt=a.report.with_suffix('.roundtrip.pptx');assert not rt.exists();prs.save(rt);assert summary(Presentation(rt))==stats
    report={'status':'pass','date':'2026-10-17','input_sha256':sha(src),'output_sha256':sha(dst),'pages':5,'unchanged_pages':[1,2,3,4],'changed_existing_parts':changed,'new_parts':sorted(added),'existing_theme_master_layout_relationships_unchanged':True,'title_shape_ids_copied_exactly':[7,9],'body_shift':DY,'matched_body_fonts_and_fills':'unchanged','configuration_rows':[[f'C{i}' for i in range(1,9)],[f'C{i}' for i in range(9,17)]],'second_row_and_pll_bottom_aligned':True,'full_bidirectional_axi_interfaces':FULL_AXI,'other_functional_connections_unchanged':True,'layout':layout,'native_paths':len(model['edges']),'explicit_bridges':len(model['bridges']),'native_save_reopen':'pass','package_relationships':'pass','slides':[{**{k:v for k,v in s.items() if k!='texts'},'text_objects':len(s['texts'])} for s in stats],'protected_files':protected,'office_tested':False,'hardware_tested':False}
    a.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:report[k] for k in ['status','output_sha256','pages','unchanged_pages','layout','slides']},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
