#!/usr/bin/env python3
"""Correct only slide 6's F3 interface dash; preserve every other native property."""
import argparse,hashlib,json
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
from pptx import Presentation
from build_page6 import NS,named

D=Path(__file__).resolve().parent
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    p=argparse.ArgumentParser();p.add_argument('--input',type=Path,required=True);p.add_argument('--out-dir',type=Path,required=True);a=p.parse_args();a.out_dir.mkdir(parents=True,exist_ok=False)
    with ZipFile(a.input) as z:parts={n:z.read(n) for n in z.namelist()}
    slide=E.fromstring(parts['ppt/slides/slide6.xml']);before=E.tostring(slide);obj=named(slide)
    ln=obj['roi-f3-request'].find('p:spPr/a:ln',NS)
    assert ln.find('a:custDash',NS) is None and ln.find('a:prstDash',NS) is None
    source=obj['mover-f1'].find('p:spPr/a:ln/a:custDash',NS);assert source is not None
    assert E.tostring(source)==E.tostring(obj['frame-f2'].find('p:spPr/a:ln/a:custDash',NS))
    assert ln.find('a:solidFill/a:srgbClr',NS).get('val')=='314D66'
    assert all(ln.find('a:'+tag,NS).get('type')=='triangle' for tag in ['headEnd','tailEnd'])
    dash=deepcopy(source);ln.insert(1,dash)
    # Removing the single added dash node must reproduce the entire slide exactly.
    check=deepcopy(slide);target=named(check)['roi-f3-request'].find('p:spPr/a:ln',NS);target.remove(target.find('a:custDash',NS));assert E.tostring(check)==before
    parts['ppt/slides/slide6.xml']=E.tostring(slide,encoding='UTF-8',xml_declaration=True,standalone=True)
    dst=a.out_dir/'soc_landscape_m5_gb.pptx'
    with ZipFile(dst,'w',ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    with ZipFile(a.input) as old,ZipFile(dst) as new:
        assert old.namelist()==new.namelist();assert new.testzip() is None
        changed=[n for n in old.namelist() if old.read(n)!=new.read(n)];assert changed==['ppt/slides/slide6.xml']
    prs=Presentation(dst);assert len(prs.slides)==6
    assert all(s.shape_type!=13 for sl in prs.slides for s in sl.shapes)
    rt=a.out_dir/'save-reopen.pptx';prs.save(rt)
    with ZipFile(rt) as z:
        restored=E.fromstring(z.read('ppt/slides/slide6.xml'))
        assert E.tostring(named(restored)['roi-f3-request'].find('p:spPr/a:ln',NS))==E.tostring(ln)
    assert len(Presentation(rt).slides)==6
    model=json.loads((D/'evidence/page6-build-02/slide-06.json').read_text())
    edge=next(e for e in model['edges'] if e['id']=='roi-f3-request');assert not edge['dash'];edge['dash']=True
    (a.out_dir/'slide-06.json').write_text(json.dumps(model,ensure_ascii=False,indent=2)+'\n')
    report={'status':'pass','input_sha256':sha(a.input),'output_sha256':sha(dst),'pages':6,'changed_package_parts':changed,'changed_shape':'roi-f3-request','only_added_property':'a:ln/a:custDash copied from F1, matching F2','color':'314D66','arrowheads':['triangle','triangle'],'all_other_shape_properties_unchanged':True,'pages_1_to_5_and_dependencies_unchanged':True,'native_save_reopen':'pass','office_tested':False}
    (a.out_dir/'native-check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
