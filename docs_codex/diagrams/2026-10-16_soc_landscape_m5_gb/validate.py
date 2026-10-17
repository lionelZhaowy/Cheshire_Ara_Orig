#!/usr/bin/env python3
"""Check the requested M5 layout correction, native deck and preserved input."""
from collections import Counter
import hashlib
import json
from pathlib import Path
import posixpath
import re
import subprocess
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
from build import NS,SCALE,REMOVED,CHANGED,CAPTION

D=Path(__file__).resolve().parent
ROOT=D.parents[2]
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()


def objects(root):
    return {int(el.find('.//p:cNvPr',NS).get('id')):el for el in root.find('p:cSld/p:spTree',NS)
            if E.QName(el).localname in ['sp','cxnSp']}


def box(el):
    xf=el.find('p:spPr/a:xfrm',NS);off=xf.find('a:off',NS);ext=xf.find('a:ext',NS)
    return [int(v)/SCALE for v in [off.get('x'),off.get('y'),ext.get('cx'),ext.get('cy')]]


def text(el):return ''.join(el.xpath('.//a:t/text()',namespaces=NS))


def points(el):
    """Flatten the actual native path for endpoint and obstacle checks."""
    x,y,w,h=box(el);xf=el.find('p:spPr/a:xfrm',NS)
    path=el.find('p:spPr/a:custGeom/a:pathLst/a:path',NS)
    if path is None:
        assert el.find('p:spPr/a:prstGeom',NS).get('prst')=='line'
        return [(x+w if xf.get('flipH')=='1' else x,y+h if xf.get('flipV')=='1' else y),
                (x if xf.get('flipH')=='1' else x+w,y if xf.get('flipV')=='1' else y+h)]
    result=[]
    for command in path:
        p=[(x+int(q.get('x'))/int(path.get('w'))*w,y+int(q.get('y'))/int(path.get('h'))*h) for q in command]
        kind=E.QName(command).localname
        if kind in ['moveTo','lnTo']:result+=p
        elif kind=='cubicBezTo':
            p0=result[-1];p1,p2,p3=p
            for i in range(1,41):
                t=i/40;s=1-t
                result.append(tuple(s**3*p0[k]+3*s*s*t*p1[k]+3*s*t*t*p2[k]+t**3*p3[k] for k in [0,1]))
        else:raise ValueError(kind)
    return result


def intersects(a,b,r,pad=0):
    """Open interior intersection of a segment with a rectangle."""
    x,y,w,h=r;low=[x-pad,y-pad];high=[x+w+pad,y+h+pad];t0,t1=0.,1.
    for k in [0,1]:
        delta=b[k]-a[k]
        if abs(delta)<1e-8:
            if a[k]<=low[k] or a[k]>=high[k]:return False
        else:
            aa=(low[k]-a[k])/delta;bb=(high[k]-a[k])/delta
            t0=max(t0,min(aa,bb));t1=min(t1,max(aa,bb))
    return t1>t0+1e-7


def summary(prs):
    result=[]
    for sl in prs.slides:
        result.append({'objects':len(sl.shapes),
                       'texts':[s.text for s in sl.shapes if s.has_text_frame and s.text],
                       'native_arrows':sum(bool(s._element.xpath('.//a:headEnd[@type="triangle"] | .//a:tailEnd[@type="triangle"]')) for s in sl.shapes),
                       'native_dashes':sum(bool(s._element.xpath('.//a:custDash')) for s in sl.shapes)})
    return result


def main():
    dest=D/'evidence/validation.json';assert not dest.exists()
    inputs=json.loads((D/'evidence/input.json').read_text())
    protected={name:sha(ROOT/name)==digest for name,digest in inputs['protected'].items()}
    assert all(protected.values()),[k for k,v in protected.items() if not v]
    assert sha(D/'evidence/before/L01_global_buffer_cpu_access.md')==inputs['handoff_before']
    input_path=ROOT/inputs['source'];output=D/'soc_landscape_m5_gb.pptx'
    edit=json.loads((D/'evidence/build-01/edit_report.json').read_text())
    assert sha(output)==edit['output_sha256']
    with ZipFile(input_path) as src,ZipFile(output) as out:
        assert out.testzip() is None
        assert set(src.namelist())==set(out.namelist())
        assert [n for n in src.namelist() if src.read(n)!=out.read(n)]==['ppt/slides/slide2.xml']
        assert not any(n.startswith('ppt/media/') for n in out.namelist())
        for n in [n for n in out.namelist() if n.endswith('.rels')]:
            rels=E.fromstring(out.read(n));ids=[r.get('Id') for r in rels];assert len(ids)==len(set(ids))
            base=posixpath.dirname(posixpath.dirname(n)) if n!='_rels/.rels' else ''
            for rel in rels:
                if rel.get('TargetMode')=='External':continue
                target=posixpath.normpath(posixpath.join(base,rel.get('Target'))).lstrip('/')
                assert target in out.namelist(),(n,target)
        before=objects(E.fromstring(src.read('ppt/slides/slide2.xml')))
        root=E.fromstring(out.read('ppt/slides/slide2.xml'));after=objects(root)
        ids=root.xpath('//p:cNvPr/@id',namespaces=NS);assert len(ids)==len(set(ids))
        assert set(before)-set(after)==REMOVED
        assert set(after)-set(before)==set(range(1300,1308))
        altered={id for id in before.keys()&after.keys() if E.tostring(before[id])!=E.tostring(after[id])}
        assert altered==CHANGED,(altered,CHANGED)
        all_text='\n'.join(text(el) for el in after.values())
        assert 'M_GB' not in all_text and not re.search(r'7\s*[×x]\s*5',all_text)
        assert text(after[886])==CAPTION
        names=[el.find('.//p:cNvPr',NS).get('name') for el in after.values()]
        assert names.count('global-buffer')==1 and 'ara-split' not in names and 'ara-gb' not in names
        for id in [880,881,889,890,894,895,1200]:assert E.tostring(before[id])==E.tostring(after[id])
        # All upper ports are initiator-facing S ports. All target ports sit on
        # the bottom edge, with seven ordered targets including M_EXT.
        ports={};xbx,xby,xbw,xbh=box(after[884])
        for el in after.values():
            t=text(el)
            if re.fullmatch(r'M[0-5]|M_EXT|S[0-6]|S_EXT',t):
                x,y,w,h=box(el)
                # The DDR Controller has its own S0/S1/S2 labels. Select only
                # the main Crossbar's port row rather than merging namespaces.
                if xbx<=x+w/2<=xbx+xbw and xby-20<y<xby+xbh+20:
                    ports[t]=(x+w/2,y+h/2)
        targets=['M0','M1','M2','M3','M4','M5','M_EXT']
        assert all(t in ports for t in targets)
        assert sorted(targets,key=lambda n:ports[n][0])==targets
        assert all(abs(ports[t][1]-1028)<8 for t in targets)
        assert all(abs(ports[t][1]-893)<8 for t in ['S0','S1','S2','S3','S4','S5','S6','S_EXT'])
        # Arrow tips terminate at the GB front-end and the existing controller.
        drop=points(after[1300]);route=points(after[1301])
        for actual,expected in [(drop[0],(2436.25,1042.1)),(drop[-1],(2436.25,1093.1)),
                                (route[0],(2566.25,1170)),(route[-1],(1506.25,638.1))]:
            assert all(abs(a-b)<.002 for a,b in zip(actual,expected)),(actual,expected)
        for id in [1300,1301]:
            assert after[id].find('p:spPr/a:ln/a:tailEnd',NS).get('type')=='triangle'
            assert after[id].find('p:spPr/a:ln/a:headEnd',NS) is None
        native_path=after[1301].find('p:spPr/a:custGeom/a:pathLst/a:path',NS)
        assert len(native_path.findall('a:moveTo',NS))==1
        assert len(native_path.findall('a:cubicBezTo',NS))==2
        # Exclude functional region backgrounds; the new routes must avoid all
        # nearby IPs, the Crossbar body and the lower target boxes.
        obstacles=[781,785,790,796,804,809,813,821,831,836,839,844,847,
                   862,865,868,871,874,877,884,915,921,927,932,939,945,1304]
        for line in [drop,route]:
            for a,b in zip(line,line[1:]):
                for id in obstacles:
                    x,y,w,h=box(after[id]);assert not intersects(a,b,(x+.01,y+.01,w-.02,h-.02)),id
    layout=json.loads((D/'evidence/preview-01/preview_layout.json').read_text())['slides'][1]
    new_texts={886,938,940,941,944,946,947,1303,1305,1306,1307}
    for t in layout['texts']:
        if int(t['shape_id']) in new_texts:
            x,y,w,h=t['box'];assert x>=0 and y>=0 and x+w<=3600 and y+h<=2025
            for u in layout['texts']:
                if u is t:continue
                xx,yy,ww,hh=u['box']
                assert not (min(x+w,xx+ww)-max(x,xx)>1 and min(y+h,yy+hh)-max(y,yy)>1),(t,u)
        for line in [drop,route]:
            assert not any(intersects(a,b,t['box']) for a,b in zip(line,line[1:])),t
    # The two compacted blocks and GB interface retain readable text inside.
    owners={938:937,940:939,941:939,944:943,946:945,947:945,1303:1302,1305:1304,1306:1304,1307:1304}
    for t in layout['texts']:
        id=int(t['shape_id'])
        if id not in owners:continue
        x,y,w,h=t['box'];bx,by,bw,bh=box(after[owners[id]])
        assert bx<=x and by<=y and x+w<=bx+bw and y+h<=by+bh,(id,t['box'],box(after[owners[id]]))
    prs=Presentation(output);assert len(prs.slides)==2;stats=summary(prs)
    roundtrip=D/'evidence/roundtrip.pptx';assert not roundtrip.exists()
    prs.save(roundtrip);assert summary(Presentation(roundtrip))==stats
    for i in [1,2]:
        svg=E.parse(str(D/f'slide-{i:02d}.svg'))
        svgtexts=svg.xpath('//s:text/text()',namespaces={'s':'http://www.w3.org/2000/svg'})
        assert ''.join(svgtexts)==''.join(stats[i-1]['texts'])
    old=input_path.parent
    assert (D/'slide-01.svg').read_bytes()==(old/'slide-01.svg').read_bytes()
    assert (D/'slide-01.png').read_bytes()==(old/'slide-01.png').read_bytes()
    diff=subprocess.run(['git','diff','--check'],cwd=ROOT,capture_output=True,text=True)
    assert diff.returncode==0,diff.stdout+diff.stderr
    report={'status':'pass','date':'2026-10-16','only_changed_package_part':'ppt/slides/slide2.xml',
            'slide1_and_dependencies_byte_identical':True,'source_file_unchanged':True,
            'page2_unchanged_shapes':len(before)-len(REMOVED)-len(CHANGED),
            'removed_shape_ids':sorted(REMOVED),'changed_shape_ids':sorted(CHANGED),'new_native_objects':8,
            'target_order':targets,'caption':CAPTION,'initiator_ports_only_on_top':True,'gb_sram_instances':1,
            'native_paths_continuous':True,'intentional_bridge_crossings':edit['route_bridges'],
            'new_path_endpoints':'pass','new_paths_avoid_ip_interiors_and_text':'pass','text_bounds':'pass',
            'slides':[{**{k:v for k,v in s.items() if k!='texts'},'text_objects':len(s['texts'])} for s in stats],
            'save_reopen':'pass','package_relationships':'pass','protected_files':protected,
            'git_diff_check':'pass','office_application_tested':False,'hardware_tested':False,
            'preview_type':'DrawingML geometry with local fallback font, not Office rendering',
            'artifact_sha256':{p.name:sha(p) for p in [output,*D.glob('slide-*.*')]}}
    dest.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k not in ['protected_files','artifact_sha256']},ensure_ascii=False,indent=2))


if __name__=='__main__':main()
