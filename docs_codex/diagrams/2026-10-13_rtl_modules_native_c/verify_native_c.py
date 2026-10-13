"""Check the C presentation against A and current navigation, not RTL behavior."""
from pathlib import Path
from collections import Counter
from urllib.parse import unquote
import xml.etree.ElementTree as ET
import hashlib, json, re, tarfile, zipfile
import pymupdf as fitz
from pptx import Presentation

ROOT=Path(__file__).resolve().parent
BASE=ROOT.parent/'2026-10-13_rtl_module_comparison'
REPO=ROOT.parents[2]
NS={'s':'http://www.w3.org/2000/svg'}
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
errors=[]; links=0; edges=0; pdfs=[]
def check(ok,*why):
    if not ok: errors.append(why)

graphs=json.loads((ROOT/'connections.json').read_text())['figures']
check(json.loads((BASE/'connections.json').read_text())['figures']==graphs,'semantic inventory changed')
manifest=json.loads((ROOT/'render_manifest.json').read_text())
visual=json.loads((ROOT/'visual_review.json').read_text())
check(len(manifest)==len(graphs)==len(visual['images'])==24,'coverage')
for item in manifest:
    name=item['name'];a=BASE/'A'/(name+'.svg');c=ROOT/'C'/(name+'.svg')
    check(sha(a)==item['A_sha256'],'A changed',name)
    check(sha(c)==item['C_sha256'],'C changed after render',name)
    check(sha(c.with_suffix('.png'))==item['png_sha256'],'PNG changed',name)
    ar=ET.parse(a).getroot(); cr=ET.parse(c).getroot()
    # Topology, stroke pattern, arrow start/end and branch dots are exact copies.
    for tag,attrs in [('polyline',['id','points','marker-start','marker-end','stroke-dasharray']),('circle',['cx','cy','r'])]:
        def values(r): return [[e.get(k) for k in attrs] for e in r.findall('s:'+tag,NS)]
        check(values(ar)==values(cr),'wire/branch topology mismatch',name,tag)
    def texts(r):
        return [(e.text.replace('A：连续走线。','C：连续走线。'),e.get('x'),e.get('y'),e.get('font-size'),e.get('text-anchor')) for e in r.findall('s:text',NS)]
    check(texts(ar)==texts(cr),'text or layout mismatch',name)
    def dashed_rects(r):
        return [[e.get(k) for k in ['x','y','width','height','stroke-dasharray']] for e in r.findall('s:rect',NS) if e.get('stroke-dasharray')]
    check(dashed_rects(ar)==dashed_rects(cr),'planned boundary mismatch',name)
    drawn={e.get('id'):e for e in cr.findall('s:polyline',NS) if re.fullmatch(r'E\d+',e.get('id',''))}
    g=next(g for g in graphs if g['name']==name)
    check(set(drawn)=={e['id'] for e in g['edges']},'missing or extra edge',name)
    edges+=len(drawn)
    for e in g['edges']:check(bool(e.get('detail',e.get('label'))),'missing connection explanation',name,e['id'])
    image=next((i for i in visual['images'] if i['name']==name),{})
    check(image.get('checked') and not image.get('issues') and image.get('sha256')==item['png_sha256'],'visual review missing or stale',name)
    with fitz.open(c.with_suffix('.pdf')) as doc:
        fonts=[f for page in doc for f in page.get_fonts()]
        check(all(not page.get_images() for page in doc),'PDF raster',name)
        check(all(doc.extract_font(f[0])[3] for f in fonts),'unembedded PDF font',name)
        pdfs.append({'file':name+'.pdf','sha256':sha(c.with_suffix('.pdf')),'embedded_fonts':len(fonts)})

source_manifest=json.loads((BASE/'source_manifest.json').read_text())
for p,digest in source_manifest['source_sha256'].items():check(sha(REPO/p)==digest,'implementation/source drift',p)

# Chapter edits in this pass are navigation and the rejected-B sentence only.
chapters=sorted((REPO/'docs_codex/rtl_modules').glob('[0-9]*.md'))
with tarfile.open(ROOT/'evidence/before_c_style.tar.gz') as archive:
    for p in chapters:
        old=archive.extractfile(str(p.relative_to(REPO))).read().decode()
        expected=old.replace('[A/B 对比](../diagrams/2026-10-13_rtl_module_comparison/index.html#','[A/C 对比](../diagrams/2026-10-13_rtl_modules_native_c/index.html#').replace('B 版中同一编号的两端为同一连接。','A/C 均以连续连线表达连接，编号只用于查表。')
        check(p.read_text()==expected,'unexpected technical chapter edit',p.name)

def link(p,target):
    global links
    if target.startswith(('http:','https:','mailto:')):return
    links+=1
    path,_,frag=unquote(target).partition('#');q=(p.parent/path).resolve() if path else p
    check(q.exists(),'broken link',str(p.relative_to(REPO)),target)
    if not q.exists():return
    if frag and q.suffix=='.html':check(f'id="{frag}"' in q.read_text(),'missing HTML anchor',target)
    if frag.startswith('L') and frag[1:].isdigit():check(int(frag[1:])<=len(q.read_text().splitlines()),'invalid source line',target)
for p in chapters+[REPO/'docs_codex/rtl_modules/README.md',ROOT/'README.md',BASE/'README.md']:
    for target in re.findall(r'\]\(([^)]+)\)',p.read_text()):link(p,target)
for p in [ROOT/'index.html',BASE/'index.html']:
    for target in re.findall(r'(?:href|src)="([^"]+)"',p.read_text()):link(p,target)
    check(not re.search(r'<img[^>]*\bB/',p.read_text()),'B still displayed',p.name)

deckfile=ROOT/'C/editable/diagrams.pptx';deck=Presentation(deckfile)
check(len(deck.slides)==24,'slide count')
for slide,src in zip(deck.slides,sorted((ROOT/'C').glob('*.svg'))):
    svgtexts=Counter(e.text for e in ET.parse(src).getroot().findall('s:text',NS))
    ppttexts=Counter(s.text for s in slide.shapes if s.has_text_frame and s.text)
    check(svgtexts==ppttexts,'PPT text mismatch',src.name)
with zipfile.ZipFile(deckfile) as z:check(not any(n.startswith('ppt/media/') for n in z.namelist()),'PPT raster image')
report={'figures':len(graphs),'continuous_connections':edges,'chapters_navigation_checked':len(chapters),'source_hashes_checked':len(source_manifest['source_sha256']),'links_checked':links,'visually_checked_images':len(visual['images']),'vector_pdfs':len(pdfs),'ppt_slides':len(deck.slides),'ppt_native_objects':sum(len(s.shapes) for s in deck.slides),'errors':errors,'office_tested':False,'browser_interaction_tested':False,'rtl_simulated':False}
(ROOT/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
(ROOT/'evidence/pdf_checks.json').write_text(json.dumps(pdfs,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
raise SystemExit(bool(errors))
