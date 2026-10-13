#!/usr/bin/env python3
"""Export both SVG masters to a two-slide, 16:9 native PowerPoint deck."""
import argparse
import hashlib
import json
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
import zipfile

sys.dont_write_bytecode = True


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--directory', type=Path, default=Path(__file__).resolve().parent)
    p.add_argument('--out-dir', type=Path, required=True)
    p.add_argument('--skill', type=Path, default=Path.home()/'.codex/skills/svg-pptx-diagrams-native')
    a = p.parse_args()
    if a.out_dir.exists() and any(a.out_dir.iterdir()):
        p.error('Output directory must be new or empty')
    a.out_dir.mkdir(parents=True, exist_ok=True)
    sources = [a.directory/(n+'.svg') for n in ('soc_reference_style','soc_rearranged')]
    sys.path.insert(0, str(a.skill/'scripts'))
    from svg_to_pptx import read_svg, create_deck
    from pptx import Presentation
    out = a.out_dir/'soc_landscape_comparison.pptx'
    reports = create_deck([read_svg(s) for s in sources], out, '16:9',
                          font_file=Path('/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'),
                          slide_width_in=32)
    prs = Presentation(out)
    assert len(prs.slides) == 2
    for slide, source, report in zip(prs.slides, sources, reports):
        expected = [e.text for e in ET.parse(source).getroot().iter() if e.tag.endswith('}text')]
        actual = [s.text for s in slide.shapes if s.has_text_frame and s.text]
        assert actual == expected and len(slide.shapes) == report['objects']
        report['source'] = source.name
        report['svg_sha256'] = hashlib.sha256(source.read_bytes()).hexdigest()
    with zipfile.ZipFile(out) as z:
        assert z.testzip() is None
        assert not any(n.startswith('ppt/media/') for n in z.namelist())
    result = dict(status='PASS', aspect='16:9', slides=reports,
                  sizing='uniform scaling and centering; no stretching',
                  native_only=True, office_display_editing_tested=False,
                  pptx_sha256=hashlib.sha256(out.read_bytes()).hexdigest())
    (a.out_dir/'soc_landscape_comparison.export.json').write_text(
        json.dumps(result, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    main()
