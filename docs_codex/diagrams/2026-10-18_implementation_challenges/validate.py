#!/usr/bin/env python3
"""Artifact acceptance: input protection, links/tables, SVG/PPT native structure and layout candidates."""
from pathlib import Path
import argparse,hashlib,json,re,zipfile,subprocess,difflib
from urllib.parse import unquote
from xml.etree import ElementTree as E
from pptx import Presentation
from PIL import ImageFont
D=Path(__file__).resolve().parent;R=D.parents[2];N={'s':'http://www.w3.org/2000/svg','a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def anchors(text):
 out=set(re.findall(r'<a\s+id="([^"]+)"',text));counts={}
 for line in text.splitlines():
  if re.match(r'^#{1,6} ',line):
   s=re.sub(r'^#+\s+','',line).strip().lower();s=re.sub(r'\[([^]]+)\]\([^)]*\)',r'\1',s);s=re.sub(r'[^\w\-\s]','',s);s=re.sub(r'\s','-',s)
   n=counts.get(s,0);counts[s]=n+1;out.add(s+(f'-{n}' if n else ''))
 return out
ap=argparse.ArgumentParser();ap.add_argument('--out',type=Path,required=True);args=ap.parse_args();errors=[];report={}
inp=json.loads((D/'evidence/input.json').read_text());allowed=set(inp['allowed_existing_edits']);protected=0;changed=[]
for rel,h in inp['protected_and_input_hashes'].items():
 p=R/rel
 if rel in allowed:
  if sha(p)!=h:changed.append(rel)
 else:
  if not p.exists() or sha(p)!=h:errors.append('Protected input changed: '+rel)
  protected+=1
report.update(protected_files_checked=protected,allowed_existing_files_changed=changed)
for rel in changed:
 old=(D/'evidence/before'/rel).read_text().splitlines(True);new=(R/rel).read_text().splitlines(True)
 # These updates are insertions only; no pre-existing text may be removed.
 diff=list(difflib.unified_diff(old,new,fromfile='before/'+rel,tofile=rel));assert not [l for l in diff if l.startswith('-') and not l.startswith('---')],rel
 (D/'evidence'/('diff-'+rel.replace('/','_')+'.txt')).write_text(''.join(diff))
docs=[R/p for p in allowed]+[R/'docs_codex/engineering/Implementation_Challenges_and_Open_Questions.md',R/'docs_codex/handoffs/L01_implementation_challenges.md',D/'README.md',D/'evidence/README.md']
link_count=0;table_rows=0
for p in docs:
 s=p.read_text();table_cols=None
 for line in s.splitlines():
  if line.startswith('|') and line.rstrip().endswith('|'):
   cols=len(re.split(r'(?<!\\)\|',line))-2
   if table_cols is None:table_cols=cols
   if cols!=table_cols:errors.append(f'Table width: {p.name}: {line}')
   table_rows+=1
  else:table_cols=None
 for target in re.findall(r'!?\[[^\]]*\]\(([^)]+)\)',s):
  if target.startswith(('http:','https:','mailto:')):continue
  link_count+=1;path,_,frag=target.partition('#');dest=(p.parent/unquote(path)).resolve() if path else p
  if not dest.exists():errors.append(f'Broken link {p.relative_to(R)} -> {target}')
  elif frag and dest.suffix=='.md' and unquote(frag) not in anchors(dest.read_text()):errors.append(f'Broken anchor {p.name} -> {target}')
report.update(markdown_files=len(docs),local_links=link_count,table_rows=table_rows)
master=json.loads((D/'masters/figures.json').read_text());prs=Presentation(D/'export/diagrams.pptx');assert len(master)==len(prs.slides)==13
main=(R/'docs_codex/engineering/Implementation_Challenges_and_Open_Questions.md').read_text();assert len(re.findall('^## [0-9]{2}',main,re.M))==13
assert len(re.findall(r'!\[',main))==13
fontpath='/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc';pages=[];layout=[]
for spec,slide in zip(master,prs.slides):
 p=D/'masters'/(spec['file']+'.svg');svg=E.parse(p).getroot();texts=svg.findall('s:text',N);actual=[s.text for s in slide.shapes if s.has_text_frame and s.text]
 assert actual==[t.text for t in texts],p.name
 assert (D/'png'/(spec['file']+'.png')).is_file()
 assert (D/'evidence/final-ppt-preview'/f'slide-{spec["number"]:02d}.png').is_file()
 assert f'第 {spec["number"]} 页' in main
 shapes=len(slide.shapes);expected=sum(1 for el in svg if el.tag.split('}')[-1] not in ['defs','title'])-1;assert shapes==expected
 arrows=sum(1 for s in slide.shapes if s._element.find('.//{'+N['a']+'}tailEnd') is not None)
 source_arrows=sum(1 for el in svg if el.get('marker-end'));assert arrows==source_arrows
 dash=sum(1 for s in slide.shapes if s._element.find('.//{'+N['a']+'}custDash') is not None);assert dash==sum(1 for el in svg if el.get('stroke-dasharray'))
 connectors=sum(1 for s in slide.shapes if s._element.tag=='{'+N['p']+'}cxnSp')
 boxes=[]
 for t in texts:
  size=float(t.get('font-size'));font=ImageFont.truetype(fontpath,round(size*10));b=font.getbbox(t.text,anchor='ls');b=[v/10 for v in b];adv=font.getlength(t.text)/10
  x,y=float(t.get('x')),float(t.get('y'));anchor=t.get('text-anchor','start');x-=adv*{'start':0,'middle':.5,'end':1}[anchor];box=(x+b[0],y+b[1],x+b[2],y+b[3])
  if box[0]<0 or box[1]<0 or box[2]>1600 or box[3]>900:layout.append({'figure':p.name,'kind':'text-outside','text':t.text,'box':box})
  boxes.append((t.text,box))
 for edge in spec['edges']:
  for (x1,y1),(x2,y2) in zip(edge['points'],edge['points'][1:]):
   for text,b in boxes:
    # Inset by 1px to avoid reporting boundary touches as glyph intersections.
    l,t,r,bottom=b;l+=1;r-=1;t+=1;bottom-=1
    cross=(x1==x2 and l<x1<r and max(min(y1,y2),t)<min(max(y1,y2),bottom)) or (y1==y2 and t<y1<bottom and max(min(x1,x2),l)<min(max(x1,x2),r))
    if cross:layout.append({'figure':p.name,'kind':'line-text-candidate','edge':edge['id'],'text':text})
 pages.append({'page':spec['number'],'objects':shapes,'texts':len(texts),'arrow_objects':arrows,'dashed_objects':dash,'native_connectors':connectors})
with zipfile.ZipFile(D/'export/diagrams.pptx') as z:
 assert z.testzip() is None;assert not [n for n in z.namelist() if n.startswith('ppt/media/')]
report.update(pages=pages,total_objects=sum(p['objects'] for p in pages),total_texts=sum(p['texts'] for p in pages),total_arrows=sum(p['arrow_objects'] for p in pages),layout_candidates=layout,pptx_sha256=sha(D/'export/diagrams.pptx'),office_tested=False,rtl_or_board_tested=False)
source=json.loads((D/'evidence/source-check.json').read_text())
for rel,h in source['files'].items():
 if sha(R/rel)!=h:errors.append('Source changed after read: '+rel)
report['targeted_source_hashes_checked']=len(source['files'])
rc=subprocess.run(['git','diff','--check'],cwd=R,capture_output=True,text=True);report['git_diff_check']={'exit_code':rc.returncode,'output':rc.stdout+rc.stderr}
if rc.returncode:errors.append('git diff --check failed')
# New files are untracked, so check their Markdown trailing whitespace separately.
for p in docs:
 for i,l in enumerate(p.read_text().splitlines(),1):
  if l.endswith((' ','\t')):errors.append(f'Trailing whitespace: {p.name}:{i}')
report['errors']=errors;report['passed']=not errors and not layout
report['deliverable_hashes']={str(p.relative_to(D)):sha(p) for folder in ['masters','png','export'] for p in sorted((D/folder).iterdir()) if p.is_file()}
args.out.parent.mkdir(parents=True,exist_ok=True);args.out.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:report[k] for k in ['protected_files_checked','markdown_files','local_links','table_rows','total_objects','total_texts','total_arrows','layout_candidates','errors','passed']},ensure_ascii=False,indent=2))
raise SystemExit(0 if report['passed'] else 1)
